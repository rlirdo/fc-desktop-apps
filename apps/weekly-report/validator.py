# -*- coding: utf-8 -*-
"""
週報成品驗收（可對任何一份 .pptx 跑）
====================================
六項必檢（來源：weekly-report-workshop/practice.html 的驗收錨點）：
  1. 中文序號（壹、貳、參、肆、伍、陸）
  2. 「本週無進度」原文（一字不差）
  3. 總覽頁列出全部區塊
  4. 「問題與請益」頁
  5. 「下週計畫」頁
  6. 每一頁都有備忘稿逐字稿
另外附帶兩項：頁數 11–15、每頁標題字級 ≥ 30pt。
"""
from pptx import Presentation
from pptx.util import Pt

CN_NUM = ["壹", "貳", "參", "肆", "伍", "陸"]
NO_PROGRESS = "本週無進度"
MIN_PAGES, MAX_PAGES = 11, 15
TITLE_MIN_PT = 30.0


def _slide_text(sl):
    buf = []
    for sh in sl.shapes:
        if sh.has_text_frame:
            buf.append(sh.text_frame.text)
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    buf.append(c.text)
    return "\n".join(buf)


def _title(sl):
    """頁標題＝這一頁字級最大（且 ≥ 24pt）的那一段文字。"""
    best, best_pt = "", 0.0
    for sh in sl.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                pt = r.font.size.pt if r.font.size is not None else 0.0
                if pt >= 24.0 and pt > best_pt and r.text.strip():
                    best, best_pt = r.text.strip(), pt
    return best


def _max_pt(sl):
    m = 0.0
    for sh in sl.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size is not None:
                    m = max(m, r.font.size.pt)
    return m


def _notes_text(sl):
    try:
        if not sl.has_notes_slide:
            return ""
        return (sl.notes_slide.notes_text_frame.text or "").strip()
    except Exception:
        return ""


def validate(path):
    """回傳 (checks, summary)。checks = [{key,name,ok,detail}]（前 6 項為必檢）。"""
    prs = Presentation(path)
    slides = list(prs.slides)
    texts = [_slide_text(s) for s in slides]
    titles = [_title(s) for s in slides]
    alltext = "\n".join(texts)
    notes = [_notes_text(s) for s in slides]
    checks = []

    def add(key, name, ok, detail):
        checks.append({"key": key, "name": name, "ok": bool(ok), "detail": detail})

    # 1 中文序號
    hit = [c for c in CN_NUM if c in alltext]
    add("c1", "中文序號（壹貳參…）", len(hit) >= 3,
        f"找到 {len(hit)} 個中文序號：{'、'.join(hit) if hit else '無'}（需 ≥ 3）")

    # 2 「本週無進度」原文
    add("c2", f"「{NO_PROGRESS}」原文", NO_PROGRESS in alltext,
        "有找到一字不差的原文" if NO_PROGRESS in alltext else "沒有找到；沒進度的區塊必須原文寫這五個字")

    # 3 總覽頁列出全部區塊
    ov_i, ov_hit = -1, []
    for i, t in enumerate(texts):
        h = [c for c in CN_NUM if c in t]
        if "總覽" in titles[i] and len(h) > len(ov_hit):
            ov_i, ov_hit = i, h
    if ov_i < 0:                       # 沒有標「總覽」時，退而找序號最齊的一頁
        for i, t in enumerate(texts):
            h = [c for c in CN_NUM if c in t]
            if len(h) > len(ov_hit):
                ov_i, ov_hit = i, h
    add("c3", "總覽頁列出全部區塊", ov_i >= 0 and len(ov_hit) >= 3,
        f"第 {ov_i + 1} 頁列出 {len(ov_hit)} 個區塊：{'、'.join(ov_hit)}" if ov_i >= 0
        else "找不到總覽頁")

    # 4 / 5 兩張必備頁
    for key, kw in (("c4", "問題與請益"), ("c5", "下週計畫")):
        pg = [i + 1 for i, t in enumerate(titles) if kw in t]
        add(key, f"「{kw}」頁", bool(pg),
            f"第 {pg[0]} 頁（以頁標題認定）" if pg
            else f"找不到標題為「{kw}」的頁（必備頁，不可省略）")

    # 6 每頁備忘稿
    empty = [i + 1 for i, t in enumerate(notes) if not t]
    add("c6", "每一頁都有備忘稿逐字稿", slides and not empty,
        f"{len(slides)} 頁全部都有逐字稿" if slides and not empty
        else f"第 {', '.join(map(str, empty))} 頁沒有備忘稿")

    # 附帶：頁數、標題字級
    n = len(slides)
    add("x1", f"頁數 {MIN_PAGES}–{MAX_PAGES} 頁", MIN_PAGES <= n <= MAX_PAGES, f"共 {n} 頁")
    bad = [i + 1 for i, s in enumerate(slides) if _max_pt(s) < TITLE_MIN_PT]
    add("x2", f"每頁標題 ≥ {int(TITLE_MIN_PT)}pt", not bad,
        "全部合格" if not bad else f"第 {', '.join(map(str, bad))} 頁最大字級不足 30pt")

    core6 = checks[:6]
    summary = {
        "path": path,
        "slides": n,
        "passed": sum(1 for c in core6 if c["ok"]),
        "total": len(core6),
        "all_ok": all(c["ok"] for c in checks),
    }
    return checks, summary


def format_report(checks, summary):
    L = [f"驗收檔案：{summary['path']}",
         f"投影片頁數：{summary['slides']} 頁",
         f"六項必檢：{summary['passed']} / {summary['total']} 通過", ""]
    for c in checks:
        L.append(("[通過] " if c["ok"] else "[未過] ") + c["name"] + "　— " + c["detail"])
    L.append("")
    L.append("全部項目通過。" if summary["all_ok"] else "尚有項目未通過，請依上面說明修正後重新產生。")
    return "\n".join(L)
