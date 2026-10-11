"""匯出：docx／pdf／xlsx（文獻清單與指標、缺口矩陣）／ris／bib／趨勢圖 png／查詢紀錄 json。

export_all(papers, gap, out_dir, name, options, progress) -> {path: ok(bool)}
  name：「文獻缺口_<關鍵字>_<YYYYMMDD>」（可用 default_name() 產生）；檔案寫入 out_dir\\name\\ 之下。
  options：
    group_by_chapter(bool)  依章節標籤分組（預設 False）
    order "apa"|"custom"    apa＝先中文筆畫後英文字母；custom＝依 Paper.order（拖曳排序）
    only_included(bool)     只匯出勾選納入者（預設 True）
    formats [..]            省略＝全部："docx","pdf","xlsx","gap_xlsx","ris","bib","png","json"
    keywords, settings, mailto, chapters   寫入查詢紀錄與檢索策略
    pdf_engine "auto"|"word"|"reportlab"
"""
from __future__ import annotations

import datetime as _dt
import json
import math
import os
import re
from typing import Optional

from .apa import format_apa7, format_apa7_runs, in_text, sort_key, clean_text, dash_range
from .models import Paper, Progress, VERSION, report, check_cancel, INDICATOR_NAMES, TYPE_LABELS, DEFAULT_CHAPTERS

FILES = {
    "docx": "參考文獻_APA7.docx",
    "pdf": "參考文獻_APA7.pdf",
    "xlsx": "文獻清單與指標.xlsx",
    "gap_xlsx": "缺口矩陣.xlsx",
    "png": "趨勢圖.png",
    "ris": "references.ris",
    "bib": "references.bib",
    "json": "查詢紀錄.json",
}
MANUAL_NOTE = "使用者自填，未經 API 驗證"
IND_ORDER = ["doi", "oa", "type", "year", "issn", "lang"]
IND_SHORT = {"doi": "①DOI", "oa": "②OA", "type": "③類型", "year": "⑥年份", "issn": "⑦ISSN/DOAJ", "lang": "⑩語言"}


def safe_name(s: str, n: int = 40) -> str:
    s = re.sub(r'[\\/:*?"<>|\r\n\t]+', " ", s or "").strip()
    s = re.sub(r"\s+", " ", s)
    return s[:n].strip() or "未命名"


def default_name(keywords, date: Optional[_dt.date] = None) -> str:
    first = []
    for g in keywords or []:
        g = g if isinstance(g, str) else ";".join(g)
        t = re.split(r"[;；]", g)[0].strip()
        if t:
            first.append(t)
    kw = safe_name("+".join(first[:3]), 40)
    return f"文獻缺口_{kw}_{(date or _dt.date.today()).strftime('%Y%m%d')}"


def ordered(papers: list, options: dict) -> list:
    ps = [p for p in papers if p.include or not options.get("only_included", True)]
    if options.get("order") == "custom":
        return sorted(ps, key=lambda p: p.order)
    return sorted(ps, key=sort_key)


def _ind_mark(p: Paper, k: str) -> str:
    v = p.indicators.get(k)
    if not v:
        return "—"
    val = v[0] if isinstance(v, (list, tuple)) else v
    return "✔" if val is True else ("✘" if val is False else "？")


def _ind_text(p: Paper, k: str) -> str:
    v = p.indicators.get(k)
    if not v:
        return ""
    val, note = (v[0], v[1]) if isinstance(v, (list, tuple)) else (v, "")
    mark = "✔" if val is True else ("✘" if val is False else "？")
    return f"{mark} {note}".strip()


# ================================================================ DOCX
def _set_font(run, size=12, italic=False, bold=False, color=None):
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.italic = italic
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "標楷體")


def _heading(doc, text, size=14, center=False):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    para = doc.add_paragraph()
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = para.add_run(text)
    _set_font(r, size=size, bold=True)
    return para


def _ref_para(doc, p: Paper):
    from docx.shared import Cm, Pt
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.left_indent = Cm(1.27)
    pf.first_line_indent = Cm(-1.27)
    pf.line_spacing = 2.0
    pf.space_after = Pt(0)
    for text, st in format_apa7_runs(p):
        r = para.add_run(text)
        _set_font(r, 12, italic=(st == "i"), bold=(st == "b"))
    if p.source == "manual":
        r = para.add_run(f"　〔{MANUAL_NOTE}〕")
        _set_font(r, 9, color=(0x80, 0x80, 0x80))
    return para


def _cell(cell, text, size=10, bold=False):
    cell.text = ""
    r = cell.paragraphs[0].add_run(str(text))
    _set_font(r, size, bold=bold)


def write_docx(papers: list, path: str, options: dict, gap: Optional[dict] = None) -> None:
    import docx
    from docx.shared import Cm, Pt
    from docx.enum.section import WD_ORIENT  # noqa: F401
    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.54))
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "標楷體")

    _heading(doc, "參考文獻", 14, center=True)
    ps = ordered(papers, options)
    if options.get("group_by_chapter"):
        chapters = list(options.get("chapters") or DEFAULT_CHAPTERS)
        for c in sorted({p.chapter or "未分類" for p in ps}):
            if c not in chapters:
                chapters.append(c)
        for c in chapters:
            sub = [p for p in ps if (p.chapter or "未分類") == c]
            if not sub:
                continue
            _heading(doc, c, 12)
            for p in sub:
                _ref_para(doc, p)
    else:
        for p in ps:
            _ref_para(doc, p)

    # 附錄一：指標證據表
    doc.add_page_break()
    _heading(doc, "附錄一　指標證據表", 14)
    para = doc.add_paragraph()
    _set_font(para.add_run("✔＝通過　✘＝未通過　？＝待人工確認；判定理由見「文獻清單與指標.xlsx」。手動＝使用者自填，未經 API 驗證。"), 9)
    cols = ["序", "文內引用", "DOI"] + [IND_SHORT[k] for k in IND_ORDER] + ["⑫相關性", "來源"]
    tbl = doc.add_table(rows=1, cols=len(cols))
    tbl.style = "Table Grid"
    for i, c in enumerate(cols):
        _cell(tbl.rows[0].cells[i], c, 9, True)
    for n, p in enumerate(ps, 1):
        row = tbl.add_row().cells
        vals = [n, in_text(p), p.doi or "—"] + [_ind_mark(p, k) for k in IND_ORDER] + [
            f"{p.relevance:.1f}", "手動" if p.source == "manual" else p.source]
        for i, v in enumerate(vals):
            _cell(row[i], v, 8)

    # 附錄二：檢索策略
    doc.add_paragraph()
    _heading(doc, "附錄二　檢索策略", 14)
    for line in strategy_lines(papers, gap, options):
        para = doc.add_paragraph()
        r = para.add_run(line)
        _set_font(r, 11)
    doc.save(path)


def strategy_lines(papers: list, gap: Optional[dict], options: dict) -> list:
    s = options.get("settings") or {}
    kw = options.get("keywords") or []
    lines = [
        f"檢索日期：{_dt.datetime.now().strftime('%Y/%m/%d %H:%M')}",
        "資料庫：Crossref REST API、OpenAlex；補充：Unpaywall（OA）、Semantic Scholar（被引／TLDR）、doi.org（DOI 解析驗證）",
        "關鍵字組（組內同義詞以 OR、組間以 AND 結合）：" + "；".join(f"〔{g if isinstance(g, str) else ' OR '.join(g)}〕" for g in kw),
        f"出版年：{s.get('year_from') or '不限'}–{s.get('year_to') or '今'}" + ("（含經典不限年）" if s.get("classic") else ""),
        f"語言：{ {'en': '英文', 'zh': '中文', 'any': '不限'}.get(s.get('lang', 'any'), s.get('lang', '不限')) }",
        "文獻類型：" + "、".join(TYPE_LABELS.get(t, t) for t in (s.get("types") or ["journal-article"])),
        "合格指標：" + "、".join(INDICATOR_NAMES[k] for k in IND_ORDER if (s.get("checks") or {}).get(k, True)),
        f"取得筆數：{len(papers)}；納入：{sum(1 for p in papers if p.include)}；手動加入：{sum(1 for p in papers if p.source == 'manual')}",
        "說明：手動加入之中文文獻為使用者自填，未經 API 驗證；其餘書目均由 API 回傳原始資料產生。",
    ]
    if gap:
        lines.append(f"缺口矩陣：OpenAlex 兩兩／三三組合命中篇數，門檻 < {gap.get('threshold')} 篇標為缺口。")
    return lines


# ================================================================ PDF
def docx_to_pdf_word(docx_path: str, pdf_path: str) -> bool:
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return False
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        d = word.Documents.Open(os.path.abspath(docx_path), ReadOnly=True, AddToRecentFiles=False)
        d.SaveAs2(os.path.abspath(pdf_path), FileFormat=17)
        d.Close(False)
        return os.path.isfile(pdf_path)
    except Exception:
        return False
    finally:
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()


def _cjk_font_path() -> Optional[str]:
    win = os.environ.get("WINDIR", r"C:\Windows")
    for f in ("kaiu.ttf", "msjh.ttc", "msjh.ttf", "mingliu.ttc"):
        p = os.path.join(win, "Fonts", f)
        if os.path.isfile(p):
            return p
    for p in ("/Library/Fonts/Arial Unicode.ttf", "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
              "/usr/share/fonts/truetype/arphic/ukai.ttc", "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"):
        if os.path.isfile(p):
            return p
    return None


def write_pdf_reportlab(papers: list, path: str, options: dict, gap: Optional[dict] = None) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from xml.sax.saxutils import escape

    cjk = "Helvetica"
    fp = _cjk_font_path()
    if fp:
        try:
            if fp.lower().endswith(".ttc"):
                pdfmetrics.registerFont(TTFont("FCCJK", fp, subfontIndex=0))
            else:
                pdfmetrics.registerFont(TTFont("FCCJK", fp))
            cjk = "FCCJK"
        except Exception:
            cjk = "Helvetica"

    lat = {"": "Times-Roman", "i": "Times-Italic", "b": "Times-Bold"}
    fdir = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
    try:
        tt = {"": "times.ttf", "i": "timesi.ttf", "b": "timesbd.ttf"}
        if all(os.path.isfile(os.path.join(fdir, f)) for f in tt.values()):
            for k, f in tt.items():
                pdfmetrics.registerFont(TTFont("FCTimes" + k, os.path.join(fdir, f)))
            lat = {k: "FCTimes" + k for k in tt}
    except Exception:
        pass

    def markup(text: str, st: str) -> str:
        out = []
        for seg in re.findall(r"[\u3000-\u303f\u3400-\u9fff\uff00-\uffef]+|[^\u3000-\u303f\u3400-\u9fff\uff00-\uffef]+", text):
            if re.match(r"[\u3000-\u303f\u3400-\u9fff\uff00-\uffef]", seg):
                f = cjk
                s = escape(seg)
                out.append(f'<font face="{f}">{"<b>" + s + "</b>" if st == "b" else s}</font>')
            else:
                f = lat.get(st, lat[""])
                out.append(f'<font face="{f}">{escape(seg)}</font>')
        return "".join(out)

    ref = ParagraphStyle("ref", fontName=lat[""], fontSize=12, leading=20, leftIndent=1.27 * cm,
                         firstLineIndent=-1.27 * cm, spaceAfter=4)
    h = ParagraphStyle("h", fontName=cjk, fontSize=14, leading=20, spaceAfter=8, alignment=1)
    h2 = ParagraphStyle("h2", fontName=cjk, fontSize=12, leading=18, spaceAfter=6)
    body = ParagraphStyle("b", fontName=lat[""], fontSize=9, leading=13)
    story = [Paragraph(markup("參考文獻", ""), h)]
    ps = ordered(papers, options)

    def add(p):
        m = "".join(markup(t, st) for t, st in format_apa7_runs(p))
        if p.source == "manual":
            m += markup(f"　〔{MANUAL_NOTE}〕", "")
        story.append(Paragraph(m, ref))

    if options.get("group_by_chapter"):
        chapters = list(options.get("chapters") or DEFAULT_CHAPTERS)
        for c in sorted({p.chapter or "未分類" for p in ps}):
            if c not in chapters:
                chapters.append(c)
        for c in chapters:
            sub = [p for p in ps if (p.chapter or "未分類") == c]
            if sub:
                story.append(Paragraph(markup(c, ""), h2))
                for p in sub:
                    add(p)
    else:
        for p in ps:
            add(p)
    story += [PageBreak(), Paragraph(markup("附錄一　指標證據表", ""), h2)]
    for n, p in enumerate(ps, 1):
        line = f"{n}. {in_text(p)}｜DOI {p.doi or '—'}｜" + "｜".join(_ind_text(p, k) for k in IND_ORDER) + f"｜⑫{p.relevance:.2f}"
        story.append(Paragraph(markup(line, ""), body))
    story += [Spacer(1, 12), Paragraph(markup("附錄二　檢索策略", ""), h2)]
    for line in strategy_lines(papers, gap, options):
        story.append(Paragraph(markup(line, ""), body))
    SimpleDocTemplate(path, pagesize=A4, leftMargin=2.54 * cm, rightMargin=2.54 * cm, topMargin=2.54 * cm,
                      bottomMargin=2.54 * cm, title="參考文獻 APA 7").build(story)


# ================================================================ XLSX
def _autosize(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def write_xlsx(papers: list, path: str, options: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "文獻清單與指標"
    head = ["序", "納入", "章節", "相關性", "年", "作者", "題名", "刊名", "卷", "期", "頁／文章號", "類型", "DOI",
            "DOI 解析", "OA 狀態", "OA PDF", "授權", "DOAJ", "ISSN", "出版社", "語言", "被引", "TLDR"] + \
           [IND_SHORT[k] for k in IND_ORDER] + ["全部通過", "APA 7", "文內引用", "來源", "備註"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1C7293")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ps = ordered(papers, dict(options, only_included=False))
    for n, p in enumerate(ps, 1):
        au = "; ".join(f"{a.get('family', '')}, {a.get('given', '')}".strip(", ") for a in p.authors)
        ws.append([n, "✔" if p.include else "", p.chapter, round(p.relevance, 3), p.year, au, clean_text(p.title),
                   p.container, p.volume, p.issue, dash_range(p.pages) or (f"Article {p.article_number}" if p.article_number else ""),
                   TYPE_LABELS.get(p.type, p.type), p.doi, {True: "可解析", False: "失敗", None: "未驗證"}[p.doi_ok],
                   p.oa_status, p.oa_pdf_url, p.oa_license, {True: "是", False: "否", None: ""}[p.doaj],
                   ", ".join(p.issn), p.publisher, p.language, p.cited_by, p.tldr] +
                  [_ind_text(p, k) for k in IND_ORDER] +
                  ["是" if p.passes_all() else "否", format_apa7(p), in_text(p), p.source,
                   (MANUAL_NOTE + " " if p.source == "manual" else "") + (p.note if p.note != MANUAL_NOTE else "")])
        if p.doi:
            ws.cell(row=n + 1, column=13).hyperlink = p.doi_url
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autosize(ws, {"A": 5, "B": 5, "C": 10, "D": 8, "E": 6, "F": 30, "G": 60, "H": 28, "M": 28, "AE": 80})
    wb.save(path)


def write_gap_xlsx(gap: dict, path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from .gap import matrix_table
    wb = Workbook()
    ws = wb.active
    ws.title = "兩兩矩陣"
    G = gap.get("groups") or []
    labels = [f"{i + 1}. " + " / ".join(g) for i, g in enumerate(G)]
    thr = gap.get("threshold", 5)
    ws.append([f"OpenAlex 命中篇數（對角線＝單組；紅底＝< {thr} 篇缺口）"] + labels)
    M = matrix_table(gap)
    red = PatternFill("solid", fgColor="F8D7DA")
    for i, row in enumerate(M):
        ws.append([labels[i]] + [("" if v is None else v) for v in row])
        for j, v in enumerate(row):
            if i != j and v is not None and v < thr:
                ws.cell(row=i + 2, column=j + 2).fill = red
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 40
    for j in range(len(G)):
        ws.column_dimensions[chr(66 + j)].width = 22

    ws2 = wb.create_sheet("組合清單")
    ws2.append(["組合", "組數", "命中篇數", "缺口", "OpenAlex 查詢式"])
    for kind, rows in (("兩兩", gap.get("pairs") or []), ("三三", gap.get("triples") or [])):
        for r in rows:
            ws2.append([" × ".join(str(i + 1) for i in r["groups"]), kind, r.get("count"),
                        "缺口" if r.get("gap") else ("" if r.get("gap") is not None else "查詢失敗"), r.get("query")])
            if r.get("gap"):
                for c in ws2[ws2.max_row]:
                    c.fill = red
    ws2.column_dimensions["E"].width = 90

    ws3 = wb.create_sheet("各組合前10篇")
    ws3.append(["組合", "名次", "年", "第一作者", "題名", "刊名", "被引", "DOI"])
    for r in (gap.get("pairs") or []) + (gap.get("triples") or []):
        for k, p in enumerate(r.get("top") or [], 1):
            ws3.append([" × ".join(str(i + 1) for i in r["groups"]), k, p.year, p.first_author_family,
                        clean_text(p.title), p.container, p.cited_by, p.doi])
    ws3.column_dimensions["E"].width = 70

    ws4 = wb.create_sheet("年度趨勢")
    tb = gap.get("trend_by_group") or {}
    years = sorted({y for d in tb.values() for y in d})
    ws4.append(["年"] + [labels[int(i)] for i in tb])
    for y in years:
        ws4.append([y] + [tb[i].get(y, 0) for i in tb])
    ws5 = wb.create_sheet("說明")
    for line in [f"資料來源：{gap.get('source', 'OpenAlex')}", f"門檻：< {thr} 篇標為缺口",
                 f"耗時：{gap.get('elapsed_sec')} 秒", "錯誤：" + ("；".join(gap.get("errors") or []) or "無")]:
        ws5.append([line])
    wb.save(path)


# ================================================================ RIS / BIB
_RIS_TY = {"journal-article": "JOUR", "proceedings-article": "CPAPER", "book-chapter": "CHAP", "book": "BOOK",
           "patent": "PAT", "webpage": "ELEC", "posted-content": "UNPB", "dissertation": "THES", "report": "RPRT"}


def write_ris(papers: list, path: str, options: dict) -> None:
    lines = []
    for p in ordered(papers, options):
        lines.append(f"TY  - {_RIS_TY.get(p.type, 'GEN')}")
        for a in p.authors:
            lines.append(f"AU  - {a.get('family', '')}" + (f", {a.get('given')}" if a.get("given") else ""))
        lines.append(f"TI  - {clean_text(p.title)}")
        if p.container:
            lines.append(f"{'T2' if p.type in ('book-chapter', 'proceedings-article') else 'JO'}  - {p.container}")
        if p.year:
            lines.append(f"PY  - {p.year}")
        if p.volume:
            lines.append(f"VL  - {p.volume}")
        if p.issue:
            lines.append(f"IS  - {p.issue}")
        if p.pages:
            sp, _, ep = dash_range(p.pages).partition("–")
            lines.append(f"SP  - {sp}")
            if ep:
                lines.append(f"EP  - {ep}")
        elif p.article_number:
            lines.append(f"SP  - {p.article_number}")
        for s in p.issn:
            lines.append(f"SN  - {s}")
        if p.publisher:
            lines.append(f"PB  - {p.publisher}")
        if p.doi:
            lines.append(f"DO  - {p.doi}")
            lines.append(f"UR  - {p.doi_url}")
        elif p.url:
            lines.append(f"UR  - {p.url}")
        if p.patent_number:
            lines.append(f"M1  - {p.patent_number}")
        if p.language:
            lines.append(f"LA  - {p.language}")
        if p.abstract:
            lines.append(f"AB  - {clean_text(p.abstract)}")
        if p.source == "manual":
            lines.append(f"N1  - {MANUAL_NOTE}")
        if p.chapter:
            lines.append(f"KW  - {p.chapter}")
        lines.append("ER  - ")
        lines.append("")
    with open(path, "w", encoding="utf-8-sig", newline="\r\n") as f:
        f.write("\n".join(lines))


def _bib_escape(s: str) -> str:
    s = clean_text(s)
    return s.replace("\\", "\\textbackslash{}").replace("{", "\\{").replace("}", "\\}").replace("&", "\\&") \
        .replace("%", "\\%").replace("$", "\\$").replace("#", "\\#").replace("_", "\\_")


_BIB_TY = {"journal-article": "article", "proceedings-article": "inproceedings", "book-chapter": "incollection",
           "book": "book", "patent": "patent", "webpage": "online", "dissertation": "phdthesis", "report": "techreport"}


def write_bib(papers: list, path: str, options: dict) -> None:
    used = set()
    out = []
    for p in ordered(papers, options):
        fam = re.sub(r"[^A-Za-z]", "", p.first_author_family.encode("ascii", "ignore").decode()) or "ref"
        base = f"{fam.lower()}{p.year or 'nd'}"
        key, k = base, 0
        while key in used:
            k += 1
            key = base + "abcdefghijklmnopqrstuvwxyz"[(k - 1) % 26]
        used.add(key)
        f = {}
        f["author"] = " and ".join(
            "{" + a.get("family", "") + "}" if not a.get("given") else f"{a.get('family', '')}, {a.get('given')}"
            for a in p.authors)
        f["title"] = "{" + _bib_escape(p.title) + "}"
        if p.type == "book-chapter":
            f["booktitle"] = _bib_escape(p.book_title or p.container)
        elif p.type == "proceedings-article":
            f["booktitle"] = _bib_escape(p.container)
        elif p.container:
            f["journal"] = _bib_escape(p.container)
        if p.year:
            f["year"] = str(p.year)
        for fld in ("volume",):
            if getattr(p, fld):
                f[fld] = getattr(p, fld)
        if p.issue:
            f["number"] = p.issue
        if p.pages:
            f["pages"] = dash_range(p.pages).replace("–", "--")
        elif p.article_number:
            f["eid"] = p.article_number
        if p.publisher:
            f["publisher"] = _bib_escape(p.publisher)
        if p.issn:
            f["issn"] = ", ".join(p.issn)
        if p.doi:
            f["doi"] = p.doi
        if p.url and not p.doi:
            f["url"] = p.url
        if p.patent_number:
            f["number"] = p.patent_number
        if p.language:
            f["langid"] = p.language
        if p.source == "manual":
            f["note"] = MANUAL_NOTE
        body = ",\n".join(f"  {k2} = {{{v}}}" if not v.startswith("{") else f"  {k2} = {v}" for k2, v in f.items() if v)
        out.append(f"@{_BIB_TY.get(p.type, 'misc')}{{{key},\n{body}\n}}\n")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(out))


# ================================================================ PNG（PIL 自繪）
def _font(size: int, bold: bool = False):
    from PIL import ImageFont
    win = os.environ.get("WINDIR", r"C:\Windows")
    for f in (("msjhbd.ttc", "msjh.ttc") if bold else ("msjh.ttc",)) + ("kaiu.ttf", "arial.ttf"):
        p = os.path.join(win, "Fonts", f)
        if os.path.isfile(p):
            try:
                return ImageFont.truetype(p, size)
            except OSError:
                continue
    for f in ("/System/Library/Fonts/PingFang.ttc", "/System/Library/Fonts/STHeiti Light.ttc",
              "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        if os.path.isfile(f):
            try:
                return ImageFont.truetype(f, size)
            except OSError:
                continue
    return ImageFont.load_default()


PALETTE = ["#1C7293", "#2E7D5B", "#C99A3E", "#065A82", "#5FA98A", "#0B1F3A"]


def write_trend_png(papers: list, gap: Optional[dict], path: str) -> None:
    from PIL import Image, ImageDraw
    W, H = 1600, 1000
    img = Image.new("RGB", (W, H), "#F5F9FA")
    d = ImageDraw.Draw(img)
    f_t, f_l, f_s = _font(34, True), _font(22), _font(18)
    ink, grey = "#1E293B", "#64748B"
    d.text((60, 30), f"文獻趨勢：近 10 個完整年度（{_dt.date.today().year - 10}–{_dt.date.today().year - 1}）每年篇數　＆　前 10 高被引",
           font=f_t, fill=ink)

    # ---- 上：折線
    this = _dt.date.today().year - 1           # 近 10 個完整年度
    series = []
    from_gap = bool(gap and gap.get("trend_by_group"))
    if from_gap:
        for i, tr in gap["trend_by_group"].items():
            g = (gap.get("groups") or [])[int(i)] if gap.get("groups") else [str(i)]
            series.append((f"{int(i) + 1}. {g[0]}", {int(y): c for y, c in tr.items()}))
    else:
        cnt = {y: 0 for y in range(this - 9, this + 1)}
        for p in papers:
            if p.year in cnt:
                cnt[p.year] += 1
        series.append(("本次檢索結果", cnt))
    years = list(range(this - 9, this + 1))
    x0, y0, x1, y1 = 140, 110, W - 440, 520
    vmax = max([1] + [v for _, s in series for v in s.values()])
    log = vmax > 50 * max(1, min([max(s.values()) or 1 for _, s in series]))   # 量級差太大改對數
    def ty(v):
        if log:
            return y1 - (y1 - y0) * (math.log10(1 + v) / math.log10(1 + vmax))
        return y1 - (y1 - y0) * v / vmax
    d.rectangle((x0, y0, x1, y1), fill="#FFFFFF", outline="#CBD5E1")
    if not log and vmax <= 4:
        ticks = list(range(0, int(vmax) + 1))
    else:
        ticks = [vmax * k / 4 if not log else (10 ** (math.log10(1 + vmax) * k / 4) - 1) for k in range(5)]
    for v in ticks:
        yy = ty(v)
        d.line((x0, yy, x1, yy), fill="#E2E8F0")
        d.text((x0 - 10, yy), f"{int(round(v)):,}", font=f_s, fill=grey, anchor="rm")
    step = (x1 - x0 - 40) / max(1, len(years) - 1)
    for i, y in enumerate(years):
        xx = x0 + 20 + i * step
        d.text((xx, y1 + 12), str(y), font=f_s, fill=grey, anchor="mt")
    for si, (name, s) in enumerate(series):
        col = PALETTE[si % len(PALETTE)]
        pts = [(x0 + 20 + i * step, ty(s.get(y, 0))) for i, y in enumerate(years)]
        d.line(pts, fill=col, width=4)
        for (px, py) in pts:
            d.ellipse((px - 6, py - 6, px + 6, py + 6), fill=col)
        ly = y0 + 10 + si * 34
        d.rectangle((x1 + 30, ly + 4, x1 + 52, ly + 22), fill=col)
        lab = name
        while d.textlength(lab, font=f_l) > W - (x1 + 62) - 20 and len(lab) > 4:
            lab = lab[:-1]
        d.text((x1 + 62, ly), lab if lab == name else lab[:-1] + "…", font=f_l, fill=ink)
    d.text((x0, y0 - 34), "每年篇數" + ("（對數刻度）" if log else "") + ("　來源：OpenAlex（各關鍵字組）" if from_gap else "　來源：本次檢索結果"),
           font=f_l, fill=grey)

    # ---- 下：前 10 高被引橫條
    top = sorted([p for p in papers if p.cited_by], key=lambda p: -p.cited_by)[:10]
    bx0, by0, bx1 = 640, 600, W - 120
    d.text((60, 560), "前 10 高被引（本次檢索結果）", font=f_l, fill=grey)
    if top:
        cmax = top[0].cited_by
        bh = 36
        for i, p in enumerate(top):
            yy = by0 + i * (bh + 2)
            lab = f"{p.first_author_family} ({p.year}) {clean_text(p.title)}"
            while d.textlength(lab, font=f_s) > bx0 - 80 and len(lab) > 10:
                lab = lab[:-2]
            if lab != f"{p.first_author_family} ({p.year}) {clean_text(p.title)}":
                lab = lab[:-1] + "…"
            d.text((60, yy + bh / 2), lab, font=f_s, fill=ink, anchor="lm")
            w = (bx1 - bx0 - 90) * p.cited_by / cmax
            d.rectangle((bx0, yy + 4, bx0 + max(2, w), yy + bh - 4), fill=PALETTE[i % 2 * 2])
            d.text((bx0 + w + 8, yy + bh / 2), f"{p.cited_by:,}", font=f_s, fill=ink, anchor="lm")
    else:
        d.text((bx0, by0), "（無被引資料）", font=f_l, fill=grey)
    d.text((W - 30, H - 24), "FC文獻缺口搜尋工具　NDHU自資系游豐兆製作", font=f_s, fill=grey, anchor="rb")
    img.save(path, "PNG")


# ================================================================ JSON
def write_query_json(papers: list, gap: Optional[dict], path: str, options: dict, results: dict) -> None:
    rec = {
        "app": "FC LitGap Finder",
        "version": VERSION,
        "time": _dt.datetime.now().isoformat(timespec="seconds"),
        "mailto": options.get("mailto", ""),
        "keywords": options.get("keywords") or [],
        "settings": {k: v for k, v in (options.get("settings") or {}).items() if k not in ("email", "mailto") and "key" not in k.lower() and "password" not in k.lower()},
        "apis": ["Crossref", "OpenAlex", "Unpaywall", "Semantic Scholar", "doi.org"],
        "counts": {"total": len(papers), "included": sum(1 for p in papers if p.include),
                   "manual": sum(1 for p in papers if p.source == "manual"),
                   "oa": sum(1 for p in papers if p.oa_pdf_url)},
        "gap": None if not gap else {
            "groups": gap.get("groups"), "threshold": gap.get("threshold"), "elapsed_sec": gap.get("elapsed_sec"),
            "pairs": [{"groups": list(r["groups"]), "count": r.get("count"), "gap": r.get("gap"), "query": r.get("query")}
                      for r in gap.get("pairs") or []],
            "triples": [{"groups": list(r["groups"]), "count": r.get("count"), "gap": r.get("gap"), "query": r.get("query")}
                        for r in gap.get("triples") or []],
            "trend": {str(k): v for k, v in (gap.get("trend") or {}).items()},
        },
        "papers": [{"order": p.order, "doi": p.doi, "title": clean_text(p.title), "year": p.year, "source": p.source,
                    "include": p.include, "chapter": p.chapter, "relevance": p.relevance} for p in papers],
        "outputs": {os.path.basename(k): v for k, v in results.items()},
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(rec, f, ensure_ascii=False, indent=2)


# ================================================================ 主函式
def export_all(papers: list, gap: Optional[dict], out_dir: str, name: str, options: Optional[dict] = None,
               progress: Progress = None, cancel_event=None) -> dict:
    options = dict(options or {})
    folder = os.path.join(out_dir, safe_name(name, 80))
    os.makedirs(folder, exist_ok=True)
    fmts = options.get("formats") or list(FILES)
    results: dict = {}
    errors: dict = {}
    steps = [f for f in ["docx", "pdf", "xlsx", "gap_xlsx", "ris", "bib", "png"] if f in fmts]
    for i, fmt in enumerate(steps):
        check_cancel(cancel_event)
        path = os.path.join(folder, FILES[fmt])
        report(progress, f"匯出 {FILES[fmt]}", i / (len(steps) + 1))
        try:
            if fmt == "docx":
                write_docx(papers, path, options, gap)
            elif fmt == "pdf":
                engine = options.get("pdf_engine", "auto")
                docx_path = os.path.join(folder, FILES["docx"])
                ok = False
                if engine in ("auto", "word") and os.path.isfile(docx_path):
                    ok = docx_to_pdf_word(docx_path, path)
                    if ok:
                        options["_pdf_engine_used"] = "Word"
                if not ok:
                    write_pdf_reportlab(papers, path, options, gap)
                    options["_pdf_engine_used"] = "reportlab"
            elif fmt == "xlsx":
                write_xlsx(papers, path, options)
            elif fmt == "gap_xlsx":
                if not gap:
                    continue
                write_gap_xlsx(gap, path)
            elif fmt == "ris":
                write_ris(papers, path, options)
            elif fmt == "bib":
                write_bib(papers, path, options)
            elif fmt == "png":
                write_trend_png(papers, gap, path)
            results[path] = os.path.isfile(path)
        except Exception as e:  # 單一格式失敗不影響其他格式
            results[path] = False
            errors[FILES[fmt]] = f"{type(e).__name__}: {e}"
    if "json" in fmts:
        path = os.path.join(folder, FILES["json"])
        options["_errors"] = errors
        try:
            write_query_json(papers, gap, path, dict(options, settings=dict(options.get("settings") or {},
                             pdf_engine=options.get("_pdf_engine_used"))), results)
            results[path] = True
        except Exception as e:
            results[path] = False
            errors[FILES["json"]] = str(e)
    export_all.last_errors = errors  # type: ignore[attr-defined]
    report(progress, f"匯出完成：{sum(results.values())}/{len(results)} 個檔案" + (f"；失敗：{', '.join(errors)}" if errors else ""), 1.0)
    return results
