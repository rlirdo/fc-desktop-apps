#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
core_mask.py — 姓名遮罩與學生編號（NameMasker 2.0 核心）
========================================================
把 Zuvio「下載數據」xlsx／名冊 xlsx／通用 CSV 轉成「可以安全拿去分析」的檔案：

  1. 找出整張工作表裡**所有**含「姓名」的表頭列（同一張表常有好幾個區塊：
     作答明細、未作答名單、各子題彙總…），每個表頭列往下到空白列／下一個
     表頭列為一個「區塊」。
  2. 在 A 欄插入一欄「學生編號」，每位學生給一個 `{學期}_{課程縮寫}_{流水號}`
     （例 `115-1_EC_1`）。同一個人在不同區塊拿**同一個編號**；編號鍵優先用
     9 碼學號，沒有學號才用姓名；流水號依**首次出現順序**。匿名作答者不編號。
  3. 「學號」欄、「電子郵件」欄的每個字元都改成 O（長度不變）。沒有該欄就略過。
  4. 姓名欄的中文姓名第 2 個字改成 O（可關閉）；所有文字儲存格裡出現的名冊
     姓名，改成該生的學生編號（「和王小明討論」→「和 115-1_EC_3 討論」）。
  5. 新增工作表「學生編號連結姓名」：學生編號｜姓名｜學號（原）｜電子郵件（原）｜首次出現列。
     ★ 這張表是「再識別鑰匙」，只能留在自己的電腦，不可上傳、不可外流。
  6. 另存「原檔名02.xlsx」（已存在就 03、04…），**原檔一個位元組都不動**。

全程離線，只讀寫本機檔案。需求：Python 3.8+、openpyxl。
"""
from __future__ import annotations

import csv
import os
import re
import sys
from dataclasses import dataclass, field

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - 封裝後不會發生
    sys.exit("缺少 openpyxl，請先執行：pip install openpyxl")

__version__ = "2.0.0"

# --------------------------------------------------------------------------
# 常數與樣式
# --------------------------------------------------------------------------
LINK_SHEET_TITLE = "學生編號連結姓名"
LINK_SHEET_HEADERS = ("學生編號", "姓名", "學號（原）", "電子郵件（原）", "首次出現列")

_CJK = "㐀-䶿一-鿿豈-﫿"
NAME_RUN = re.compile(f"[{_CJK}]{{2,}}")            # 連續 2 個以上中文字
STUDENT_ID_RE = re.compile(r"^\d{9}$")               # 9 碼學號
DIGIT_RUN_RE = re.compile(r"\d{9,}")                  # 文字內 ≥9 碼數字串（學號黏在句子裡也抓）
EMAIL_IN_TEXT_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
ID_HEADER_PAT = re.compile(r"(學號|學生證號|student\s*(id|no))", re.I)
EMAIL_HEADER_PAT = re.compile(r"(電子郵件|電子信箱|電郵|信箱|e-?mail|mail)", re.I)
NAME_HEADER = "姓名"

# 這些不是人名，不編號、不做姓名遮罩
NON_PERSON = {
    "姓名", "學號", "匿名", "匿名作答者", "匿名使用者", "未作答", "未作答學生",
    "小計", "合計", "總計", "平均", "人數", "無", "N/A", "n/a", "-", "—",
}
ANON_PREFIX = ("匿名",)

SUPPORTED_EXT = (".xlsx", ".xlsm", ".csv")
COURSE_RE = re.compile(r"^[A-Za-z]{2}$")
SEMESTER_RE = re.compile(r"^\d{3}-[12]$")


class MaskError(Exception):
    """可直接顯示給使用者看的錯誤訊息。"""


# --------------------------------------------------------------------------
# 報告物件
# --------------------------------------------------------------------------
@dataclass
class Report:
    src: str = ""
    dst: str = ""
    sheet_title: str = ""
    course_code: str = ""
    semester: str = ""
    mask_names: bool = True
    blocks: int = 0
    block_info: list = field(default_factory=list)   # [(表頭列, 起, 迄, 學號欄, 郵件欄)]
    students: int = 0
    anon_rows: int = 0
    id_masked: int = 0
    email_masked: int = 0
    name_masked: int = 0
    text_replacements: int = 0
    link_rows: int = 0
    notes: list = field(default_factory=list)
    student_numbers: list = field(default_factory=list)   # 只放編號字串，不含姓名

    def summary_lines(self):
        """給 GUI／CLI 顯示的摘要（**不含任何姓名**）。"""
        L = [
            f"找到 {self.blocks} 個資料區塊（工作表「{self.sheet_title}」）",
            f"編號學生 {self.students} 位　（{self.semester}_{self.course_code}_1 … "
            f"{self.semester}_{self.course_code}_{self.students}）" if self.students else
            "編號學生 0 位",
            f"匿名／未具名列 {self.anon_rows} 列（不編號）",
            f"遮罩學號 {self.id_masked} 筆、電子郵件 {self.email_masked} 筆（每字元改 O，長度不變）",
            f"姓名欄遮罩 {self.name_masked} 筆" + ("" if self.mask_names else "（本次未勾選姓名遮罩）"),
            f"文字內姓名改成學生編號 {self.text_replacements} 處（含整格為姓名的統計儲存格）",
            f"對照表「{LINK_SHEET_TITLE}」{self.link_rows} 筆",
            f"輸出檔：{self.dst}",
        ]
        for n in self.notes:
            L.append("※ " + n)
        return L


# --------------------------------------------------------------------------
# 小工具
# --------------------------------------------------------------------------
def validate_course(code):
    c = (code or "").strip()
    if not COURSE_RE.match(c):
        raise MaskError("「課程縮寫」必須是 2 個英文字母（例如 EC、GC、BI）。")
    return c.upper()


def validate_semester(sem):
    s = (sem or "").strip()
    if not SEMESTER_RE.match(s):
        raise MaskError("「學期」格式必須像 115-1（3 位數字 + 連字號 + 1 或 2）。")
    return s


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    return str(v)


def is_formula(v):
    return isinstance(v, str) and v.startswith("=")


def normalize_id(v):
    """學號正規化：int 411354043 → '411354043'；float 411354043.0 → '411354043'。"""
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if float(v).is_integer() else str(v)
    return str(v).strip()


def is_person_name(name):
    """是否為可編號的真實姓名（≥2 字、非匿名／標籤字樣）。"""
    t = (name or "").strip()
    if len(t) < 2:
        return False
    if t in NON_PERSON:
        return False
    for p in ANON_PREFIX:
        if t.startswith(p):
            return False
    if STUDENT_ID_RE.match(t):
        return False
    return True


def is_anonymous(name):
    t = (name or "").strip()
    return any(t.startswith(p) for p in ANON_PREFIX)


def mask_name(text):
    """把字串中第一段連續中文字的第 2 個字改為 O；無中文則原樣返回。"""
    if not isinstance(text, str):
        return text, False
    m = NAME_RUN.search(text)
    if not m:
        return text, False
    run = m.group(0)
    masked = run[0] + "O" + run[2:]
    return text[: m.start()] + masked + text[m.start() + len(run):], True


def mask_all_chars(value):
    """每個字元改成 O，長度不變（int/float 學號先正規化成字串）。"""
    s = normalize_id(value) if not isinstance(value, str) else value
    if s.strip() == "":
        return value, False
    return "O" * len(s), True


def replace_names_in_text(text, pairs):
    """把文字中的名冊姓名換成學生編號，回傳 (新字串, 替換處數)。

    - 3 字以上的姓名：直接子字串替換（「和王小明討論」→「和 115-1_EC_1 討論」）。
    - 2 字姓名：只在「整段連續中文剛好等於該姓名」時才換，避免把「答對」「原理」
      這類一般詞誤判成人名（測驗題目型匯出的答案欄全是 2 字中文）。
    """
    if not isinstance(text, str) or not text:
        return text, 0
    hits = 0
    # 自由文字裡夾帶的學號（≥9 碼數字串，可能與其他字黏在一起，例「4113xxxxx115-1 EC」）
    # 與電子郵件，一律換成等長的 O（2026/09/16 W02 實測：知情同意書/智財題學生會自己打學號）
    def _o(m):
        return "O" * len(m.group(0))
    text, n1 = DIGIT_RUN_RE.subn(_o, text)
    text, n2 = EMAIL_IN_TEXT_RE.subn(_o, text)
    hits += n1 + n2
    long_pairs = [(n, v) for n, v in pairs if len(n) >= 3]
    short_pairs = dict((n, v) for n, v in pairs if len(n) < 3)

    for nm, no in long_pairs:
        if nm in text:
            hits += text.count(nm)
            text = text.replace(nm, no)

    if short_pairs:
        out_parts, last = [], 0
        for m in NAME_RUN.finditer(text):
            run = m.group(0)
            if run in short_pairs:
                out_parts.append(text[last:m.start()])
                out_parts.append(short_pairs[run])
                last = m.end()
                hits += 1
        if last:
            out_parts.append(text[last:])
            text = "".join(out_parts)
    return text, hits


def next_output_path(src, ext=".xlsx"):
    """另存新檔路徑：原檔名 + 02（已存在則 03、04…），與原檔同資料夾。"""
    folder, fname = os.path.split(os.path.abspath(src))
    stem, _ = os.path.splitext(fname)
    n = 2
    while True:
        dst = os.path.join(folder, f"{stem}{n:02d}{ext}")
        if not os.path.exists(dst):
            return dst
        n += 1


# --------------------------------------------------------------------------
# 區塊偵測
# --------------------------------------------------------------------------
@dataclass
class Block:
    header_row: int
    name_col: int
    id_col: int = 0
    email_col: int = 0
    first: int = 0
    last: int = 0

    @property
    def has_rows(self):
        return self.last >= self.first


def row_is_blank(ws, r, maxc):
    for c in range(1, maxc + 1):
        if cell_text(ws.cell(r, c).value).strip() != "":
            return False
    return True


def find_header_rows(ws, maxr, maxc):
    """回傳 [(列號, 姓名欄索引)]；優先取精確等於「姓名」的表頭，找不到才用模糊比對。"""
    exact, fuzzy = [], []
    for r in range(1, maxr + 1):
        hit_e = hit_f = 0
        for c in range(1, maxc + 1):
            t = cell_text(ws.cell(r, c).value).strip()
            if not t:
                continue
            if t == NAME_HEADER:
                hit_e = c
                break
            if not hit_f and NAME_HEADER in t and len(t) <= 10:
                hit_f = c
        if hit_e:
            exact.append((r, hit_e))
        elif hit_f:
            fuzzy.append((r, hit_f))
    return exact if exact else fuzzy


def detect_blocks(ws, maxr, maxc):
    heads = find_header_rows(ws, maxr, maxc)
    if not heads:
        raise MaskError(
            "整張工作表都找不到「姓名」欄位。\n"
            "請確認這是 Zuvio「下載數據」的 xlsx，或名冊的第 1 列有「姓名」標題；\n"
            "也可以把姓名欄的標題改成「姓名」後再執行一次。")
    head_rows = [h[0] for h in heads]
    blocks = []
    for i, (hr, ncol) in enumerate(heads):
        idc = emc = 0
        for c in range(1, maxc + 1):
            if c == ncol:
                continue
            t = cell_text(ws.cell(hr, c).value).strip()
            if not t or len(t) > 24:
                continue
            if not idc and ID_HEADER_PAT.search(t):
                idc = c
            if not emc and EMAIL_HEADER_PAT.search(t):
                emc = c
        nxt = head_rows[i + 1] if i + 1 < len(heads) else None
        r, last = hr + 1, hr
        while r <= maxr:
            if nxt and r >= nxt:
                break
            if row_is_blank(ws, r, maxc):
                break
            last = r
            r += 1
        blocks.append(Block(hr, ncol, idc, emc, hr + 1, last))
    return blocks


# --------------------------------------------------------------------------
# 主流程（單一工作表）
# --------------------------------------------------------------------------
def process_sheet(ws, course, sem, do_mask_names, rep):
    maxr, maxc = ws.max_row, ws.max_column
    blocks = detect_blocks(ws, maxr, maxc)
    rep.blocks = len(blocks)
    rep.sheet_title = ws.title

    if not any(b.id_col for b in blocks):
        rep.notes.append("這個檔案沒有「學號」欄，改用姓名當編號鍵。")
    if not any(b.email_col for b in blocks):
        rep.notes.append("本檔無電子郵件欄，略過（正常，Zuvio 匯出通常沒有；名冊類匯出才有）。")

    # ---- 第 1 遍：蒐集所有資料列（姓名、學號、郵件） ----------------------
    entries = []          # (列, block, 姓名, 學號, 郵件)
    alias = {}            # 姓名 -> 9 碼學號（同一人在別的區塊沒有學號時可對上）
    for b in blocks:
        if not b.has_rows:
            continue
        for r in range(b.first, b.last + 1):
            name = cell_text(ws.cell(r, b.name_col).value).strip()
            if not name:
                continue
            sid = normalize_id(ws.cell(r, b.id_col).value) if b.id_col else ""
            eml = cell_text(ws.cell(r, b.email_col).value).strip() if b.email_col else ""
            entries.append((r, b, name, sid, eml))
            if is_person_name(name) and STUDENT_ID_RE.match(sid) and name not in alias:
                alias[name] = sid

    # ---- 指派學生編號（首次出現順序、跨區塊同人同號） --------------------
    numbers = {}          # key -> 學生編號
    link_rows = []        # (學生編號, 姓名, 學號, 郵件, 首次出現列)
    row_number = {}       # 列 -> A 欄要寫的字串
    name_to_number = {}   # 姓名 -> 學生編號（文字替換用）
    for (r, b, name, sid, eml) in entries:
        if not is_person_name(name):
            row_number[r] = "匿名" if is_anonymous(name) else ""
            rep.anon_rows += 1
            continue
        key = alias.get(name) or (sid if STUDENT_ID_RE.match(sid) else name)
        if key not in numbers:
            no = f"{sem}_{course}_{len(numbers) + 1}"
            numbers[key] = no
            link_rows.append((no, name, sid, eml, r))
            name_to_number.setdefault(name, no)
        else:
            no = numbers[key]
            name_to_number.setdefault(name, no)
            # 補齊對照表上缺的學號／郵件（同一人在別的區塊才出現）
            for i, rec in enumerate(link_rows):
                if rec[0] == no:
                    link_rows[i] = (rec[0], rec[1], rec[2] or sid, rec[3] or eml, rec[4])
                    break
        row_number[r] = numbers[key]

    rep.students = len(numbers)
    rep.student_numbers = [rec[0] for rec in link_rows]

    # ---- 插入 A 欄「學生編號」（原 A → B，欄索引全部 +1） ----------------
    try:
        ws.insert_cols(1)
    except Exception as e:  # pragma: no cover
        raise MaskError(f"無法在 A 欄插入「學生編號」欄：{e}")
    for b in blocks:
        b.name_col += 1
        if b.id_col:
            b.id_col += 1
        if b.email_col:
            b.email_col += 1
    maxc += 1
    entries = [(r, b, n, s, e) for (r, b, n, s, e) in entries]

    for b in blocks:
        ws.cell(b.header_row, 1).value = "學生編號"
    for r, val in row_number.items():
        if val:
            ws.cell(r, 1).value = val

    # ---- 遮罩學號／電子郵件（每字元 O、長度不變） ------------------------
    for (r, b, name, sid, eml) in entries:
        if b.id_col:
            v = ws.cell(r, b.id_col).value
            if not is_formula(v):
                nv, ok = mask_all_chars(v)
                if ok:
                    ws.cell(r, b.id_col).value = nv
                    rep.id_masked += 1
        if b.email_col:
            v = ws.cell(r, b.email_col).value
            if not is_formula(v):
                nv, ok = mask_all_chars(v)
                if ok:
                    ws.cell(r, b.email_col).value = nv
                    rep.email_masked += 1

    # ---- 文字儲存格裡的姓名 → 學生編號 ----------------------------------
    protected = set()      # 不做文字替換的儲存格（姓名／學號／郵件欄的資料列）
    for (r, b, name, sid, eml) in entries:
        protected.add((r, b.name_col))
        if b.id_col:
            protected.add((r, b.id_col))
        if b.email_col:
            protected.add((r, b.email_col))

    pairs = sorted(name_to_number.items(), key=lambda kv: -len(kv[0]))
    if pairs:
        for r in range(1, maxr + 1):
            for c in range(2, maxc + 1):
                if (r, c) in protected:
                    continue
                v = ws.cell(r, c).value
                if not isinstance(v, str) or is_formula(v) or not v:
                    continue
                new, hits = replace_names_in_text(v, pairs)
                if hits:
                    ws.cell(r, c).value = new
                    rep.text_replacements += hits

    # ---- 姓名欄遮罩（第 2 字 → O） --------------------------------------
    if do_mask_names:
        for (r, b, name, sid, eml) in entries:
            if not is_person_name(name):
                continue
            v = ws.cell(r, b.name_col).value
            if is_formula(v):
                continue
            nv, ok = mask_name(cell_text(v))
            if ok:
                ws.cell(r, b.name_col).value = nv
                rep.name_masked += 1

    # ---- 版面小調整 -----------------------------------------------------
    try:
        ws.column_dimensions["A"].width = 16
    except Exception:
        pass

    rep.block_info = [(b.header_row, b.first, b.last,
                       get_column_letter(b.id_col) if b.id_col else "",
                       get_column_letter(b.email_col) if b.email_col else "")
                      for b in blocks]
    return link_rows


def add_link_sheet(wb, link_rows, rep):
    title = LINK_SHEET_TITLE
    i = 2
    while title in wb.sheetnames:
        title = f"{LINK_SHEET_TITLE}{i}"
        i += 1
    ls = wb.create_sheet(title)
    ls.append(list(LINK_SHEET_HEADERS))
    for no, name, sid, eml, row in link_rows:
        ls.append([no, name, sid, eml, row])
    for col, w in zip("ABCDE", (18, 14, 14, 26, 12)):
        try:
            ls.column_dimensions[col].width = w
        except Exception:
            pass
    rep.link_rows = len(link_rows)
    if title != LINK_SHEET_TITLE:
        rep.notes.append(f"已有同名工作表，對照表改名為「{title}」。")
    return title


# --------------------------------------------------------------------------
# CSV 支援
# --------------------------------------------------------------------------
def workbook_from_csv(path):
    raw = None
    for enc in ("utf-8-sig", "cp950", "utf-8"):
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                raw = list(csv.reader(f))
            break
        except UnicodeDecodeError:
            continue
    if raw is None:
        with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
            raw = list(csv.reader(f))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CSV"
    for row in raw:
        ws.append([c if c != "" else None for c in row])
    return wb


# --------------------------------------------------------------------------
# 對外主函式
# --------------------------------------------------------------------------
def process_workbook(src, course_code, semester, mask_names=True):
    """處理單一檔案，回傳 Report。失敗丟 MaskError（訊息可直接顯示給使用者）。"""
    course = validate_course(course_code)
    sem = validate_semester(semester)

    src = os.path.abspath(src)
    if not os.path.isfile(src):
        raise MaskError(f"找不到檔案：{src}")
    low = src.lower()
    if not low.endswith(SUPPORTED_EXT):
        raise MaskError("只支援 .xlsx / .xlsm / .csv（舊版 .xls 請先用 Excel 另存為 .xlsx）")

    rep = Report(src=src, course_code=course, semester=sem, mask_names=bool(mask_names))

    if low.endswith(".csv"):
        wb = workbook_from_csv(src)
        rep.notes.append("輸入為 CSV，輸出轉存為 xlsx。")
    else:
        try:
            wb = openpyxl.load_workbook(src, keep_vba=low.endswith(".xlsm"), data_only=False)
        except Exception as e:
            raise MaskError("無法開啟檔案（可能不是有效的 Excel 檔，或正被 Excel 開啟中）：\n"
                            f"{e}")

    # 選擇工作表：預設第 1 個；第 1 個沒有「姓名」時退而找第一個有的
    ws = wb.worksheets[0]
    target = ws if find_header_rows(ws, ws.max_row, ws.max_column) else None
    if target is None:
        for cand in wb.worksheets[1:]:
            if find_header_rows(cand, cand.max_row, cand.max_column):
                target = cand
                rep.notes.append(f"第 1 個工作表沒有「姓名」欄，改處理「{cand.title}」。")
                break
    if target is None:
        raise MaskError(
            "整個活頁簿都找不到「姓名」欄位。\n"
            "請確認這是 Zuvio「下載數據」的 xlsx，或名冊的第 1 列有「姓名」標題。")
    if len(wb.worksheets) > 1:
        rep.notes.append(f"這個活頁簿有 {len(wb.worksheets)} 張工作表，只處理「{target.title}」，其餘原樣保留。")

    link_rows = process_sheet(target, course, sem, mask_names, rep)
    add_link_sheet(wb, link_rows, rep)

    dst = next_output_path(src, ".xlsx")
    try:
        wb.save(dst)
    except Exception as e:
        raise MaskError(f"無法儲存新檔（請確認資料夾可寫入、檔案沒有被 Excel 開著）：\n{e}")
    finally:
        try:
            wb.close()
        except Exception:
            pass
    rep.dst = dst
    return rep


# --------------------------------------------------------------------------
def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    files = [a for a in argv if not a.startswith("-")]
    if not files:
        sys.exit("用法：python core_mask.py 檔案.xlsx [--course EC --semester 115-1]")
    course, sem, mk = "EC", "115-1", True
    for i, a in enumerate(argv):
        if a == "--course" and i + 1 < len(argv):
            course = argv[i + 1]
        if a == "--semester" and i + 1 < len(argv):
            sem = argv[i + 1]
        if a == "--no-mask-names":
            mk = False
    files = [f for f in files if f not in (course, sem)]
    for f in files:
        try:
            rep = process_workbook(f, course, sem, mk)
        except MaskError as e:
            print(f"失敗：{f}\n{e}")
            continue
        for line in rep.summary_lines():
            print(line)
    return 0


if __name__ == "__main__":
    main()
