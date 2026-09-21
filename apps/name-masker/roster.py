#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
roster.py — 原始名單 → 固定學生編號對照表（NameMasker 2.2 核心之一）
====================================================================
TA 先給一份「原始名單」，本模組負責：

  1. 解析名單：
     * Excel／CSV：任一工作表第 1–5 列內，表頭同時有「學號」與「姓名」即可
       （欄位順序不拘、允許 None 表頭；另可有「序號」「班級／系級／系所＋年級」
       「電子郵件」）。
     * Word（.docx）：先找文件中的表格（表頭同列含「學號」「姓名」），
       表格都不行就把段落＋儲存格文字逐行丟給 PDF 那一套行解析。
     * PDF：東華教務系統「選課名單／成績登記表」（文字型 PDF，用 pypdf 抽字）。
     * 若檔案本身就是「對照表」（有「學生編號」欄）→ 直接採用，不重新編號。
  2. 給固定學生編號：名單有「序號」（或 PDF 的序號）且每位都有、不重複
     → `{學期}_{課程縮寫}_{序號}`（**允許缺號**）；否則依**學號字串遞增排序**給 1..N。
  3. 讀寫對照表 `{學期}_{課程縮寫}_學生名單與學生編號對照表.xlsx`：
     * 工作表「學生編號對照」：序號｜學生編號｜學號｜姓名｜班級
     * 工作表「名單外作答者」：學生編號｜學號｜姓名｜首次出現檔案
       （不在名單的作答者＝退選等，編號自 101 起遞增，**持久登記**、跨檔跨週一致）
     * 工作表「說明」：規則與警語
  4. 名單外作答者的持久登記（`Codebook.register_outsider`）。

★ 對照表含真實姓名與學號，是「再識別鑰匙」，只能留在自己的電腦。

全程離線，只讀寫本機檔案。需求：Python 3.8+、openpyxl；
讀 PDF 才需要 pypdf，讀 Word 才需要 python-docx。

本檔與 `HomeworkWordCloud\\core\\roster.py` 完全同步（只有檔頭這段說明不同）。
"""
from __future__ import annotations

import csv
import os
import re
import sys
from dataclasses import dataclass, field

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("缺少 openpyxl，請先執行：pip install openpyxl")

__version__ = "2.2.0"

# --------------------------------------------------------------------------
# 常數
# --------------------------------------------------------------------------
CODE_SHEET = "學生編號對照"
CODE_HEADERS = ("序號", "學生編號", "學號", "姓名", "班級")
OUT_SHEET = "名單外作答者"
OUT_HEADERS = ("學生編號", "學號", "姓名", "首次出現檔案")
NOTE_SHEET = "說明"

CODEBOOK_NAME_FMT = "{sem}_{course}_學生名單與學生編號對照表.xlsx"
OUTSIDER_START = 101            # 名單外作答者從 101 號起

ROSTER_EXT = (".xlsx", ".xlsm", ".csv", ".pdf", ".docx")

# ---- 學號形狀（東華：一律 9 個字元；日後要加新形狀就加在這裡） -------------
SID_LEN = 9
SID_SHAPES = (
    r"\d{9}",                 # 990054043
    r"\d{5}[A-Za-z]\d{3}",    # 99001A001（第 6 碼是英文字母）
    r"[A-Za-z]\d{8}",         # A99000001（第 1 碼是英文字母）
)
SID_SHAPE_RE = re.compile("^(?:%s)$" % "|".join(SID_SHAPES))
SID_IN_TEXT_RE = re.compile(r"(?<![A-Za-z0-9])(?:%s)(?![A-Za-z0-9])"
                            % "|".join(SID_SHAPES))
ALNUM_RUN_RE = re.compile(r"[A-Za-z0-9]+")
SEQ_MAX = 999                   # 名單序號的上限（超過就不可能是序號）

_CJK = "㐀-䶿一-鿿豈-﫿"
CJK_RE = re.compile("[%s]" % _CJK)
NAME_TAIL_RE = re.compile("[%s·・．A-Za-z]{2,12}$" % _CJK)
PURE_NAME_RE = re.compile("^[%s·・．]{2,6}$" % _CJK)
SID9_RE = re.compile(r"\d{9}")
NOTE_MARK_RE = re.compile(r"[【\[（(]\s*註[^】\])）]*[】\])）]")
# 跨行被切斷的【註*】：`【` 與 `】` 之間可夾換行，最多約 8 個字元
NOTE_SPAN_RE = re.compile(r"【[^【】]{0,5}註[^【】]{0,8}】")
LEAD_INT_RE = re.compile(r"^\s*(\d{1,3})(?!\d)")
PURE_INT_RE = re.compile(r"^\s*(\d{1,3})\s*$")
CLASS_HINT_RE = re.compile("[系班級年大碩博所]|學程")
# 分行版面才用的嚴格判定（避免把下一位學生的姓名誤當成班級）
CLASS_STRONG_RE = re.compile("[系班所]|學程|年級")
PURE_CJK_RE = re.compile("^[%s ]{2,16}$" % _CJK)

# 表頭／雜訊字樣：出現這些字的「姓名」一律不採用
HEADER_WORDS = (
    "姓名", "學號", "序號", "班級", "系級", "系所", "年級", "學分", "任課教師",
    "課程", "名單", "大學", "學年度", "學期", "備註", "電子郵件", "信箱",
    "選課", "人數", "列印", "頁次", "合計", "小計", "註冊", "教務",
)

# 表頭比對用
ID_HDR_RE = re.compile(r"(學號|學生證號|student\s*(id|no))", re.I)
NAME_HDR_RE = re.compile(r"^(姓名|學生姓名|name)$", re.I)
SEQ_HDR_RE = re.compile(r"^(序號|編號|項次|No\.?|#)$", re.I)
CLASS_HDR_RE = re.compile(r"(班級|系級|班別|系所|科系|年級)")
CODE_HDR_RE = re.compile(r"^(學生編號|學生代號|代號)$")


class RosterError(Exception):
    """可直接顯示給使用者看的錯誤訊息。"""


PDF_FAIL_MSG = ("這份名單讀不出任何學生，請改用 Excel 名單"
                "（需有 學號、姓名 欄）")
NO_NAME_NOTE_FMT = ("序號 {seq} 的學生在名單檔中讀不到姓名（可能是罕用字），"
                    "請在對照表補上姓名；遮罩仍會依學號進行。")
MISSING_SEQ_NOTE_FMT = ("名單序號缺 {seqs}，可能有學生沒讀到，"
                        "請對照原名單人數。")
DUP_SEQ_NOTE = ("名單的序號有重複或有些學生讀不到序號，"
                "改依學號遞增排序給 1..N（編號可能與原名單序號不同）。")
NO_SEQ_NOTE = "這份名單沒有讀到序號，改依學號遞增排序給 1..N。"
DOCX_NEED_PKG = ("要讀 Word 名單需要 python-docx 套件"
                 "（請執行：pip install python-docx）；\n"
                 "或改用 Excel／CSV 名單（需有 學號、姓名 欄）。")
DOC_OLD_MSG = ("這是 Word 舊格式（.doc），本程式讀不了。\n"
               "請先用 Word 開啟後「另存新檔」為 .docx，再選一次。")


# --------------------------------------------------------------------------
# 小工具
# --------------------------------------------------------------------------
def normalize_id(v):
    """學號正規化：int 990054043 → '990054043'；float → 去掉 .0；去頭尾空白。"""
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if float(v).is_integer() else str(v)
    return str(v).strip()


def text_of(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, float) and float(v).is_integer():
        return str(int(v))
    return str(v).strip()


def _looks_like_name(t):
    t = (t or "").strip()
    if len(t) < 2 or len(t) > 12:
        return False
    for w in HEADER_WORDS:
        if w in t:
            return False
    if re.search(r"\d", t):
        return False
    return bool(re.search("[%s A-Za-z]" % _CJK, t))


def _name_ok(t):
    """比 `_looks_like_name` 寬鬆一點：允許含空白的族名、中文＋英文名（最長 24 字）。"""
    t = (t or "").strip()
    if len(t) < 2 or len(t) > 24:
        return False
    for w in HEADER_WORDS:
        if w in t:
            return False
    if re.search(r"\d", t):
        return False
    if CJK_RE.search(t):
        return True
    return len(re.findall(r"[A-Za-z]", t)) >= 2


def _clean_name(t):
    """姓名字串清理：去掉註記、全形空白，連續空白縮成一個。"""
    t = NOTE_MARK_RE.sub(" ", t or "")
    t = t.replace("　", " ")
    t = re.sub(r"[·・．]", "·", t)
    t = re.sub(r"\s+", " ", t).strip(" 、,，.")
    return t.strip()


# --------------------------------------------------------------------------
# 資料物件
# --------------------------------------------------------------------------
@dataclass
class Student:
    seq: int = 0
    code: str = ""
    sid: str = ""
    name: str = ""
    klass: str = ""


@dataclass
class Outsider:
    code: str = ""
    sid: str = ""
    name: str = ""
    src: str = ""


@dataclass
class RosterData:
    """解析完成、尚未編號的名單。"""
    records: list = field(default_factory=list)     # [Student]（code 可能為空）
    source: str = ""
    kind: str = ""              # 'excel' / 'csv' / 'pdf' / 'docx' / 'codebook'
    rule: str = ""              # '依序號' / '依學號排序' / '沿用對照表'
    notes: list = field(default_factory=list)

    @property
    def count(self):
        return len(self.records)


# --------------------------------------------------------------------------
# PDF 解析（東華選課名單／成績登記表）
# --------------------------------------------------------------------------
def pdf_lines(path):
    """用 pypdf 抽出所有頁的文字行。"""
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RosterError(
            "要讀 PDF 名單需要 pypdf 套件（請執行：pip install pypdf）；\n"
            "或改用 Excel／CSV 名單（需有 學號、姓名 欄）。")
    try:
        reader = PdfReader(path)
    except Exception as e:
        raise RosterError(f"這個 PDF 打不開（可能不是有效的 PDF 或已加密）：\n{e}")
    lines = []
    for page in reader.pages:
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        for ln in txt.split("\n"):
            lines.append(ln.rstrip())
    if not any(s.strip() for s in lines):
        raise RosterError(
            "這份 PDF 抽不出任何文字（可能是掃描影像檔）。\n" + PDF_FAIL_MSG)
    return lines


def normalize_lines(lines):
    """抽字後的正規化：把跨行被切斷的 `【註*】` 整段移除，並把兩行接起來。"""
    txt = "\n".join((ln or "").replace("　", " ").rstrip() for ln in lines)
    txt = NOTE_SPAN_RE.sub("", txt)
    return [ln.rstrip() for ln in txt.split("\n")]


def _split_glued(run):
    """把「(姓名英文字母)＋學號＋序號」黏成一串的英數字串拆開。

    回傳候選 list：`(k, prefix, sid, seq)`，k = 黏在學號後面的序號位數。
    依 k = 2 → 1 → 3 → 0 的順序排列（多個候選時由 `_pick_glued` 決定）。
    """
    cands = []
    L = len(run)
    for k in (2, 1, 3, 0):
        if L < SID_LEN + k:
            continue
        sid = run[L - SID_LEN - k:L - k]
        prefix = run[:L - SID_LEN - k]
        tail = run[L - k:] if k else ""
        if prefix and not (prefix.isascii() and prefix.isalpha()):
            continue            # 學號左邊只能是「屬於姓名的英文字母」或空字串
        if not SID_SHAPE_RE.match(sid):
            continue
        if k:
            if not tail.isdigit():
                continue
            seq = int(tail)
            if seq < 1 or seq > SEQ_MAX:
                continue
        else:
            seq = 0
        cands.append((k, prefix, sid, seq))
    return cands


def _pick_glued(cands, last_seq):
    """多個候選都成立時：先選「序號＝上一筆＋1」，再選純數字學號，最後照 k 順序。"""
    if not cands:
        return None
    if last_seq:
        for c in cands:
            if c[3] == last_seq + 1:
                return c
    # 有「序號黏在學號後面」的候選時，不要退回 k=0（那代表整行沒讀到序號）
    pool = [c for c in cands if c[0] > 0] or cands
    for c in pool:
        if c[2].isdigit():
            return c
    return pool[0]


def _prev_name_line(lines, i, used, span=1):
    """往上找「像姓名」的短行（不含數字）。span=1 只看緊鄰的前一個非空行。"""
    looked = 0
    for j in range(i - 1, -1, -1):
        if j in used:
            continue
        t = _clean_name(lines[j])
        if not t:
            continue
        looked += 1
        if looked > span:
            return -1, ""
        if re.search(r"\d", lines[j] or ""):
            return -1, ""
        if _name_ok(t) and CJK_RE.search(t):
            return j, t
        return -1, ""
    return -1, ""


def _lookback_name(lines, i, used):
    """PyMuPDF 式版面：姓名可能單獨在學號的前幾行。"""
    for j in range(i - 1, max(-1, i - 5), -1):
        if j in used:
            continue
        t = NOTE_MARK_RE.sub("", lines[j]).strip()
        t = re.sub(r"\s+", "", t)
        if not t:
            continue
        if PURE_NAME_RE.match(t) and _looks_like_name(t):
            used.add(j)
            return t
    return ""


def _tail_name(before):
    """從學號左邊的字串取出姓名（取結尾那一段中文／英文）。"""
    t = NOTE_MARK_RE.sub(" ", before or "")
    t = t.replace("　", " ").strip()
    t = re.sub(r"\s+", "", t)
    if not t:
        return ""
    m = NAME_TAIL_RE.search(t)
    if not m:
        return ""
    cand = m.group(0)
    # 姓名前面若黏了「班級」等字樣，逐字往右縮短找出合理的姓名
    while cand and not _looks_like_name(cand):
        cand = cand[1:]
    return cand if _looks_like_name(cand) else ""


def _seq_around(lines, i, before, after):
    """k=0（序號沒有黏在學號後面）時，往左右與下方找序號。"""
    mm = LEAD_INT_RE.match(after or "")
    if mm:
        return int(mm.group(1)), (after or "")[mm.end():]
    head = LEAD_INT_RE.match((before or "").lstrip())
    if head:
        return int(head.group(1)), after
    for j in range(i + 1, min(len(lines), i + 3)):
        pm = PURE_INT_RE.match(lines[j] or "")
        if pm:
            return int(pm.group(1)), after
        if (lines[j] or "").strip():
            break
    return 0, after


def _class_from(rest, lines, i, glued=False):
    t = (rest or "").strip()
    t = re.sub(r"\s+", "", t)
    if t and len(t) <= 16 and (CLASS_HINT_RE.search(t)
                               or (glued and PURE_CJK_RE.match(t))):
        return t
    for j in range(i + 1, min(len(lines), i + 4)):
        s = re.sub(r"\s+", "", NOTE_MARK_RE.sub("", lines[j]).strip())
        if not s or PURE_INT_RE.match(s):
            continue                       # 空行／單獨一行的序號 → 繼續往下看
        if SID9_RE.search(s):
            break                          # 已經是下一位學生
        if CLASS_STRONG_RE.search(s) and len(s) <= 16 and not re.search(r"\d{4,}", s):
            return s
        break
    return ""


def parse_pdf_records(lines):
    """把「抽字後的行列表」解析成 [Student]（不編號）。對行序與版面有容忍度。

    支援的列形狀（中＝中文字、9＝數字、A＝英文字母）：

        中中中99999999999 中中中中中     姓名＋9 碼學號＋2 碼序號 黏在一起，後面班級
        中中中9999999999 中中中中中      序號 1 碼
        中中中 【註*】9999999999 中…     註記夾在姓名與學號之間（可跨行被切斷）
        中中中 中中99999999999 中…       姓名含空白（族名、複姓分寫）
        中中中99999A9999 中…             學號第 6 碼是英文字母
        中中中A9999999999 中…            學號＝1 字母＋8 數字
        99999999999 中中中中中           這一列沒有姓名（文字層抽不到，罕用字）
        中中中 Aaaaa / Aaaaaa99999…      姓名折行（中文名＋英文名被拆成兩行）
        中中中 / 990054001 / 1 / 自資系大三   姓名／學號／序號／班級各自一行
    """
    lines = normalize_lines(lines)
    used_name_lines = set()
    recs = []
    last_seq = 0
    for i, raw in enumerate(lines):
        line = NOTE_MARK_RE.sub(" ", raw or "")
        best = None
        for m in ALNUM_RUN_RE.finditer(line):
            if len(m.group(0)) < SID_LEN:
                continue
            if best is None or len(m.group(0)) > len(best.group(0)):
                best = m
        if best is None:
            continue
        cand = _pick_glued(_split_glued(best.group(0)), last_seq)
        if cand is None:
            continue
        k, prefix, sid, seq = cand
        before = line[:best.start()]
        after = line[best.end():]
        if not seq:
            seq, after = _seq_around(lines, i, before, after)

        # ---- 姓名 -------------------------------------------------------
        head = _clean_name(before)
        if CJK_RE.search(head):
            name = _clean_name(head + prefix)
            if not _name_ok(name):
                name = _tail_name(head + prefix)
        elif line.strip() == best.group(0):
            # 整行就是一個學號 → 分行版面，姓名在前幾行
            name = _lookback_name(lines, i, used_name_lines)
        else:
            # 學號左邊只有英文字母／空白 → 可能是「姓名折行」
            j, prev = _prev_name_line(lines, i, used_name_lines, span=1)
            tail = _clean_name((head + " " + prefix).strip())
            if j >= 0:
                used_name_lines.add(j)
                name = (prev + " " + tail).strip() if tail else prev
            else:
                name = tail if _name_ok(tail) else ""
        if name and not _name_ok(name):
            name = ""

        recs.append(Student(seq=seq or 0, sid=sid, name=name,
                            klass=_class_from(after, lines, i, glued=bool(k))))
        if seq:
            last_seq = seq
    return recs


def seq_plan(recs):
    """決定要「依序號」還是「依學號排序」，並回傳提醒文字。

    * 每位都有序號且不重複 → 依序號（**允許缺號**，缺號只提醒不擋）。
    * 序號有重複或部分缺漏 → 依學號排序給 1..N。
    """
    notes = []
    seqs = [r.seq for r in recs if r.seq]
    if seqs and len(seqs) == len(recs) and len(set(seqs)) == len(recs):
        missing = [n for n in range(1, max(seqs) + 1) if n not in set(seqs)]
        if missing:
            shown = "、".join(str(n) for n in missing[:20])
            if len(missing) > 20:
                shown += "…"
            notes.append(MISSING_SEQ_NOTE_FMT.format(seqs=shown))
        return True, notes
    notes.append(DUP_SEQ_NOTE if seqs else NO_SEQ_NOTE)
    return False, notes


def records_to_data(recs, source="", kind="pdf"):
    """把解析出來的紀錄包成 RosterData。只有「一位學生都讀不到」才丟錯。"""
    if not recs:
        raise RosterError(PDF_FAIL_MSG)
    notes = []
    for r in recs:
        if not r.name:
            notes.append(NO_NAME_NOTE_FMT.format(seq=r.seq or "（未知）"))
    _, seq_notes = seq_plan(recs)
    notes.extend(seq_notes)
    return RosterData(records=recs, source=os.path.abspath(source) if source else "",
                      kind=kind, notes=notes)


def parse_pdf(path):
    return records_to_data(parse_pdf_records(pdf_lines(path)), path, "pdf")


# --------------------------------------------------------------------------
# Word（.docx）解析
# --------------------------------------------------------------------------
def _docx_module():
    try:
        import docx                               # noqa: F401
    except ImportError:
        raise RosterError(DOCX_NEED_PKG)
    return docx


def docx_document(path):
    docx = _docx_module()
    try:
        return docx.Document(path)
    except Exception as e:
        raise RosterError(
            f"這個 Word 檔打不開（可能不是有效的 .docx，或正被 Word 開著）：\n{e}")


def _cell_text(cell):
    parts = [p.text for p in cell.paragraphs]
    for t in getattr(cell, "tables", []):         # 巢狀表格
        for row in t.rows:
            for c in row.cells:
                parts.append(_cell_text(c))
    return "\n".join(x for x in parts if x is not None).strip()


def _iter_tables(container):
    for t in getattr(container, "tables", []):
        yield t
        for row in t.rows:
            for c in row.cells:
                for sub in _iter_tables(c):
                    yield sub


def _docx_table_rows(table):
    rows = []
    for row in table.rows:
        try:
            cells = list(row.cells)
        except Exception:
            continue
        rows.append([_cell_text(c) for c in cells])
    return rows


def docx_text_lines(doc):
    """段落＋表格儲存格的文字，逐行攤平（給 PDF 那一套行解析用）。"""
    lines = []
    for p in doc.paragraphs:
        lines.extend((p.text or "").split("\n"))
    for t in _iter_tables(doc):
        for row in _docx_table_rows(t):
            cells = [c for c in row if c.strip()]
            if not cells:
                continue
            lines.append(" ".join(c.replace("\n", " ") for c in cells))
            lines.extend(cells)
    return lines


def parse_docx(path):
    doc = docx_document(path)
    for t in _iter_tables(doc):
        rows = _docx_table_rows(t)
        hdr, col = _find_header(rows)
        if hdr < 0:
            continue
        recs = _records_from_rows(rows, hdr, col)
        if not recs:
            continue
        notes = [NO_NAME_NOTE_FMT.format(seq=r.seq or "（未知）")
                 for r in recs if not r.name]
        if col["code"] >= 0 and all(r.code for r in recs):
            return RosterData(records=recs, source=os.path.abspath(path),
                              kind="codebook", rule="沿用對照表", notes=notes)
        return RosterData(records=recs, source=os.path.abspath(path), kind="docx",
                          notes=notes)
    recs = parse_pdf_records(docx_text_lines(doc))
    if recs:
        return records_to_data(recs, path, "docx")
    raise RosterError(
        "這份 Word 名單讀不出「學號」與「姓名」。\n"
        "請確認名單是一張表格，且表頭同一列有「學號」與「姓名」"
        "（可另有 序號、班級、電子郵件）；或改用 Excel／CSV 名單。")


# --------------------------------------------------------------------------
# Excel／CSV 解析
# --------------------------------------------------------------------------
def _rows_from_csv(path):
    raw = None
    for enc in ("utf-8-sig", "cp950", "utf-8"):
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                raw = list(csv.reader(f))
            break
        except UnicodeDecodeError:
            continue
    if raw is None:
        with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
            raw = list(csv.reader(f))
    return [[c if c != "" else None for c in row] for row in raw]


def _sheets_from_excel(path):
    """回傳 [(工作表名, rows)]；read_only 開檔後務必關閉，否則檔案會被自己鎖住。"""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise RosterError("無法開啟名單檔（可能不是有效的 Excel 檔，或正被 Excel 開著）：\n"
                          f"{e}")
    out = []
    try:
        for ws in wb.worksheets:
            out.append((ws.title, [list(r) for r in ws.iter_rows(values_only=True)]))
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return out


def _find_header(rows):
    """在第 1–5 列內找同時有「學號」與「姓名」的表頭列。回傳 (列索引, 欄位對應)。"""
    for r in range(min(5, len(rows))):
        cells = [text_of(c) for c in rows[r]]
        cid = cname = cseq = ccls = ccode = -1
        cdept = cyear = -1
        for c, t in enumerate(cells):
            if not t:
                continue
            if ccode < 0 and CODE_HDR_RE.match(t):
                ccode = c
            if cid < 0 and ID_HDR_RE.search(t) and not CODE_HDR_RE.match(t):
                cid = c
            if cname < 0 and NAME_HDR_RE.match(t):
                cname = c
            if cseq < 0 and SEQ_HDR_RE.match(t):
                cseq = c
            if CLASS_HDR_RE.search(t):
                if "系所" in t or "科系" in t:
                    if cdept < 0:
                        cdept = c
                elif "年級" in t:
                    if cyear < 0:
                        cyear = c
                elif ccls < 0:
                    ccls = c
        if cid >= 0 and cname >= 0:
            return r, {"sid": cid, "name": cname, "seq": cseq, "klass": ccls,
                       "dept": cdept, "year": cyear, "code": ccode}
    return -1, None


def _records_from_rows(rows, hdr_row, col):
    recs = []
    for r in range(hdr_row + 1, len(rows)):
        row = rows[r]

        def get(key):
            c = col.get(key, -1)
            if c < 0 or c >= len(row):
                return None
            return row[c]

        sid = normalize_id(get("sid"))
        name = text_of(get("name"))
        if not sid:
            continue
        if not name and not SID_SHAPE_RE.match(sid):
            continue          # 沒有姓名又不像學號 → 統計列／雜訊列
        if ID_HDR_RE.search(name) or NAME_HDR_RE.match(name):
            continue          # 同一張表裡重複出現的表頭列
        klass = text_of(get("klass"))
        if not klass:
            dept, year = text_of(get("dept")), text_of(get("year"))
            klass = (dept + year).strip()
        seq_raw = text_of(get("seq"))
        seq = int(seq_raw) if seq_raw.isdigit() else 0
        code = text_of(get("code"))
        recs.append(Student(seq=seq, code=code, sid=sid, name=name, klass=klass))
    return recs


def parse_table(path):
    low = path.lower()
    if low.endswith(".csv"):
        sheets = [("CSV", _rows_from_csv(path))]
        kind = "csv"
    else:
        sheets = _sheets_from_excel(path)
        kind = "excel"
    for title, rows in sheets:
        hdr, col = _find_header(rows)
        if hdr < 0:
            continue
        recs = _records_from_rows(rows, hdr, col)
        if not recs:
            continue
        notes = []
        if len(sheets) > 1:
            notes.append(f"名單取自工作表「{title}」。")
        notes.extend(NO_NAME_NOTE_FMT.format(seq=r.seq or "（未知）")
                     for r in recs if not r.name)
        if col["code"] >= 0 and all(r.code for r in recs):
            return RosterData(records=recs, source=os.path.abspath(path),
                              kind="codebook", rule="沿用對照表", notes=notes)
        return RosterData(records=recs, source=os.path.abspath(path), kind=kind,
                          notes=notes)
    raise RosterError(
        "這份名單讀不出「學號」與「姓名」欄。\n"
        "請確認名單的前 5 列內有一列同時含有「學號」與「姓名」標題"
        "（可另有 序號、班級、電子郵件）。")


# --------------------------------------------------------------------------
# 解析 + 編號
# --------------------------------------------------------------------------
def parse_roster(path):
    p = os.path.abspath(path or "")
    if not os.path.isfile(p):
        raise RosterError(f"找不到名單檔：{p}")
    low = p.lower()
    if low.endswith(".doc"):
        raise RosterError(DOC_OLD_MSG)
    if not low.endswith(ROSTER_EXT):
        raise RosterError("名單只支援 .xlsx / .xlsm / .csv / .pdf / .docx"
                          "（舊版 .xls、.doc 請先用 Office 另存為新格式）。")
    if low.endswith(".pdf"):
        data = parse_pdf(p)
    elif low.endswith(".docx"):
        data = parse_docx(p)
    else:
        data = parse_table(p)
    if not data.records:
        raise RosterError("這份名單一位學生都讀不到，請確認檔案內容。")
    return data


def assign_codes(data, semester, course):
    """依 §1.2 規則給固定學生編號，就地寫回 data.records 並回傳 list[Student]。"""
    recs = data.records
    seen = {}
    for r in recs:
        r.sid = normalize_id(r.sid)
        if not r.sid:
            raise RosterError(f"名單中有一位學生沒有學號（姓名第 1 字「{r.name[:1]}」），"
                              "請補上學號後再試。")
        if r.sid in seen:
            raise RosterError("名單裡有重複的學號（同一個學號出現 2 次以上），"
                              "請先修正名單再產生對照表。")
        seen[r.sid] = True

    if data.kind == "codebook":
        data.rule = "沿用對照表"
        return recs

    use_seq, notes = seq_plan(recs)
    for n in notes:
        if n not in data.notes:
            data.notes.append(n)
    if use_seq:
        recs.sort(key=lambda r: r.seq)
        data.rule = "依序號"
    else:
        recs.sort(key=lambda r: r.sid)
        for i, r in enumerate(recs, start=1):
            r.seq = i
        data.rule = "依學號排序"
    for r in recs:
        r.code = f"{semester}_{course}_{r.seq}"
    return recs


# --------------------------------------------------------------------------
# 對照表讀寫
# --------------------------------------------------------------------------
def codebook_path_for(roster_path, semester, course):
    folder = os.path.dirname(os.path.abspath(roster_path)) if roster_path else os.getcwd()
    return os.path.join(folder, CODEBOOK_NAME_FMT.format(sem=semester, course=course))


def _note_lines(semester, course, source, rule):
    return [
        ["這是什麼"],
        [f"　{semester} 學年課程「{course}」的『學生名單與學生編號對照表』。"],
        ["　由「姓名遮罩與學生編號 NameMasker」自動產生／維護。"],
        [""],
        ["編號規則"],
        [f"　學生編號 = {semester}_{course}_{{序號}}；本表的編號來源：{rule}。"],
        ["　1. 名單每位都有「序號」且不重複 → 直接用序號（允許缺號）。"],
        ["　2. 否則依「學號」字串遞增排序給 1..N。"],
        [f"　3. 不在名單的作答者（退選、旁聽等）自 {OUTSIDER_START} 號起遞增，"
         f"登記在「{OUT_SHEET}」工作表，跨檔跨週一致。"],
        ["　4. 名單檔讀不到姓名的學生（罕用字）仍會保留，姓名欄空白，請自行補上。"],
        [""],
        ["名單來源"],
        [f"　{os.path.basename(source) if source else '（未記錄）'}"],
        [""],
        ["⚠ 含真名學號＝再識別鑰匙，只留本機"],
        ["　本檔可以把學生編號換回真實姓名與學號，屬於個人資料。"],
        ["　請只留在自己的電腦，不要上傳雲端、不要寄給別人、不要跟分析結果一起交出去。"],
    ]


def _fit_widths(ws, widths):
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, start=1):
        try:
            ws.column_dimensions[get_column_letter(i)].width = w
        except Exception:
            pass


def write_codebook(path, students, semester, course, source="", rule="", outsiders=None):
    """全新建立（或覆寫）對照表。"""
    folder = os.path.dirname(os.path.abspath(path))
    if folder:
        os.makedirs(folder, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = CODE_SHEET
    ws.append(list(CODE_HEADERS))
    for s in students:
        ws.append([s.seq, s.code, s.sid, s.name, s.klass])
    _fit_widths(ws, (8, 18, 14, 14, 16))
    try:
        ws.freeze_panes = "A2"
    except Exception:
        pass

    ws2 = wb.create_sheet(OUT_SHEET)
    ws2.append(list(OUT_HEADERS))
    for o in (outsiders or []):
        ws2.append([o.code, o.sid, o.name, o.src])
    _fit_widths(ws2, (18, 14, 14, 34))

    ws3 = wb.create_sheet(NOTE_SHEET)
    for row in _note_lines(semester, course, source, rule):
        ws3.append(row)
    _fit_widths(ws3, (72,))

    try:
        wb.save(path)
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return path


def read_codebook(path):
    """讀既有對照表 → (students, outsiders)。"""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise RosterError(f"無法開啟對照表（可能正被 Excel 開著）：\n{e}")
    try:
        if CODE_SHEET in wb.sheetnames:
            rows = [list(r) for r in wb[CODE_SHEET].iter_rows(values_only=True)]
        else:
            rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        out_rows = ([list(r) for r in wb[OUT_SHEET].iter_rows(values_only=True)]
                    if OUT_SHEET in wb.sheetnames else [])
    finally:
        try:
            wb.close()
        except Exception:
            pass

    hdr, col = _find_header(rows)
    if hdr < 0 or col["code"] < 0:
        raise RosterError(
            f"這個檔案看起來不是對照表：工作表「{CODE_SHEET}」要有 "
            "「學生編號」「學號」「姓名」欄。")
    students = _records_from_rows(rows, hdr, col)
    for s in students:
        if not s.seq:
            m = re.search(r"_(\d+)$", s.code or "")
            s.seq = int(m.group(1)) if m else 0
    students.sort(key=lambda s: (s.seq or 0))

    outsiders = []
    if out_rows:
        h2 = [text_of(c) for c in out_rows[0]]
        try:
            ic, ii, inm = h2.index("學生編號"), h2.index("學號"), h2.index("姓名")
            isrc = h2.index("首次出現檔案") if "首次出現檔案" in h2 else -1
        except ValueError:
            ic = ii = inm = isrc = -1
        if ic >= 0:
            for r in out_rows[1:]:
                code = text_of(r[ic]) if ic < len(r) else ""
                if not code:
                    continue
                outsiders.append(Outsider(
                    code=code,
                    sid=normalize_id(r[ii]) if 0 <= ii < len(r) else "",
                    name=text_of(r[inm]) if 0 <= inm < len(r) else "",
                    src=text_of(r[isrc]) if 0 <= isrc < len(r) else ""))
    return students, outsiders


def _update_outsider_sheet(path, outsiders):
    """只改寫「名單外作答者」工作表，其餘工作表與欄位原樣保留。"""
    wb = openpyxl.load_workbook(path)
    try:
        if OUT_SHEET in wb.sheetnames:
            del wb[OUT_SHEET]
        ws = wb.create_sheet(OUT_SHEET)
        ws.append(list(OUT_HEADERS))
        for o in outsiders:
            ws.append([o.code, o.sid, o.name, o.src])
        _fit_widths(ws, (18, 14, 14, 34))
        wb.save(path)
    finally:
        try:
            wb.close()
        except Exception:
            pass


# --------------------------------------------------------------------------
# Codebook：遮罩流程實際使用的物件
# --------------------------------------------------------------------------
class Codebook:
    """對照表（名單內固定編號 + 名單外作答者持久登記）。"""

    def __init__(self, path, semester, course, students, outsiders=None,
                 source="", rule="", created=False):
        self.path = os.path.abspath(path) if path else ""
        self.semester = semester
        self.course = course
        self.students = list(students)
        self.outsiders = list(outsiders or [])
        self.source = source
        self.rule = rule
        self.created = created          # 這次是否新建／覆寫了檔案
        self.dirty = False
        self._rebuild()

    # ---- 索引 ----
    def _rebuild(self):
        self.stu_by_id, self.stu_by_name = {}, {}
        self.out_by_id, self.out_by_name = {}, {}
        for s in self.students:
            if s.sid:
                self.stu_by_id[normalize_id(s.sid)] = s.code
            if s.name:
                self.stu_by_name.setdefault(s.name, s.code)
        for o in self.outsiders:
            if o.sid:
                self.out_by_id.setdefault(normalize_id(o.sid), o.code)
            if o.name:
                self.out_by_name.setdefault(o.name, o.code)

    @property
    def by_id(self):
        d = dict(self.out_by_id)
        d.update(self.stu_by_id)
        return d

    # HomeworkWordCloud 相容別名
    by_sid = by_id

    @property
    def by_name(self):
        d = dict(self.out_by_name)
        d.update(self.stu_by_name)
        return d

    @property
    def outside(self):
        """HomeworkWordCloud 相容：名單外作答者清單。"""
        return self.outsiders

    @property
    def count(self):
        return len(self.students)

    # ---- 查詢 ----
    def roster_code_for(self, sid, name=""):
        """只查「名單內」學生。"""
        sid = normalize_id(sid)
        if sid and sid in self.stu_by_id:
            return self.stu_by_id[sid]
        if name and name in self.stu_by_name:
            return self.stu_by_name[name]
        return ""

    def outsider_code_for(self, sid, name=""):
        sid = normalize_id(sid)
        if sid and sid in self.out_by_id:
            return self.out_by_id[sid]
        if name and name in self.out_by_name:
            return self.out_by_name[name]
        return ""

    def code_for(self, sid, name="", source_file=None):
        """查編號。給了 `source_file`（可以是空字串）→ 查不到就登記為名單外作答者。"""
        code = self.roster_code_for(sid, name) or self.outsider_code_for(sid, name)
        if code or source_file is None:
            return code
        if not normalize_id(sid) and not (name or "").strip():
            return ""
        return self.register_outsider(sid, name, source_file)

    def in_roster(self, sid, name=""):
        return bool(self.roster_code_for(sid, name))

    def roster_names(self):
        """{姓名: 學生編號}（名單內所有人，供自由文字替換用）。"""
        return {s.name: s.code for s in self.students if s.name}

    def all_names(self):
        d = self.roster_names()
        for o in self.outsiders:
            if o.name:
                d.setdefault(o.name, o.code)
        return d

    def roster_ids(self):
        return {normalize_id(s.sid): s.code for s in self.students if s.sid}

    # HomeworkWordCloud 相容介面
    def names(self):
        return ([s.name for s in self.students if s.name]
                + [o.name for o in self.outsiders if o.name])

    def sids(self):
        return ([normalize_id(s.sid) for s in self.students if s.sid]
                + [normalize_id(o.sid) for o in self.outsiders if o.sid])

    def name_to_code(self):
        return dict(self.all_names())

    def codes(self):
        return [s.code for s in self.students]

    # ---- 名單外作答者持久登記 ----
    def _next_outsider_no(self):
        n = OUTSIDER_START - 1
        for o in self.outsiders:
            m = re.search(r"_(\d+)$", o.code or "")
            if m:
                n = max(n, int(m.group(1)))
        return max(n + 1, OUTSIDER_START)

    def register_outsider(self, sid, name="", src=""):
        """不在名單的作答者：已登記過就回既有編號，否則新編（101 起）並標記需回寫。"""
        code = self.roster_code_for(sid, name) or self.outsider_code_for(sid, name)
        if code:
            return code
        code = f"{self.semester}_{self.course}_{self._next_outsider_no()}"
        o = Outsider(code=code, sid=normalize_id(sid), name=name, src=src)
        self.outsiders.append(o)
        if o.sid:
            self.out_by_id[o.sid] = code
        if o.name:
            self.out_by_name.setdefault(o.name, code)
        self.dirty = True
        return code

    # ---- 存檔 ----
    def save(self, force=False):
        """把名單外登記回寫檔案；沒有變動就不動檔案。回傳是否真的寫了。"""
        if not self.path:
            return False
        if not (self.dirty or force):
            return False
        if os.path.exists(self.path):
            _update_outsider_sheet(self.path, self.outsiders)
        else:
            write_codebook(self.path, self.students, self.semester, self.course,
                           self.source, self.rule, self.outsiders)
        self.dirty = False
        return True

    def save_if_dirty(self, log=print):
        """HomeworkWordCloud 相容：有新登記才回寫，並印一行訊息。"""
        if not (self.dirty and self.path):
            return False
        self.save()
        log(f"  對照表已回寫（名單外作答者 {len(self.outsiders)} 人）：{self.path}")
        return True

    def summary_lines(self):
        L = [f"對照表：{self.path}",
             f"名單內學生 {len(self.students)} 人（編號規則：{self.rule or '沿用既有對照表'}）"]
        if self.outsiders:
            L.append(f"名單外作答者 {len(self.outsiders)} 人"
                     f"（{OUTSIDER_START} 號起，已持久登記）")
        return L


# --------------------------------------------------------------------------
# 一站式入口：由名單／對照表取得 Codebook
# --------------------------------------------------------------------------
def build_codebook(roster_path=None, codebook_path=None, semester="115-1", course="EC",
                   overwrite=False):
    """回傳 (Codebook, 訊息行 list)。

    * 只給 codebook_path → 直接讀既有對照表。
    * 只給 roster_path   → 對照表預設放在名單同資料夾；已存在且未指定覆寫 → 沿用既有檔。
    * 兩者都給           → 以 codebook_path 為輸出位置，名單用來建檔（若檔案不存在或覆寫）。
    """
    msgs = []
    if not roster_path and not codebook_path:
        raise RosterError("請先選一份「原始名單」（Excel／CSV／Word／選課名單 PDF）"
                          "或既有的「學生編號對照表」。")

    if not roster_path:
        cb_path = os.path.abspath(codebook_path)
        if not os.path.isfile(cb_path):
            raise RosterError(f"找不到對照表：{cb_path}")
        students, outsiders = read_codebook(cb_path)
        if not students:
            raise RosterError(f"對照表「{os.path.basename(cb_path)}」裡沒有任何學生資料。")
        msgs.append(f"沿用既有對照表（{len(students)} 人）：{cb_path}")
        return Codebook(cb_path, semester, course, students, outsiders,
                        source=cb_path, rule="沿用既有對照表"), msgs

    data = parse_roster(roster_path)
    if data.kind == "codebook" and not codebook_path:
        # TA 直接把「對照表」當名單載入 → 原檔就是對照表，直接採用、不重編
        cb_path = os.path.abspath(roster_path)
        students, outsiders = read_codebook(cb_path)
        msgs.append(f"載入的檔案本身就是對照表，直接採用（{len(students)} 人）：{cb_path}")
        return Codebook(cb_path, semester, course, students, outsiders,
                        source=cb_path, rule="沿用既有對照表"), msgs
    cb_path = os.path.abspath(codebook_path) if codebook_path else \
        codebook_path_for(roster_path, semester, course)

    if os.path.isfile(cb_path) and not overwrite:
        students, outsiders = read_codebook(cb_path)
        if students:
            msgs.extend(data.notes)
            msgs.append(f"對照表已存在，沿用既有編號（{len(students)} 人）：{cb_path}")
            msgs.append("（要用新名單重新編號，請勾選／加上 --overwrite-codebook）")
            return Codebook(cb_path, semester, course, students, outsiders,
                            source=cb_path, rule="沿用既有對照表"), msgs

    students = assign_codes(data, semester, course)
    msgs.extend(data.notes)             # assign_codes 也會補上序號相關提醒
    outsiders = []
    if os.path.isfile(cb_path):
        try:
            _, outsiders = read_codebook(cb_path)
            if outsiders:
                msgs.append(f"保留既有的名單外作答者登記 {len(outsiders)} 筆。")
        except RosterError:
            outsiders = []
    write_codebook(cb_path, students, semester, course, data.source, data.rule, outsiders)
    msgs.append(f"已產生對照表（{len(students)} 人、{data.rule}）：{cb_path}")
    msgs.append("⚠ 對照表含真名與學號＝再識別鑰匙，只留本機，不要上傳。")
    cb = Codebook(cb_path, semester, course, students, outsiders,
                  source=data.source, rule=data.rule, created=True)
    return cb, msgs


# --------------------------------------------------------------------------
# HomeworkWordCloud 相容層（兩支 App 共用同一份 roster.py）
# --------------------------------------------------------------------------
OUTSIDE_START = OUTSIDER_START
CODEBOOK_SHEET = CODE_SHEET
OUTSIDE_SHEET = OUT_SHEET
CODEBOOK_COLS = list(CODE_HEADERS)
OUTSIDE_COLS = list(OUT_HEADERS)

norm_text = text_of
normalize_sid = normalize_id


def norm_course_code(s, default="EC"):
    s = text_of(s).upper()
    return s if re.match(r"^[A-Z]{2}$", s) else default


def norm_semester(s, default="115-1"):
    s = text_of(s).replace("／", "/").replace("/", "-")
    return s if re.match(r"^\d{3}-[12]$", s) else default


def make_code(semester, course, n):
    return f"{norm_semester(semester)}_{norm_course_code(course)}_{int(n)}"


def codebook_filename(semester, course):
    return CODEBOOK_NAME_FMT.format(sem=norm_semester(semester),
                                    course=norm_course_code(course))


def parse_pdf_lines(lines):
    """HomeworkWordCloud 相容：回傳 [{序號, 學號, 姓名, 班級, 學生編號}]。"""
    return [{"序號": r.seq or "", "學號": r.sid, "姓名": r.name,
             "班級": r.klass, "學生編號": r.code} for r in parse_pdf_records(lines)]


def _guess_sem_course(students, semester="", course=""):
    """從既有編號反推學期與課程縮寫（對照表可能來自別的課）。"""
    for s in students:
        m = re.match(r"^(\d{3}-[12])_([A-Z]{2})_\d+$", text_of(s.code))
        if m:
            return m.group(1), m.group(2)
    return norm_semester(semester), norm_course_code(course)


def load_codebook(path):
    """直接讀一份既有的對照表 xlsx。"""
    students, outsiders = read_codebook(path)
    sem, crs = _guess_sem_course(students)
    return Codebook(os.path.abspath(path), sem, crs, students, outsiders,
                    source=os.path.abspath(path), rule="沿用既有對照表")


def prepare_codebook(path, semester, course, overwrite=False, out_dir=None,
                     log=print):
    """TA 給一份「原始名單或學生編號對照表」→ 回傳可用的 Codebook。

    `out_dir` 指定對照表的輸出資料夾（預設放在名單同資料夾）。
    """
    path = os.path.abspath(path)
    semester, course = norm_semester(semester), norm_course_code(course)
    data = parse_roster(path)

    if data.kind == "codebook":
        cb = load_codebook(path)
        log(f"  對照表：{os.path.basename(path)}（{cb.count} 人，"
            f"名單外 {len(cb.outsiders)} 人）")
        log(f"  學生編號規則：{cb.semester}_{cb.course}_{{序號}}（沿用既有對照表，不重編）")
        return cb

    students = assign_codes(data, semester, course)
    for n in data.notes:
        log("  [提醒] " + n)
    folder = out_dir or os.path.dirname(path) or os.getcwd()
    target = os.path.join(folder, codebook_filename(semester, course))

    if os.path.isfile(target) and not overwrite:
        try:
            old = load_codebook(target)
        except RosterError:
            old = None
        if old is not None and old.count:
            log(f"  對照表已存在，直接沿用（不重編）：{target}")
            log(f"    名單檔 {os.path.basename(path)} 讀到 {len(students)} 人、"
                f"既有對照表 {old.count} 人、名單外 {len(old.outsiders)} 人")
            if old.count != len(students):
                log("    [提醒] 兩者人數不同（可能有加退選）。要改用新名單重編，"
                    "請勾選／加上「覆寫對照表」。")
            return old

    outsiders = []
    if os.path.isfile(target):
        try:
            _, outsiders = read_codebook(target)
        except RosterError:
            outsiders = []
    write_codebook(target, students, semester, course, path, data.rule, outsiders)
    log(f"  名單檔：{os.path.basename(path)}　讀到 {len(students)} 人")
    log(f"  學生編號規則：{semester}_{course}_{{序號}}（{data.rule}）")
    log(f"  對照表已產生：{target}")
    return Codebook(target, semester, course, students, outsiders,
                    source=path, rule=data.rule, created=True)


def describe(cb):
    """給 GUI 顯示的一行摘要。"""
    if cb is None:
        return "尚未載入名單／對照表"
    return (f"讀到 {cb.count} 人　編號規則：{cb.semester}_{cb.course}_{{序號}}"
            f"（{cb.rule}）　名單外 {len(cb.outsiders)} 人")


# --------------------------------------------------------------------------
def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    if not argv:
        print("用法：python roster.py 名單.(xlsx|csv|pdf|docx) "
              "[--course EC --semester 115-1] [--overwrite-codebook]")
        return 0
    course, sem, ow = "EC", "115-1", False
    files = []
    i = 0
    while i < len(argv):
        a = argv[i]
        if a == "--course" and i + 1 < len(argv):
            course = argv[i + 1]
            i += 2
            continue
        if a == "--semester" and i + 1 < len(argv):
            sem = argv[i + 1]
            i += 2
            continue
        if a == "--overwrite-codebook":
            ow = True
            i += 1
            continue
        if a.startswith("-"):
            i += 1
            continue
        files.append(a)
        i += 1
    for f in files:
        try:
            cb, msgs = build_codebook(roster_path=f, semester=sem, course=course,
                                      overwrite=ow)
        except RosterError as e:
            print(f"失敗：{f}\n{e}")
            continue
        for m in msgs:
            print(m)
        for line in cb.summary_lines():
            print(line)
    return 0


if __name__ == "__main__":
    sys.exit(main())
