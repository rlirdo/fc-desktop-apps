#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
mask_names.py — 將 Excel「姓名」欄的中文姓名第二個字改為 O
============================================================
流程（固定、確定性，規則與原版完全一致）：
  1. 開啟你拖進來（或指令列指定）的 .xlsx 檔
  2. 取「第 1 個工作表」，在「第 1 列」尋找內容為「姓名」的儲存格（找不到才退而求其次：含「姓名」字樣）
  3. 該欄第 2 列起，所有中文姓名的第 2 個字改成 O（王小明→王O明、李明→李O、歐陽小花→歐O小花）
  4. 遮罩結果放在新工作表「原表名_遮罩」；原工作表完全不動
  5. 整本另存新檔：「原檔名02.xlsx」，存在與原檔相同的資料夾（原檔案完全不被修改）

隱私聲明：本程式完全離線，只讀寫你電腦上的檔案，不連網、不上傳任何資料。
需求：Python 3.8+、openpyxl（pip install openpyxl）

用法：
  python mask_names.py 檔案.xlsx
  或使用 GUI：python app.py
"""
import os
import re
import sys

# Windows 主控台編碼保險
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - 封裝後不會發生
    sys.exit("缺少 openpyxl，請先執行：pip install openpyxl")

NAME_RUN = re.compile(r"[一-鿿㐀-䶿]{2,}")   # 連續 2 個以上中文字

SUPPORTED_EXT = (".xlsx", ".xlsm")


class MaskError(Exception):
    """可直接顯示給使用者看的錯誤訊息。"""


def mask_name(text):
    """把字串中第一段連續中文字的第 2 個字改為 O；無中文則原樣返回。"""
    if not isinstance(text, str):
        return text, False
    m = NAME_RUN.search(text)
    if not m:
        return text, False
    run = m.group(0)
    masked = run[0] + "O" + run[2:]
    return text[: m.start()] + masked + text[m.start() + len(run):], True


def find_name_col(ws):
    """在第 1 個工作表的第 1 列找「姓名」欄；回傳 (欄索引, 標題文字)。"""
    exact, fuzzy = None, None
    for cell in ws[1]:
        v = cell.value
        if not isinstance(v, str):
            continue
        t = v.strip()
        if t == "姓名" and exact is None:
            exact = (cell.column, t)
        elif "姓名" in t and fuzzy is None:
            fuzzy = (cell.column, t)
    return exact or fuzzy


def next_output_path(src):
    """另存新檔路徑：原檔名 + 02（已存在則 03、04…），與原檔同資料夾。"""
    folder, fname = os.path.split(os.path.abspath(src))
    stem, ext = os.path.splitext(fname)
    n = 2
    while True:
        dst = os.path.join(folder, f"{stem}{n:02d}{ext}")
        if not os.path.exists(dst):
            return dst
        n += 1


def mask_file(path):
    """對單一 Excel 檔案執行遮罩，回傳結果 dict。

    回傳欄位：
      src         原檔完整路徑
      dst         另存的新檔完整路徑
      masked      遮罩筆數
      sheet       新工作表名稱
      column      姓名欄的欄位字母（如 "C"）
      header      姓名欄標題文字
    失敗時丟出 MaskError（訊息可直接顯示給使用者）。
    """
    src = os.path.abspath(path)
    if not os.path.isfile(src):
        raise MaskError(f"找不到檔案：{src}")
    if not src.lower().endswith(SUPPORTED_EXT):
        raise MaskError("只支援 .xlsx / .xlsm 檔案（舊版 .xls 請先用 Excel 另存為 .xlsx）")

    keep_vba = src.lower().endswith(".xlsm")
    try:
        wb = openpyxl.load_workbook(src, keep_vba=keep_vba)         # 完整保留原內容
        wbv = openpyxl.load_workbook(src, data_only=True)           # 公式取計算值
    except Exception as e:
        raise MaskError(f"無法開啟檔案（可能不是有效的 Excel 檔或正被 Excel 開啟中）：\n{e}")

    ws = wb.worksheets[0]                                           # 第 1 個工作表
    wsv = wbv.worksheets[0]

    hit = find_name_col(ws)
    if hit is None:
        headers = [f"{get_column_letter(c.column)}={c.value!r}" for c in ws[1] if c.value is not None]
        raise MaskError("第 1 個工作表的第 1 列找不到「姓名」欄位。\n"
                        f"第 1 列現有標題：{'、'.join(headers) or '(整列皆空白)'}\n"
                        "請把姓名欄的標題改為「姓名」後再執行一次。")
    col, header = hit

    # 建立遮罩工作表（值複製整個第 1 個工作表，僅姓名欄第 2 列起遮罩）
    base = ws.title[:25] + "_遮罩"
    new_name, i = base, 2
    while new_name in wb.sheetnames:
        new_name, i = f"{base}{i}", i + 1
    out = wb.create_sheet(new_name)

    masked = 0
    for row in wsv.iter_rows():
        for cell in row:
            v = cell.value
            if v is not None and cell.column == col and cell.row >= 2:
                v, ok = mask_name(v)
                if ok:
                    masked += 1
            out.cell(cell.row, cell.column).value = v

    dst = next_output_path(src)
    try:
        wb.save(dst)
    except Exception as e:
        raise MaskError(f"無法儲存新檔（請確認資料夾可寫入、檔案未被開啟）：\n{e}")

    return {
        "src": src,
        "dst": dst,
        "masked": masked,
        "sheet": new_name,
        "column": get_column_letter(col),
        "header": header,
    }


def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    if not argv:
        sys.exit("用法：python mask_names.py 檔案.xlsx（或改用 GUI：python app.py）")
    try:
        r = mask_file(argv[0])
    except MaskError as e:
        sys.exit(str(e))
    print(f"已找到姓名欄：{r['column']} 欄（標題「{r['header']}」）")
    print(f"完成：遮罩 {r['masked']} 個姓名 → 新工作表「{r['sheet']}」")
    print(f"已另存新檔：{r['dst']}")
    print("原檔案完全未更動；全程離線，資料未離開你的電腦。")
    return 0


if __name__ == "__main__":
    main()
