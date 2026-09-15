#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
姓名遮罩與學生編號 NameMasker 2.0 — GUI／CLI 入口
====================================================
把 Zuvio「下載數據」xlsx（或名冊 xlsx、通用 CSV）變成可以安全分析的檔案：
  • A 欄插入「學生編號」（115-1_EC_1、115-1_EC_2…），同一個人跨區塊同號
  • 學號欄、電子郵件欄每個字元改成 O（長度不變）
  • 姓名欄第 2 字改成 O；自由文字裡的同學姓名改成該生的學生編號
  • 新增工作表「學生編號連結姓名」（再識別鑰匙，只能留在自己電腦）
  • 另存「原檔名02.xlsx」，原檔一個位元組都不動

用法：
  雙擊 NameMasker.exe                          → 開啟視窗
  把 Excel 檔拖到 NameMasker.exe 的圖示上      → 直接處理（用上次的課程／學期設定）
  NameMasker.exe 檔案.xlsx --course EC --semester 115-1 [--no-mask-names]
  NameMasker.exe --selftest                    → 自我測試（成功印 SELFTEST OK 並 exit 0）
  NameMasker.exe --version                     → 顯示版本

環境變數 NAMEMASKER_NO_DIALOG=1：命令列模式不跳訊息框（供自動化／CI 使用）。
"""
import json
import os
import sys
import tempfile
import traceback

APP_NAME = "姓名遮罩與學生編號"
EXE_NAME = "NameMasker"
VERSION = "2.0.0"

DEFAULT_COURSE = "EC"
DEFAULT_SEMESTER = "115-1"

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass
try:
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from core_mask import (  # noqa: E402
    LINK_SHEET_HEADERS, LINK_SHEET_TITLE, MaskError, SUPPORTED_EXT,
    process_workbook, validate_course, validate_semester,
)

_LOG_LINES = []
SETTINGS_DIR = os.path.join(os.path.expanduser("~"), ".namemasker")
SETTINGS_PATH = os.path.join(SETTINGS_DIR, "settings.json")


# --------------------------------------------------------------------------
# 基礎工具
# --------------------------------------------------------------------------
def resource_path(rel):
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, rel)


def out(msg=""):
    """console=False 封裝時 stdout 可能不存在；一律記錄到 log，有 console 就順便印。"""
    _LOG_LINES.append(str(msg))
    try:
        if sys.stdout is not None:
            print(msg)
            sys.stdout.flush()
    except Exception:
        pass


def write_log(name):
    path = os.path.join(tempfile.gettempdir(), name)
    try:
        with open(path, "w", encoding="utf-8") as f:
            f.write("\n".join(_LOG_LINES) + "\n")
    except Exception:
        return None
    return path


def load_settings():
    s = {"course_code": DEFAULT_COURSE, "semester": DEFAULT_SEMESTER, "mask_names": True}
    try:
        with open(SETTINGS_PATH, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict):
            for k in s:
                if k in data:
                    s[k] = data[k]
    except Exception:
        pass
    try:
        validate_course(s["course_code"])
    except Exception:
        s["course_code"] = DEFAULT_COURSE
    try:
        validate_semester(s["semester"])
    except Exception:
        s["semester"] = DEFAULT_SEMESTER
    s["mask_names"] = bool(s["mask_names"])
    return s


def save_settings(s):
    try:
        os.makedirs(SETTINGS_DIR, exist_ok=True)
        with open(SETTINGS_PATH, "w", encoding="utf-8") as f:
            json.dump(s, f, ensure_ascii=False, indent=2)
        return True
    except Exception:
        return False


def parse_args(args):
    """回傳 (檔案清單, course, semester, mask_names)；旗標值不會被當成檔案。"""
    s = load_settings()
    course, sem, mk = s["course_code"], s["semester"], s["mask_names"]
    files, i = [], 0
    while i < len(args):
        a = args[i]
        if a in ("--course", "-c") and i + 1 < len(args):
            course = args[i + 1]
            i += 2
            continue
        if a in ("--semester", "-s") and i + 1 < len(args):
            sem = args[i + 1]
            i += 2
            continue
        if a == "--no-mask-names":
            mk = False
            i += 1
            continue
        if a == "--mask-names":
            mk = True
            i += 1
            continue
        if a.startswith("-"):
            i += 1
            continue
        files.append(a)
        i += 1
    return files, course, sem, mk


# --------------------------------------------------------------------------
# 自我測試
# --------------------------------------------------------------------------
def _fmt(template, order_map):
    """把 '{n}' 換成第 n 號學生編號。"""
    s = template
    for n, no in order_map.items():
        s = s.replace("{%d}" % n, no)
    return s


def run_selftest():  # noqa: C901 - 測試流程刻意寫得很直白
    import hashlib
    import shutil

    import openpyxl
    import sample_data as SD

    out(f"{EXE_NAME} {VERSION} — 自我測試開始（全部使用合成資料，零真實個資）")
    tmp = tempfile.mkdtemp(prefix="NameMasker_selftest_")
    SEM, CRS = "115-1", "EC"
    NO = {n: f"{SEM}_{CRS}_{n}" for n in range(1, 20)}
    try:
        # ================= [1/8] 參數驗證 ==================================
        for bad in ("E", "ECC", "E1", "", "電化"):
            try:
                validate_course(bad)
                raise AssertionError(f"課程縮寫 {bad!r} 應該被擋下來")
            except MaskError:
                pass
        assert validate_course("ec") == "EC", "課程縮寫應自動轉大寫"
        for bad in ("115", "115-3", "11-1", "1151", ""):
            try:
                validate_semester(bad)
                raise AssertionError(f"學期 {bad!r} 應該被擋下來")
            except MaskError:
                pass
        assert validate_semester(" 115-1 ") == "115-1"
        out("[1/8] 課程縮寫／學期格式驗證通過")

        # ================= [2/8] 產生 Zuvio 合成檔 =========================
        src = os.path.join(tmp, "合成Zuvio.xlsx")
        SD.write_zuvio_sample(src)
        before = hashlib.sha256(open(src, "rb").read()).hexdigest()
        out(f"[2/8] 已產生 Zuvio 格式合成檔（{len(SD.ZUVIO_ROWS)} 列、5 個表頭區塊）")

        # ================= [3/8] 執行主流程 ================================
        rep = process_workbook(src, CRS, SEM, mask_names=True)
        E = SD.ZUVIO_EXPECT
        for k, want in E.items():
            got = getattr(rep, k)
            assert got == want, f"統計 {k} 應為 {want}，實際 {got}"
        assert os.path.basename(rep.dst) == "合成Zuvio02.xlsx", \
            f"輸出檔名應為 合成Zuvio02.xlsx，實際 {os.path.basename(rep.dst)}"
        assert os.path.dirname(rep.dst) == tmp, "輸出應與原檔同資料夾"
        out(f"[3/8] 主流程統計全部符合：區塊 {rep.blocks}、學生 {rep.students}、"
            f"學號遮罩 {rep.id_masked}、姓名遮罩 {rep.name_masked}、"
            f"文字替換 {rep.text_replacements}、對照表 {rep.link_rows}")

        # 原檔位元組不得更動
        after = hashlib.sha256(open(src, "rb").read()).hexdigest()
        assert before == after, "原檔被更動了（不允許）"

        # ================= [4/8] 讀回輸出檔逐項驗證 ========================
        wb = openpyxl.load_workbook(rep.dst)
        ws = wb[SD.ZUVIO_SHEET]

        # (a) A 欄學生編號
        for row, want in SD.ZUVIO_A_COL.items():
            got = ws.cell(row, 1).value
            if isinstance(want, int):
                want_s = NO[want]
            elif want == "":
                want_s = None
            else:
                want_s = want
            assert got == want_s, f"第 {row} 列 A 欄應為 {want_s!r}，實際 {got!r}"
        for hr in SD.ZUVIO_HEADER_ROWS:
            assert ws.cell(hr, 1).value == "學生編號", f"第 {hr} 列 A 欄應是表頭「學生編號」"

        # (b) 跨區塊同人同號（1 號出現在 14/28/36/44 列）
        for row in (14, 28, 36, 44):
            assert ws.cell(row, 1).value == NO[1], f"第 {row} 列應同為 {NO[1]}"
        for row in (38, 43):
            assert ws.cell(row, 1).value == NO[7], f"第 {row} 列應同為 {NO[7]}"

        # (c) 匿名列不編號
        for row in (13, 31):
            assert ws.cell(row, 1).value == "匿名", f"第 {row} 列（匿名）不可拿到學生編號"

        # (d) 學號欄（插欄後 B 欄）全部 O 且長度不變
        id_rows = {}
        for r0, row in enumerate(SD.ZUVIO_ROWS, start=1):
            id_rows[r0] = row[0]
        checked = 0
        for row in list(SD.ZUVIO_A_COL):
            if row in SD.ZUVIO_HEADER_ROWS:
                continue
            raw = id_rows[row]
            got = ws.cell(row, 2).value
            want_len = len(str(int(raw)) if isinstance(raw, (int, float)) else str(raw))
            assert isinstance(got, str) and set(got) == {"O"} and len(got) == want_len, \
                f"第 {row} 列學號應為 {want_len} 個 O，實際 {got!r}"
            checked += 1
        assert checked == E["id_masked"], f"學號遮罩檢查列數 {checked} ≠ {E['id_masked']}"

        # (e) 姓名欄（插欄後 C 欄）遮罩
        for row, want in SD.ZUVIO_MASKED_NAMES.items():
            got = ws.cell(row, 3).value
            assert got == want, f"第 {row} 列姓名應為 {want!r}，實際 {got!r}"

        # (f) 文字儲存格內容
        for (row, col), tpl in SD.ZUVIO_TEXT.items():
            want = _fmt(tpl, NO)
            got = ws.cell(row, col).value
            assert got == want, f"第 {row} 列第 {col} 欄應為 {want!r}，實際 {got!r}"

        # (g) 全表不得殘留任何合成姓名（姓名欄已遮罩、連結表除外）
        names = [rec[1] for rec in SD.ZUVIO_ORDER]
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if not isinstance(v, str):
                    continue
                for nm in names:
                    assert nm not in v, f"第 {r} 列第 {c} 欄仍殘留姓名（{len(nm)} 字）"

        # (h) 連結表
        assert LINK_SHEET_TITLE in wb.sheetnames, f"找不到工作表「{LINK_SHEET_TITLE}」"
        ls = wb[LINK_SHEET_TITLE]
        for c, h in enumerate(LINK_SHEET_HEADERS, start=1):
            assert ls.cell(1, c).value == h, f"連結表第 {c} 欄標題應為 {h!r}"
        assert ls.max_row == len(SD.ZUVIO_ORDER) + 1, \
            f"連結表應有 {len(SD.ZUVIO_ORDER)} 筆，實際 {ls.max_row - 1}"
        for i, (n, name, sid, first_row) in enumerate(SD.ZUVIO_ORDER, start=2):
            assert ls.cell(i, 1).value == NO[n], f"連結表第 {i} 列編號應為 {NO[n]}"
            assert ls.cell(i, 2).value == name, f"連結表第 {i} 列姓名不符"
            assert str(ls.cell(i, 3).value) == sid, f"連結表第 {i} 列學號不符"
            assert ls.cell(i, 5).value == first_row, \
                f"連結表第 {i} 列首次出現列應為 {first_row}，實際 {ls.cell(i, 5).value}"
        wb.close()
        out("[4/8] 輸出檔逐項驗證通過（A 欄編號、跨區塊同號、匿名不編號、"
            "學號全 O、姓名遮罩、文字替換、無姓名殘留、連結表）")

        # ================= [5/8] 02 → 03 遞增、--no-mask-names ==============
        rep2 = process_workbook(src, CRS, SEM, mask_names=True)
        assert os.path.basename(rep2.dst) == "合成Zuvio03.xlsx", \
            f"第二次輸出應為 合成Zuvio03.xlsx，實際 {os.path.basename(rep2.dst)}"

        src2 = os.path.join(tmp, "不遮姓名.xlsx")
        SD.write_zuvio_sample(src2)
        rep3 = process_workbook(src2, CRS, SEM, mask_names=False)
        assert rep3.name_masked == 0, "取消勾選時不應遮罩姓名欄"
        assert rep3.id_masked == E["id_masked"], "取消姓名遮罩時，學號仍必須遮罩"
        assert rep3.text_replacements == E["text_replacements"], \
            "取消姓名遮罩時，文字內姓名仍必須換成學生編號"
        wb3 = openpyxl.load_workbook(rep3.dst)
        ws3 = wb3[SD.ZUVIO_SHEET]
        assert ws3.cell(14, 3).value == "王小明", "取消勾選時姓名欄應維持原樣"
        wb3.close()
        out("[5/8] 輸出編號遞增（02→03）與「不遮姓名欄」選項驗證通過")

        # ================= [6/8] 名冊型（含電子郵件欄、None 表頭） ==========
        rsrc = os.path.join(tmp, "合成名冊.xlsx")
        SD.write_roster_sample(rsrc)
        rrep = process_workbook(rsrc, CRS, SEM, mask_names=True)
        for k, want in SD.ROSTER_EXPECT.items():
            got = getattr(rrep, k)
            assert got == want, f"名冊統計 {k} 應為 {want}，實際 {got}"
        wbr = openpyxl.load_workbook(rrep.dst)
        wsr = wbr[SD.ROSTER_SHEET]
        for i, row in enumerate(SD.ROSTER_ROWS[1:], start=2):
            assert wsr.cell(i, 1).value == NO[i - 1], f"名冊第 {i} 列應有學生編號"
            eml = wsr.cell(i, 5).value           # 電子郵件欄（原 4 → 插欄後 5）
            assert isinstance(eml, str) and set(eml) == {"O"} and len(eml) == len(row[3]), \
                f"名冊第 {i} 列電子郵件應為 {len(row[3])} 個 O"
            assert wsr.cell(i, 6).value == row[4], "非郵件的 15 字說明欄不可被動到"
        wbr.close()
        out("[6/8] 名冊型（第 1 列表頭、None 表頭、電子郵件欄）驗證通過："
            f"郵件遮罩 {rrep.email_masked} 筆，長度不變")

        # ================= [7/8] 通用 CSV ==================================
        csrc = os.path.join(tmp, "合成通用.csv")
        SD.write_csv_sample(csrc)
        crep = process_workbook(csrc, CRS, SEM, mask_names=True)
        for k, want in SD.CSV_EXPECT.items():
            got = getattr(crep, k)
            assert got == want, f"CSV 統計 {k} 應為 {want}，實際 {got}"
        assert crep.dst.endswith("02.xlsx"), "CSV 應輸出成 xlsx"
        assert crep.email_masked == 0 and any("電子郵件" in n for n in crep.notes), \
            "沒有電子郵件欄時應在說明中註記"
        out(f"[7/8] 通用 CSV → xlsx 驗證通過（{crep.students} 位學生）")

        # ================= [8/8] 錯誤處理 ==================================
        bad = os.path.join(tmp, "沒有姓名欄.xlsx")
        wbb = openpyxl.Workbook()
        wbb.active.append(["座號", "分數"])
        wbb.active.append([1, 100])
        wbb.save(bad)
        wbb.close()
        try:
            process_workbook(bad, CRS, SEM)
            raise AssertionError("缺「姓名」欄時應丟出 MaskError")
        except MaskError as e:
            assert "姓名" in str(e)
        try:
            process_workbook(src, "E1", SEM)
            raise AssertionError("錯誤的課程縮寫應丟出 MaskError")
        except MaskError:
            pass
        try:
            process_workbook(os.path.join(tmp, "不存在.xlsx"), CRS, SEM)
            raise AssertionError("檔案不存在應丟出 MaskError")
        except MaskError:
            pass
        out("[8/8] 錯誤處理（無姓名欄、課程縮寫格式、檔案不存在）驗證通過")

        out("SELFTEST OK")
        return 0
    except Exception:
        out("SELFTEST FAILED")
        out(traceback.format_exc())
        return 1
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# --------------------------------------------------------------------------
# 命令列（含拖放到 exe 圖示）模式
# --------------------------------------------------------------------------
def run_cli(paths, course, sem, mask_names):
    results, errors = [], []
    try:
        course = validate_course(course)
        sem = validate_semester(sem)
    except MaskError as e:
        out(str(e))
        if os.environ.get("NAMEMASKER_NO_DIALOG") != "1":
            _msgbox(f"{APP_NAME}　參數錯誤", str(e), warn=True)
        return 1

    for p in paths:
        try:
            rep = process_workbook(p, course, sem, mask_names)
            results.append(rep)
            out(f"完成：{os.path.basename(rep.src)}")
            for line in rep.summary_lines():
                out("    " + line)
        except MaskError as e:
            errors.append((p, str(e)))
            out(f"失敗：{os.path.basename(p)}\n{e}")
        except Exception as e:  # noqa: BLE001
            errors.append((p, f"未預期的錯誤：{e}"))
            out(f"失敗：{os.path.basename(p)}\n{traceback.format_exc()}")

    lines = []
    for rep in results:
        lines.append(f"✔ {os.path.basename(rep.src)}")
        lines.append(f"    區塊 {rep.blocks} 個、編號學生 {rep.students} 位、"
                     f"遮罩學號 {rep.id_masked} 筆、文字替換 {rep.text_replacements} 處")
        lines.append(f"    → {os.path.basename(rep.dst)}")
    for p, msg in errors:
        lines.append(f"✘ {os.path.basename(p)}\n    {msg}")
    lines.append("")
    lines.append(f"設定：學期 {sem}、課程 {course}、"
                 f"{'含' if mask_names else '不含'}姓名欄遮罩")
    lines.append(f"※ 工作表「{LINK_SHEET_TITLE}」是再識別鑰匙，只能留在自己的電腦，不要上傳。")
    lines.append("原檔案完全未更動；全程離線，資料未離開你的電腦。")
    body = "\n".join(lines)

    if os.environ.get("NAMEMASKER_NO_DIALOG") == "1":
        return 1 if errors else 0
    _msgbox(f"{APP_NAME}　處理結果", body, warn=bool(errors))
    return 1 if errors else 0


def _msgbox(title, body, warn=False):
    try:
        import tkinter as tk
        from tkinter import messagebox
        root = tk.Tk()
        root.withdraw()
        (messagebox.showwarning if warn else messagebox.showinfo)(title, body)
        root.destroy()
    except Exception:
        pass


# --------------------------------------------------------------------------
# GUI
# --------------------------------------------------------------------------
def pick_font():
    import tkinter.font as tkfont
    try:
        fams = set(tkfont.families())
    except Exception:
        return "TkDefaultFont"
    for f in ("Microsoft JhengHei UI", "微軟正黑體", "Microsoft JhengHei",
              "PingFang TC", "Heiti TC", "Noto Sans CJK TC"):
        if f in fams:
            return f
    return "TkDefaultFont"


def run_gui(preset_files=None, preset=None):  # noqa: C901
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    cfg = preset or load_settings()

    root = tk.Tk()
    root.title(f"{APP_NAME}　{EXE_NAME} v{VERSION}")
    root.geometry("820x640")
    root.minsize(720, 560)
    try:
        ico = resource_path("icon.ico")
        if os.path.exists(ico):
            root.iconbitmap(ico)
    except Exception:
        pass

    fam = pick_font()
    F = (fam, 11)
    FB = (fam, 11, "bold")
    FT = (fam, 16, "bold")
    FS = (fam, 9)

    files = list(preset_files or [])
    last_out_dir = {"path": None}

    tk.Label(root, text=APP_NAME, font=FT, anchor="w").pack(fill="x", padx=14, pady=(12, 0))
    tk.Label(root,
             text="把 Zuvio 下載數據（或名冊 xlsx／CSV）轉成可以安全分析的檔案："
                  "A 欄加「學生編號」、學號與電子郵件改成 OOOO、姓名第 2 字改 O、"
                  "自由文字裡的同學姓名改成學生編號，另存「原檔名02.xlsx」。原檔不會被更動。",
             font=F, anchor="w", justify="left", wraplength=780, fg="#333333"
             ).pack(fill="x", padx=14, pady=4)
    tk.Label(root, text="🔒 完全離線：只讀寫你電腦上的檔案，不連網、不上傳任何資料。",
             font=FB, anchor="w", fg="#1b6b3a").pack(fill="x", padx=14, pady=(0, 6))

    # ---- 設定列 ----
    setbox = tk.LabelFrame(root, text=" 這批檔案的設定 ", font=FB, fg="#204a87")
    setbox.pack(fill="x", padx=14, pady=4)
    inner = tk.Frame(setbox)
    inner.pack(fill="x", padx=10, pady=8)

    tk.Label(inner, text="學期", font=FB).pack(side="left")
    v_sem = tk.StringVar(value=cfg["semester"])
    tk.Entry(inner, textvariable=v_sem, font=F, width=9).pack(side="left", padx=(4, 2))
    tk.Label(inner, text="（例 115-1）", font=FS, fg="#777777").pack(side="left", padx=(0, 14))

    tk.Label(inner, text="課程縮寫", font=FB).pack(side="left")
    v_crs = tk.StringVar(value=cfg["course_code"])
    e_crs = tk.Entry(inner, textvariable=v_crs, font=F, width=6)
    e_crs.pack(side="left", padx=(4, 2))
    tk.Label(inner, text="（2 個英文字母，例 EC 環化）", font=FS, fg="#777777"
             ).pack(side="left", padx=(0, 14))

    def upper_course(*_):
        s = v_crs.get()
        if s != s.upper():
            v_crs.set(s.upper())
    v_crs.trace_add("write", upper_course)

    v_mask = tk.BooleanVar(value=cfg["mask_names"])
    tk.Checkbutton(inner, text="同時遮罩姓名欄與文字內姓名（建議勾選）",
                   variable=v_mask, font=F).pack(side="left")

    lbl_preview = tk.Label(setbox, text="", font=FS, fg="#555555", anchor="w")
    lbl_preview.pack(fill="x", padx=12, pady=(0, 8))

    def refresh_preview(*_):
        s, c = v_sem.get().strip(), v_crs.get().strip()
        lbl_preview.config(text=f"學生編號會長成：{s}_{c}_1、{s}_{c}_2 …　"
                                "（同一位同學不管出現在哪個區塊，都是同一個編號）")
    v_sem.trace_add("write", refresh_preview)
    v_crs.trace_add("write", refresh_preview)
    refresh_preview()

    # ---- 檔案清單 ----
    box = tk.LabelFrame(root, text=" 待處理的檔案（.xlsx / .xlsm / .csv，可多選） ",
                        font=FB, fg="#204a87")
    box.pack(fill="both", expand=False, padx=14, pady=4)
    lb = tk.Listbox(box, font=F, height=6, activestyle="none")
    sb = tk.Scrollbar(box, command=lb.yview)
    lb.config(yscrollcommand=sb.set)
    lb.pack(side="left", fill="both", expand=True, padx=(8, 0), pady=8)
    sb.pack(side="right", fill="y", padx=(0, 8), pady=8)

    def refresh():
        lb.delete(0, "end")
        for p in files:
            lb.insert("end", p)
        btn_run.config(state=("normal" if files else "disabled"))

    def add_files():
        picked = filedialog.askopenfilenames(
            title="選擇檔案（可按住 Ctrl 多選）",
            filetypes=[("Excel／CSV", "*.xlsx *.xlsm *.csv"), ("所有檔案", "*.*")])
        for p in picked:
            if p not in files:
                files.append(p)
        refresh()

    def clear_files():
        files.clear()
        refresh()

    def log(msg=""):
        txt.config(state="normal")
        txt.insert("end", msg + "\n")
        txt.see("end")
        txt.config(state="disabled")
        root.update_idletasks()

    def do_run():
        if not files:
            return
        try:
            course = validate_course(v_crs.get())
            sem = validate_semester(v_sem.get())
        except MaskError as e:
            messagebox.showwarning(APP_NAME, str(e))
            return
        mk = bool(v_mask.get())
        save_settings({"course_code": course, "semester": sem, "mask_names": mk})

        btn_run.config(state="disabled")
        txt.config(state="normal")
        txt.delete("1.0", "end")
        txt.config(state="disabled")
        log(f"設定：學期 {sem}、課程縮寫 {course}、"
            f"{'含' if mk else '不含'}姓名遮罩　→ 學生編號格式 {sem}_{course}_1")
        log("")
        ok = fail = 0
        for p in list(files):
            try:
                rep = process_workbook(p, course, sem, mk)
                ok += 1
                last_out_dir["path"] = os.path.dirname(rep.dst)
                log(f"✔ {os.path.basename(rep.src)}")
                for line in rep.summary_lines():
                    log("    " + line)
            except MaskError as e:
                fail += 1
                log(f"✘ {os.path.basename(p)}")
                for line in str(e).splitlines():
                    log("    " + line)
            except Exception as e:  # noqa: BLE001
                fail += 1
                log(f"✘ {os.path.basename(p)}　未預期的錯誤：{e}")
            log("")
        log(f"── 全部完成：成功 {ok} 個、失敗 {fail} 個。原檔案完全未更動。")
        log(f"※ 工作表「{LINK_SHEET_TITLE}」可以把編號換回真名，是再識別鑰匙，")
        log("   只能留在自己的電腦，不要上傳雲端、不要寄給別人。")
        btn_run.config(state="normal")
        btn_open.config(state=("normal" if last_out_dir["path"] else "disabled"))
        if fail == 0:
            messagebox.showinfo(APP_NAME, f"已完成 {ok} 個檔案。\n"
                                          "輸出檔就放在原檔案的同一個資料夾裡。")
        else:
            messagebox.showwarning(APP_NAME, f"完成 {ok} 個，{fail} 個未能處理。\n"
                                             "請看下方訊息區的說明。")

    def open_folder():
        d = last_out_dir["path"]
        if not d or not os.path.isdir(d):
            return
        try:
            if sys.platform.startswith("win"):
                os.startfile(d)  # noqa: S606
            elif sys.platform == "darwin":
                import subprocess
                subprocess.Popen(["open", d])
            else:
                import subprocess
                subprocess.Popen(["xdg-open", d])
        except Exception as e:  # noqa: BLE001
            messagebox.showerror(APP_NAME, f"無法開啟資料夾：{e}")

    def about():
        messagebox.showinfo(
            f"關於 {APP_NAME}",
            f"{APP_NAME}（{EXE_NAME}）v{VERSION}\n\n"
            "1. 掃描整張工作表的所有「姓名」表頭，切出多個資料區塊。\n"
            "2. A 欄插入「學生編號」（學期_課程_流水號），同一個人跨區塊同號。\n"
            "3. 學號欄、電子郵件欄每個字元改成 O（長度不變）。\n"
            "4. 姓名欄第 2 字改成 O；文字裡的同學姓名改成學生編號。\n"
            f"5. 新增工作表「{LINK_SHEET_TITLE}」可換回真名 → 只能留本機。\n"
            "6. 另存「原檔名02.xlsx」，原檔完全不動。\n\n"
            "隱私：完全離線，不連網、不上傳任何資料。")

    bar = tk.Frame(root)
    bar.pack(fill="x", padx=14, pady=(2, 6))
    tk.Button(bar, text="選擇檔案（可多選）", font=FB, command=add_files,
              width=18).pack(side="left")
    tk.Button(bar, text="清除清單", font=F, command=clear_files, width=9).pack(side="left", padx=6)
    btn_run = tk.Button(bar, text="開始處理", font=FB, command=do_run, width=12,
                        state="disabled", bg="#204a87", fg="white",
                        activebackground="#16345f", activeforeground="white")
    btn_run.pack(side="left", padx=6)
    btn_open = tk.Button(bar, text="開啟輸出資料夾", font=F, command=open_folder,
                         width=14, state="disabled")
    btn_open.pack(side="left", padx=6)
    tk.Button(bar, text="關於", font=F, command=about, width=6).pack(side="right")

    outbox = tk.LabelFrame(root, text=" 處理結果 ", font=FB, fg="#204a87")
    outbox.pack(fill="both", expand=True, padx=14, pady=(0, 6))
    txt = tk.Text(outbox, font=(fam, 10), wrap="word", state="disabled",
                  bg="#fbfbfb", relief="flat")
    sb2 = ttk.Scrollbar(outbox, command=txt.yview)
    txt.config(yscrollcommand=sb2.set)
    txt.pack(side="left", fill="both", expand=True, padx=(8, 0), pady=8)
    sb2.pack(side="right", fill="y", padx=(0, 8), pady=8)

    tk.Label(root, text=f"{EXE_NAME} v{VERSION}　｜　也可以把檔案直接拖到本程式的圖示上執行"
                        "　｜　東華大學自然資源與環境學系　仿生與環境工作坊",
             font=FS, fg="#777777", anchor="w").pack(fill="x", padx=14, pady=(0, 10))

    refresh()
    root.mainloop()
    return 0


# --------------------------------------------------------------------------
def main():
    args = sys.argv[1:]
    if "--version" in args or "-v" in args:
        out(f"{EXE_NAME} {VERSION}")
        write_log(f"{EXE_NAME}_version.log")
        return 0
    if "--help" in args or "-h" in args:
        out(__doc__)
        return 0
    if "--selftest" in args:
        code = run_selftest()
        p = os.path.join(tempfile.gettempdir(), f"{EXE_NAME}_selftest.log")
        out(f"（測試紀錄：{p}）")
        write_log(f"{EXE_NAME}_selftest.log")
        return code

    files, course, sem, mk = parse_args(args)
    files = [f for f in files if f.lower().endswith(SUPPORTED_EXT) and os.path.isfile(f)]
    if files:
        code = run_cli(files, course, sem, mk)
        write_log(f"{EXE_NAME}_cli.log")
        return code
    return run_gui(preset_files=[], preset={"course_code": course, "semester": sem,
                                            "mask_names": mk})


if __name__ == "__main__":
    sys.exit(main())
