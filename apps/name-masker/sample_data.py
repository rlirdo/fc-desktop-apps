# -*- coding: utf-8 -*-
"""合成測試資料（零真實個資，全部為虛構人名與虛構學號）。

供 `app.py --selftest` 使用，也用來產生 `selftest\\` 資料夾裡的示範檔。
資料格式仿照 Zuvio「下載數據」xlsx（多個表頭區塊、匿名作答者、未作答名單、
測驗題目型答案欄、自由文字提到同學姓名），以及名冊型 xlsx（含電子郵件欄、
含 None 表頭）與通用 CSV。
"""
import csv

import openpyxl

BLANK = None

# --------------------------------------------------------------------------
# 1) Zuvio 題組問答＋測驗題目 合成檔（5 個表頭區塊）
# --------------------------------------------------------------------------
ZUVIO_SHEET = "sheet1"

# 每一列就是實際要寫進去的內容；None = 空白儲存格
ZUVIO_ROWS = [
    ["資料夾名稱", "合成測試資料夾", BLANK, BLANK, BLANK, BLANK],                 # R1
    ["問題題型", "題組問答", BLANK, BLANK, BLANK, BLANK],                         # R2
    ["是否分組", "否", BLANK, BLANK, BLANK, BLANK],                               # R3
    ["是否匿名", "否", BLANK, BLANK, BLANK, BLANK],                               # R4
    ["題幹", "題組問答_合成資料_不含真實個資", BLANK, BLANK, BLANK, BLANK],        # R5
    [BLANK, BLANK, BLANK, BLANK, BLANK, BLANK],                                   # R6
    [BLANK, BLANK, BLANK, BLANK, BLANK, BLANK],                                   # R7
    ["問題類型", BLANK, BLANK, "第1題:問答題", "第2題:單選題", BLANK],             # R8
    ["問題敘述", BLANK, BLANK, "請描述一個與課程相關的生活例子。", "請選一個選項。", BLANK],  # R9
    ["正解", BLANK, BLANK, "無正解", "無正解", BLANK],                             # R10
    [BLANK, BLANK, BLANK, BLANK, BLANK, BLANK],                                   # R11
    ["學號", "姓名", "作答時間", "回答", "回答", BLANK],                           # R12  ← 區塊 1
    ["匿名作答者", "匿名作答者", "2026-09-09 15:00:00", "這是匿名的回答", "(2)乙", BLANK],   # R13
    [990054043, "王小明", "2026-09-07 16:00:00", "我和林大同一起討論了這題", "(1)甲", BLANK],  # R14  int 學號
    [990054044, "林大同", "2026-09-08 10:00:00", "參考了陳曉華學姊的報告", "(3)丙", BLANK],    # R15  int 學號
    ["990054045", "陳曉華", "2026-09-08 11:00:00", "自己查資料完成", "(2)乙", BLANK],          # R16  str 學號
    [990054046, "歐陽小花", "2026-09-09 09:00:00", "和王小明討論過", "(1)甲", BLANK],          # R17
    [990054047, "李明", "2026-09-10 08:00:00", "沒有跟別人討論", "(2)乙", BLANK],              # R18
    [BLANK, BLANK, BLANK, BLANK, BLANK, BLANK],                                   # R19
    ["未作答學生", BLANK, BLANK, BLANK, BLANK, BLANK],                            # R20
    ["學號", "姓名", BLANK, BLANK, BLANK, BLANK],                                 # R21  ← 區塊 2
    [990054048, "張三豐", BLANK, BLANK, BLANK, BLANK],                            # R22
    [BLANK, BLANK, BLANK, BLANK, BLANK, BLANK],                                   # R23
    [BLANK, BLANK, BLANK, BLANK, BLANK, BLANK],                                   # R24
    ["第1題:問答題", BLANK, BLANK, BLANK, BLANK, BLANK],                          # R25
    ["問題敘述", "請描述一個與課程相關的生活例子。", BLANK, BLANK, BLANK, BLANK],   # R26
    ["學號", "姓名", "回答", BLANK, BLANK, BLANK],                                # R27  ← 區塊 3
    [990054043, "王小明", "我和林大同一起討論了這題", BLANK, BLANK, BLANK],        # R28
    [990054044, "林大同", "參考了陳曉華學姊的報告", BLANK, BLANK, BLANK],          # R29
    [990054046, "歐陽小花", "和王小明討論過", BLANK, BLANK, BLANK],                # R30
    ["匿名作答者", "匿名作答者", "這是匿名的回答", BLANK, BLANK, BLANK],           # R31
    [BLANK, BLANK, BLANK, BLANK, BLANK, BLANK],                                   # R32
    [BLANK, BLANK, BLANK, BLANK, BLANK, BLANK],                                   # R33
    ["測驗題目", "合成小考", BLANK, BLANK, BLANK, BLANK],                         # R34
    ["學號", "姓名", "作答時間", "總分", "第1題", "第2題"],                        # R35  ← 區塊 4
    [990054043, "王小明", "2026-09-12 09:00:00", 80, "答對", "答錯"],              # R36
    ["990054045", "陳曉華", "2026-09-12 09:05:00", 100, "答對", "答對"],           # R37
    [990054049, "林明", "2026-09-12 09:10:00", 60, "答錯", "答對"],                # R38  2 字姓名
    [BLANK, BLANK, BLANK, BLANK, BLANK, BLANK],                                   # R39
    [BLANK, BLANK, BLANK, BLANK, BLANK, BLANK],                                   # R40
    ["分組討論彙總", BLANK, BLANK, BLANK, BLANK, BLANK],                          # R41
    ["學號", "姓名", "回答", "同組成員", BLANK, BLANK],                            # R42  ← 區塊 5
    [990054049, "林明", "討論主題是水質檢測", "李明", BLANK, BLANK],               # R43
    [990054043, "王小明", "我和歐陽小花同組", "林明", BLANK, BLANK],               # R44
]

# ---- 期望值（selftest 逐項比對用） ---------------------------------------
ZUVIO_EXPECT = {
    "blocks": 5,
    "students": 7,
    "anon_rows": 2,
    "id_masked": 16,
    "email_masked": 0,
    "name_masked": 14,
    "text_replacements": 9,
    "link_rows": 7,
}

# 學生編號（依首次出現順序）→ 對應的合成姓名／首次出現列
ZUVIO_ORDER = [
    (1, "王小明", "990054043", 14),
    (2, "林大同", "990054044", 15),
    (3, "陳曉華", "990054045", 16),
    (4, "歐陽小花", "990054046", 17),
    (5, "李明", "990054047", 18),
    (6, "張三豐", "990054048", 22),
    (7, "林明", "990054049", 38),
]

# 表頭列（原始列號，插欄後列號不變）
ZUVIO_HEADER_ROWS = [12, 21, 27, 35, 42]

# (列, 期望的 A 欄值)：n → 第 n 號；"匿名"；"" 代表空白
ZUVIO_A_COL = {
    12: "學生編號", 13: "匿名", 14: 1, 15: 2, 16: 3, 17: 4, 18: 5,
    21: "學生編號", 22: 6,
    27: "學生編號", 28: 1, 29: 2, 30: 4, 31: "匿名",
    35: "學生編號", 36: 1, 37: 3, 38: 7,
    42: "學生編號", 43: 7, 44: 1,
}

# 姓名遮罩後（B 欄以後皆右移 1 欄，姓名欄 = 原 2 → 3）
ZUVIO_MASKED_NAMES = {
    14: "王O明", 15: "林O同", 16: "陳O華", 17: "歐O小花", 18: "李O",
    22: "張O豐", 28: "王O明", 29: "林O同", 30: "歐O小花",
    36: "王O明", 37: "陳O華", 38: "林O", 43: "林O", 44: "王O明",
    13: "匿名作答者", 31: "匿名作答者",     # 匿名列不遮罩
}

# 文字替換後的期望內容（列, 插欄後的欄索引）→ 期望字串（{n} 代表第 n 號學生編號）
ZUVIO_TEXT = {
    (14, 5): "我和{2}一起討論了這題",
    (15, 5): "參考了{3}學姊的報告",
    (17, 5): "和{1}討論過",
    (28, 4): "我和{2}一起討論了這題",
    (29, 4): "參考了{3}學姊的報告",
    (30, 4): "和{1}討論過",
    (43, 4): "討論主題是水質檢測",      # 不含姓名，不可被動到
    (43, 5): "{5}",                     # 同組成員「李明」整段中文＝姓名 → 換編號
    (44, 4): "我和{4}同組",
    (44, 5): "{7}",
    (36, 6): "答對",                    # 測驗答案欄不可被當成姓名
    (36, 7): "答錯",
    (38, 6): "答錯",
    (13, 5): "這是匿名的回答",
}


def write_zuvio_sample(path):
    """產生 Zuvio 格式合成檔，回傳路徑。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ZUVIO_SHEET
    for row in ZUVIO_ROWS:
        ws.append(list(row))
    wb.save(path)
    wb.close()
    return path


# --------------------------------------------------------------------------
# 2) 名冊型合成檔（第 1 列就是表頭、含電子郵件欄、含 None 表頭）
# --------------------------------------------------------------------------
ROSTER_SHEET = "名冊"
ROSTER_ROWS = [
    [BLANK, "學號", "姓名", "電子郵件", BLANK],                                    # R1（第 1、5 欄表頭是 None）
    [BLANK, 990054043, "王小明", "wang.ming@gms.ndhu.edu.tw", "環境化學修課生甲班A"],
    [BLANK, 990054044, "林大同", "lin.datong@gms.ndhu.edu.tw", "環境化學修課生乙班B"],
    [BLANK, "990054045", "陳曉華", "chen.hsiaohua@gms.ndhu.edu.tw", "環境化學修課生丙班C"],
]
ROSTER_EXPECT = {
    "blocks": 1,
    "students": 3,
    "id_masked": 3,
    "email_masked": 3,
    "name_masked": 3,
    "text_replacements": 0,
    "link_rows": 3,
}


def write_roster_sample(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ROSTER_SHEET
    for row in ROSTER_ROWS:
        ws.append(list(row))
    wb.save(path)
    wb.close()
    return path


# --------------------------------------------------------------------------
# 3) 通用 CSV
# --------------------------------------------------------------------------
CSV_ROWS = [
    ["姓名", "學號", "作答"],
    ["王小明", "990054043", "我和林大同一起做實驗"],
    ["林大同", "990054044", "自己完成"],
]
CSV_EXPECT = {"blocks": 1, "students": 2, "id_masked": 2, "text_replacements": 1, "link_rows": 2}


def write_csv_sample(path):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows(CSV_ROWS)
    return path


# --------------------------------------------------------------------------
# 4) 2.1：原始名單（固定學生編號）相關合成資料
# --------------------------------------------------------------------------
# 合成班級（5 人）；學號刻意與「列出順序」不同，用來驗證「依學號排序」規則
CLASS_STUDENTS = [
    # (序號, 學號, 姓名, 班級)
    (1, "990054001", "王小明", "自資系大三"),
    (2, "990054002", "林大同", "自資系大三"),
    (3, "990054003", "陳曉華", "資工系大二"),
    (4, "990054004", "歐陽小花", "自資系大二"),
    (5, "990054005", "李明", "自資系大三"),
]
OUTSIDER = ("990054099", "張三豐")        # 不在名單的作答者（退選等）

# ---- (4a) 含「序號」欄的 Excel 名單 --------------------------------------
ROSTER_SEQ_SHEET = "名單"
ROSTER_SEQ_ROWS = [["序號", "學號", "姓名", "班級"]] + [
    [seq, int(sid), name, klass] for seq, sid, name, klass in CLASS_STUDENTS]


def write_roster_seq_sample(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ROSTER_SEQ_SHEET
    for row in ROSTER_SEQ_ROWS:
        ws.append(list(row))
    wb.save(path)
    wb.close()
    return path


# ---- (4b) 沒有「序號」欄的 Excel 名單（列出順序與學號順序不同） -----------
# 表頭也故意含 None（e學苑匯出常見），欄位順序不拘
ROSTER_NOSEQ_ROWS = [
    [BLANK, "學號", "班級", "姓名", "電子郵件"],
    [BLANK, "990054003", "資工系大二", "陳曉華", "s003@gms.ndhu.edu.tw"],
    [BLANK, "990054001", "自資系大三", "王小明", "s001@gms.ndhu.edu.tw"],
    [BLANK, "990054005", "自資系大三", "李明", "s005@gms.ndhu.edu.tw"],
    [BLANK, "990054004", "自資系大二", "歐陽小花", "s004@gms.ndhu.edu.tw"],
    [BLANK, "990054002", "自資系大三", "林大同", "s002@gms.ndhu.edu.tw"],
]


def write_roster_noseq_sample(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "e學苑匯出"
    for row in ROSTER_NOSEQ_ROWS:
        ws.append(list(row))
    wb.save(path)
    wb.close()
    return path


# ---- (4c) 名單 CSV -------------------------------------------------------
ROSTER_CSV_ROWS = [["序號", "學號", "姓名"]] + [
    [seq, sid, name] for seq, sid, name, _ in CLASS_STUDENTS]


def write_roster_csv_sample(path):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows(ROSTER_CSV_ROWS)
    return path


# ---- (4d) 東華「選課名單」PDF 抽字後的行列表 -----------------------------
# pypdf 版面：一位學生一行 → 姓名 + 9 碼學號 + 序號 + 空白 + 班級
PDF_LINES_ONELINE = [
    "國立東華大學115學年度第1學期 選課名單，合成測試",
    "NRES99999合成課程科目: 合成課程:",
    "學    分: 3/3",
    "任課教師: /測試老師",
    "班    級 姓      名 1 2 3 4 5 6 7 8 9 10 出席 平時 期中 期末 總分",
    "未列姓名，請通知學生",
    "本名單僅供合成測試",
    "正。千萬勿自行加填，",
] + [
    "%s%s%d %s" % (name, sid, seq, klass) for seq, sid, name, klass in CLASS_STUDENTS
] + [
    "註*:前一學期GPA平均未達2.0者。 共1頁,第1頁2026/9/15 列印 12:59:02",
]
# 第 3 位學生加上【註*】標記（實際名單會有）
PDF_LINES_ONELINE[8 + 2] = "陳曉華 【註*】9900540033 資工系大二"

# PyMuPDF 版面：姓名／學號／序號／班級各自一行（行序容忍度測試）
PDF_LINES_SPLIT = ["班    級 姓      名", "國立東華大學115學年度第1學期"]
for _seq, _sid, _name, _klass in CLASS_STUDENTS:
    PDF_LINES_SPLIT += [_name, "【註*】" if _seq == 3 else "", _sid, str(_seq), _klass]
PDF_LINES_SPLIT.append("註*:前一學期GPA平均未達2.0者。")

# 序號不連續（壞檔）→ 應丟出白話錯誤
PDF_LINES_BAD = [
    "班    級 姓      名",
    "王小明9900540011 自資系大三",
    "林大同9900540023 自資系大三",       # 序號跳號
    "陳曉華9900540037 資工系大二",
]

PDF_EXPECT = [(seq, sid, name, klass) for seq, sid, name, klass in CLASS_STUDENTS]
PDF_SAMPLE_NAME = "樣本_選課名單_合成.pdf"


# ---- (4d-2) 2.2：「黏連格式」選課名單（姓名／學號／序號中間不留空白） ------
# 涵蓋實際檔案看得到的每一種列形狀（中＝中文字、9＝數字、A＝英文字母）：
#   中中中99999999999      姓名＋9 碼學號＋2 碼序號 黏在一起
#   中中中9999999999       序號 1 碼
#   中中中 【註*】999…      註記夾在姓名與學號之間
#   中中中 中中999…        姓名含空白
#   中中中99999A9999       學號第 6 碼是英文字母
#   中中中A9999999999      學號＝1 字母＋8 數字（要靠「序號＝上一筆＋1」判斷）
#   99999999999           這一列沒有姓名（文字層抽不到，罕用字）
#   中中中 Aaaaa / Aaaaaa999…   姓名折行
#   中中中 中中中【註 / *】999…  註記被換行切斷
GLUED_STUDENTS = [
    # (序號, 學號, 姓名, 班級)
    (1, "990001001", "王小明", "自資系大三"),
    (2, "990001002", "林大同", "自資系大三"),
    (3, "990001003", "王 小美", "資工系大二"),
    (4, "99000A004", "李大華", "自資系大二"),
    (5, "A99000005", "張美玲", "自資系大三"),
    (6, "990001006", "", "自資系大二"),                  # 文字層讀不到姓名
    (7, "990001007", "趙小強 Jason Chao", "資工系大三"),  # 姓名折行
    (8, "990001008", "孫 大偉", "自資系大四"),
    (9, "99001A009", "周小芳", "環科系碩一"),
    (10, "99001A010", "吳志豪", "自資系大二"),
    (11, "A99000011", "鄭雅婷", "環科系碩二"),
    (12, "990001012", "許文彥", "自資系大三"),
]

_GLUED_HEAD_1 = [
    "國立東華大學115學年度第1學期 選課名單（合成測試，非真實名單）",
    "NRES99001合成課程科目名稱: 科目代號:",
    "學    分: 3/3",
    "任課教師: /測試老師",
    "班    級: 合成班",
    "上課時間: 星期三 3 4 5",
    "上課教室: 合成大樓A101",
    "頁    次: 1/2",
    "班    級 姓      名 序號 學    號 1 2 3 4 5 6 7 8 9 10 出席 平時 期中 期末 總分",
    "未列姓名者請通知教務處，本名單僅供合成測試之用。",
]
_GLUED_HEAD_2 = [
    "國立東華大學115學年度第1學期 選課名單（合成測試，非真實名單）",
    "NRES99001合成課程科目名稱: 科目代號:",
    "頁    次: 2/2",
    "班    級 姓      名 序號 學    號 1 2 3 4 5 6 7 8 9 10 出席 平時 期中 期末 總分",
]
_GLUED_FOOT = "註*:前一學期GPA平均未達2.0者。 共2頁,第{p}頁2026/9/15 列印 12:59:02"

# 每位學生實際被 pypdf 抽出來的那一（或兩）行
GLUED_BODY_1 = [
    "王小明9900010011 自資系大三",
    "林大同 【註*】9900010022 自資系大三",
    "王 小美9900010033 資工系大二",
    "李大華99000A0044 自資系大二",
    "張美玲A990000055 自資系大三",
    "9900010066 自資系大二",
    "趙小強 Jason",
    "Chao9900010077 資工系大三",
]
GLUED_BODY_2 = [
    "孫 大偉【註",
    "*】9900010088 自資系大四",
    "周小芳99001A0099 環科系碩一",
    "吳志豪99001A01010 自資系大二",
    "鄭雅婷A9900001111 環科系碩二",
    "許文彥99000101212 自資系大三",
]

PDF_LINES_GLUED = (_GLUED_HEAD_1 + GLUED_BODY_1 + [_GLUED_FOOT.format(p=1)]
                   + _GLUED_HEAD_2 + GLUED_BODY_2 + [_GLUED_FOOT.format(p=2)])
PDF_GLUED_EXPECT = [(seq, sid, name, klass)
                    for seq, sid, name, klass in GLUED_STUDENTS]
PDF_GLUED_SAMPLE_NAME = "樣本_選課名單_合成_黏連格式.pdf"

# 缺號（序號不連續）但仍應照序號編號：刪掉序號 2、4 那兩位
PDF_LINES_GAP = _GLUED_HEAD_1 + [
    ln for i, ln in enumerate(GLUED_BODY_1) if i not in (1, 3)
] + [_GLUED_FOOT.format(p=1)]


def write_glued_pdf_via_word(path):
    """用 Word COM 把 PDF_LINES_GLUED 印成兩頁的合成 PDF（只在重新產生樣本時用）。

    需要 Windows + Microsoft Word + pywin32。產生的檔案請放在
    `selftest\\樣本_選課名單_合成_黏連格式.pdf`。
    """
    import os

    import win32com.client as win32           # noqa: N813

    path = os.path.abspath(path)
    page1 = _GLUED_HEAD_1 + GLUED_BODY_1 + [_GLUED_FOOT.format(p=1)]
    page2 = _GLUED_HEAD_2 + GLUED_BODY_2 + [_GLUED_FOOT.format(p=2)]
    body = "\r".join(page1) + "\r" + chr(12) + "\r".join(page2)
    word = win32.gencache.EnsureDispatch("Word.Application")
    word.Visible = False
    doc = None
    try:
        doc = word.Documents.Add()
        doc.PageSetup.LeftMargin = 28
        doc.PageSetup.RightMargin = 28
        rng = doc.Content
        rng.Font.Name = "微軟正黑體"
        rng.Font.Size = 9
        rng.ParagraphFormat.SpaceAfter = 0
        rng.Text = body
        doc.SaveAs2(path, FileFormat=17)       # 17 = wdFormatPDF
    finally:
        if doc is not None:
            doc.Close(False)
        word.Quit()
    return path


def write_roster_pdf_via_word(path):
    """用 Word COM 把 PDF_LINES_ONELINE 印成一份小型合成 PDF（只在重新產生樣本時用）。

    需要 Windows + Microsoft Word + pywin32；失敗會丟例外，由呼叫端決定要不要忽略。
    產生的檔案請放在 `selftest\\樣本_選課名單_合成.pdf`。
    """
    import os

    import win32com.client as win32           # noqa: N813

    path = os.path.abspath(path)
    word = win32.gencache.EnsureDispatch("Word.Application")
    word.Visible = False
    doc = None
    try:
        doc = word.Documents.Add()
        rng = doc.Content
        rng.Font.Name = "微軟正黑體"
        rng.Font.Size = 11
        rng.Text = "\r".join(PDF_LINES_ONELINE)
        doc.SaveAs2(path, FileFormat=17)       # 17 = wdFormatPDF
    finally:
        if doc is not None:
            doc.Close(False)
        word.Quit()
    return path


# ---- (4e) 兩個「出現順序不同」的 Zuvio 合成檔（含名單外作答者） -----------
ZUVIO_W1_ROWS = [
    ["問題題型", "題組問答", BLANK],
    [BLANK, BLANK, BLANK],
    ["學號", "姓名", "回答"],
    [990054003, "陳曉華", "我和王小明一起查資料"],
    [990054001, "王小明", "自己完成"],
    [int(OUTSIDER[0]), OUTSIDER[1], "我已經退選但還是作答了"],
    [990054005, "李明", "和歐陽小花討論過"],
]
ZUVIO_W2_ROWS = [
    ["問題題型", "題組問答", BLANK],
    [BLANK, BLANK, BLANK],
    ["學號", "姓名", "回答"],
    [int(OUTSIDER[0]), OUTSIDER[1], "第二週我還是有作答"],
    [990054002, "林大同", "參考了陳曉華的報告"],
    [990054001, "王小明", "這週和歐陽小花同組"],
]


def _write_rows(path, rows, title="sheet1"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    wb.close()
    return path


# --------------------------------------------------------------------------
# 5) 2.2：Word（.docx）合成檔 — 名單與「要遮罩的目標檔」
# --------------------------------------------------------------------------
def _new_docx():
    try:
        import docx
    except ImportError:  # pragma: no cover
        raise RuntimeError("需要 python-docx（pip install python-docx）")
    return docx


def write_roster_docx_sample(path):
    """Word 名單（一張表格：序號｜學號｜姓名｜班級）。"""
    docx = _new_docx()
    doc = docx.Document()
    doc.add_paragraph("國立東華大學115學年度第1學期 選課名單（合成測試）")
    t = doc.add_table(rows=1, cols=4)
    hdr = ["序號", "學號", "姓名", "班級"]
    for i, h in enumerate(hdr):
        t.rows[0].cells[i].text = h
    for seq, sid, name, klass in CLASS_STUDENTS:
        c = t.add_row().cells
        for i, v in enumerate((str(seq), sid, name, klass)):
            c[i].text = v
    doc.add_paragraph("註：本表為合成資料，非真實名單。")
    doc.save(path)
    return path


def write_roster_docx_lines_sample(path):
    """Word 名單（沒有表格，只有一行一位學生的段落 → 走 PDF 那一套行解析）。"""
    docx = _new_docx()
    doc = docx.Document()
    for ln in PDF_LINES_GLUED:
        doc.add_paragraph(ln)
    doc.save(path)
    return path


# 目標檔裡會用到的合成句子（{n} = 第 n 位 CLASS_STUDENTS 的姓名）
DOCX_TARGET_BODY = [
    "第一組的組長是{1}，紀錄是{2}。",
    "{3}的學號是990054003，電子郵件 s003@gms.ndhu.edu.tw。",
    "這次報告我和{4}、{5}一起完成。",
    "本段沒有任何個資，不應該被動到。",
]
DOCX_TARGET_TABLE = [
    ["學號", "姓名", "備註"],
    ["990054001", "王小明", "和林大同同組"],
    ["990054002", "林大同", "缺席一次"],
]
DOCX_TARGET_NESTED = [["陳曉華", "990054003"]]
DOCX_TARGET_HEADER = "環境化學 分組名單　承辦：歐陽小花（990054004）"
DOCX_TARGET_FOOTER = "聯絡人：李明 990054005"
# 這一段刻意把姓名拆成兩個 run（Word 實際存檔常常這樣）
DOCX_SPLIT_RUNS = ("交件時請找王小", "明確認，學號 9900", "54001。")


def write_target_docx_sample(path):
    """要遮罩的 Word 目標檔：本文、拆 run 的姓名、表格、巢狀表格、頁首頁尾。"""
    docx = _new_docx()
    doc = docx.Document()
    names = {i: s[2] for i, s in enumerate(CLASS_STUDENTS, start=1)}
    doc.add_paragraph("環境化學 第 3 週 分組回饋（合成測試檔）")
    for tpl in DOCX_TARGET_BODY:
        s = tpl
        for i, nm in names.items():
            s = s.replace("{%d}" % i, nm)
        doc.add_paragraph(s)

    p = doc.add_paragraph()
    for chunk in DOCX_SPLIT_RUNS:
        p.add_run(chunk)

    t = doc.add_table(rows=0, cols=3)
    for row in DOCX_TARGET_TABLE:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    inner = t.rows[1].cells[2].add_table(rows=0, cols=2)
    for row in DOCX_TARGET_NESTED:
        cells = inner.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v

    sec = doc.sections[0]
    sec.header.paragraphs[0].text = DOCX_TARGET_HEADER
    sec.footer.paragraphs[0].text = DOCX_TARGET_FOOTER
    doc.save(path)
    return path


def write_zuvio_week1(path):
    return _write_rows(path, ZUVIO_W1_ROWS)


def write_zuvio_week2(path):
    return _write_rows(path, ZUVIO_W2_ROWS)


if __name__ == "__main__":
    import os
    import sys
    folder = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "selftest")
    os.makedirs(folder, exist_ok=True)
    made = [
        write_zuvio_sample(os.path.join(folder, "樣本_Zuvio合成資料.xlsx")),
        write_roster_sample(os.path.join(folder, "樣本_名冊合成資料.xlsx")),
        write_csv_sample(os.path.join(folder, "樣本_通用格式.csv")),
        write_roster_seq_sample(os.path.join(folder, "樣本_原始名單_含序號.xlsx")),
        write_roster_noseq_sample(os.path.join(folder, "樣本_原始名單_無序號.xlsx")),
        write_roster_csv_sample(os.path.join(folder, "樣本_原始名單.csv")),
    ]
    for p in made:
        print("已產生合成測試檔：" + p)
    if "--pdf" in sys.argv:
        try:
            print("已產生合成測試檔：" +
                  write_roster_pdf_via_word(os.path.join(folder, PDF_SAMPLE_NAME)))
        except Exception as e:  # noqa: BLE001
            print(f"（略過合成 PDF：{e}）")
