#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
doc_mask.py — Word（.docx）／PDF 目標檔的姓名遮罩（NameMasker 2.2）
=====================================================================
2.2 起，TA 要遮罩的「目標檔」除了 Excel／CSV 之外，也可以是 **Word 或 PDF**
（例如作業回饋單、分組名單、點名表、教務系統印出來的選課名單）。

遮罩規則與 `core_mask.py` 完全相同，只是換成在文件的文字上做：

  1. 電子郵件 → 每個字元改成 O（長度不變）。
  2. 名單上的學號（9 碼，可能含英文字母）→ 每個字元改成 O（長度不變）。
  3. 文字裡任何 ≥9 碼的數字串 → 等長的 O（學生自己打的學號也抓得到）。
  4. 名單上的姓名 → 改成該生的**學生編號**（「和王小明討論」→「和 115-1_EC_1 討論」）。
     比對時會同時試「完整姓名」「去掉空白的姓名」「中文部分（≥2 字）」，長的先比。
  5. 遮罩完做一次自我檢查：輸出文字裡不得殘留名單內的姓名或學號。

輸出檔名沿用 `core_mask.next_output_path()`：`原檔名02.docx`（已存在就 03、04…），
與原檔同資料夾，**原檔一個位元組都不動**。

  * `.docx` 目標 → 輸出 `原檔名02.docx`，本文段落、表格（含巢狀表格）、
    頁首頁尾都會遮罩，版面與格式原樣保留。
  * `.pdf`  目標 → pypdf 抽文字 → 逐行遮罩 → 輸出 `原檔名02.docx`（**文字版**，
    不保留原版面）。如果這份 PDF 本身就是一份可解析的名單，另外附一張
    「學生編號｜遮罩學號｜遮罩姓名｜班級」的遮罩後表格。

★ 對照表（{學期}_{課程}_學生名單與學生編號對照表.xlsx）才是再識別鑰匙，
  輸出的 docx 裡**不會**寫入任何真實姓名與學號的對照。

全程離線。需求：python-docx；讀 PDF 目標檔才需要 pypdf。
"""
from __future__ import annotations

import os
import re
import sys
from dataclasses import dataclass, field

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from core_mask import (  # noqa: E402
    DIGIT_RUN_RE, EMAIL_IN_TEXT_RE, MaskError, mask_name, next_output_path,
    normalize_id, validate_course, validate_semester,
)
import roster as R  # noqa: E402

__version__ = "2.2.0"

TARGET_EXT = (".docx", ".pdf")

DOCX_NEED_PKG = ("要處理 Word 檔需要 python-docx 套件"
                 "（請執行：pip install python-docx）。\n"
                 "或先用 Word 把內容另存成 Excel／純文字再處理。")
DOC_OLD_MSG = ("這是 Word 舊格式（.doc），本程式讀不了。\n"
               "請先用 Word 開啟後「另存新檔」為 .docx，再選一次。")
PDF_NO_TEXT_MSG = ("這份 PDF 抽不出任何文字，看起來是**掃描的影像檔**。\n"
                   "影像裡的姓名沒辦法用文字方式遮罩，請改用原始的 Word／Excel 檔，\n"
                   "或先用有 OCR 功能的軟體把它轉成文字型 PDF 之後再試。")
PDF_NOTE_LINE = ("本檔為 PDF 抽出文字後的遮罩版，不保留原版面。"
                 "（原始 PDF 的排版、圖片、印章都不會出現在這裡。）")

TABLE_HEADERS = ("學生編號", "遮罩學號", "遮罩姓名", "班級")


# --------------------------------------------------------------------------
# 遮罩規則
# --------------------------------------------------------------------------
_CJK = "㐀-䶿一-鿿豈-﫿"
CJK_RUN_RE = re.compile("[%s·・．]{2,}" % _CJK)


def _name_variants(name):
    """一個姓名要比對的所有寫法：完整、去空白、中文部分（≥2 字）。"""
    out = []
    n = (name or "").strip()
    if len(n) < 2:
        return out
    out.append(n)
    flat = re.sub(r"\s+", "", n)
    if flat and flat != n and len(flat) >= 2:
        out.append(flat)
    for m in CJK_RUN_RE.finditer(flat):
        run = m.group(0)
        if len(run) >= 2 and run not in out:
            out.append(run)
    return out


@dataclass
class Rules:
    """一份可重複使用的遮罩規則（由 Codebook 產生）。"""
    name_pairs: list = field(default_factory=list)   # [(要比對的姓名寫法, 學生編號)]
    sid_pairs: list = field(default_factory=list)    # [(學號, 等長的 O)]
    names: list = field(default_factory=list)        # 自我檢查用（原始姓名寫法）
    sids: list = field(default_factory=list)

    @property
    def empty(self):
        return not (self.name_pairs or self.sid_pairs)


def build_rules(codebook):
    """由 `roster.Codebook` 產生遮罩規則。codebook=None → 只遮郵件與長數字串。"""
    rules = Rules()
    if codebook is None:
        return rules
    seen_name, seen_sid = set(), set()
    for person in list(codebook.students) + list(codebook.outsiders):
        code = getattr(person, "code", "") or ""
        name = (getattr(person, "name", "") or "").strip()
        sid = normalize_id(getattr(person, "sid", ""))
        if name:
            for v in _name_variants(name):
                if v and v not in seen_name:
                    seen_name.add(v)
                    rules.name_pairs.append((v, code))
            rules.names.append(name)
        if sid and sid not in seen_sid:
            seen_sid.add(sid)
            rules.sid_pairs.append((sid, "O" * len(sid)))
            rules.sids.append(sid)
    # 長的先比，避免「王小明」被「王小」先吃掉
    rules.name_pairs.sort(key=lambda kv: -len(kv[0]))
    rules.sid_pairs.sort(key=lambda kv: -len(kv[0]))
    return rules


def _o(m):
    return "O" * len(m.group(0))


def mask_text(text, rules, counts=None):
    """把一段文字依規則遮罩，回傳 (新字串, 命中處數)。

    順序：電子郵件 → 名單上的學號 → ≥9 碼數字串 → 姓名（長的先換）。
    """
    if not isinstance(text, str) or not text:
        return text, 0
    hits = 0
    c = counts if counts is not None else {}

    text, n = EMAIL_IN_TEXT_RE.subn(_o, text)
    hits += n
    c["email"] = c.get("email", 0) + n

    for sid, dots in rules.sid_pairs:
        if sid in text:
            k = text.count(sid)
            text = text.replace(sid, dots)
            hits += k
            c["sid"] = c.get("sid", 0) + k

    text, n = DIGIT_RUN_RE.subn(_o, text)
    hits += n
    c["digits"] = c.get("digits", 0) + n

    for nm, code in rules.name_pairs:
        if nm and code and nm in text:
            k = text.count(nm)
            text = text.replace(nm, code)
            hits += k
            c["name"] = c.get("name", 0) + k
    return text, hits


def scan_residue(text, rules):
    """自我檢查：輸出文字裡還殘留哪些姓名／學號（回傳的是「筆數」，不含真名）。"""
    bad_names, bad_sids = 0, 0
    if not isinstance(text, str) or not text:
        return 0, 0
    for nm in rules.names:
        for v in _name_variants(nm):
            if len(v) >= 2 and v in text:
                bad_names += 1
                break
    for sid in rules.sids:
        if sid and sid in text:
            bad_sids += 1
    return bad_names, bad_sids


# --------------------------------------------------------------------------
# 報告物件
# --------------------------------------------------------------------------
@dataclass
class DocReport:
    src: str = ""
    dst: str = ""
    kind: str = ""                 # 'docx' / 'pdf'
    paragraphs: int = 0            # 檢查過的段落數
    changed: int = 0               # 實際被改過的段落數
    lines: int = 0                 # PDF 抽出的文字行數
    name_masked: int = 0
    id_masked: int = 0
    email_masked: int = 0
    digit_masked: int = 0
    table_rows: int = 0            # 附加的遮罩名單表列數
    residue_names: int = 0
    residue_sids: int = 0
    roster_mode: bool = False
    codebook_path: str = ""
    notes: list = field(default_factory=list)

    @property
    def ok(self):
        return self.residue_names == 0 and self.residue_sids == 0

    def summary_lines(self):
        L = []
        if self.kind == "pdf":
            L.append(f"PDF 抽出 {self.lines} 行文字（輸出為文字版 Word，不保留原版面）")
        else:
            L.append(f"檢查 {self.paragraphs} 個段落（含表格、頁首頁尾），"
                     f"遮罩 {self.changed} 個")
        L.append(f"姓名改成學生編號 {self.name_masked} 處、"
                 f"名單學號遮罩 {self.id_masked} 處")
        L.append(f"電子郵件 {self.email_masked} 處、文字內 ≥9 碼數字串 "
                 f"{self.digit_masked} 處（每字元改 O，長度不變）")
        if self.table_rows:
            L.append(f"附加遮罩後名單表 {self.table_rows} 列"
                     "（學生編號｜遮罩學號｜遮罩姓名｜班級）")
        if self.ok:
            L.append("自我檢查：輸出檔沒有殘留名單內的姓名或學號 ✔")
        else:
            L.append(f"⚠ 自我檢查：還有 {self.residue_names} 位同學的姓名、"
                     f"{self.residue_sids} 個學號殘留在輸出檔，請人工再看一次！")
        L.append(f"輸出檔：{self.dst}")
        if self.codebook_path:
            L.append("學生編號對照表：" + self.codebook_path)
        for n in self.notes:
            L.append("※ " + n)
        return L


# --------------------------------------------------------------------------
# python-docx 小工具
# --------------------------------------------------------------------------
def _docx():
    try:
        import docx
    except ImportError:
        raise MaskError(DOCX_NEED_PKG)
    return docx


def _qn(tag):
    from docx.oxml.ns import qn
    return qn(tag)


def _open_docx(path):
    docx = _docx()
    try:
        return docx.Document(path)
    except Exception as e:
        raise MaskError(
            f"這個 Word 檔打不開（可能不是有效的 .docx，或正被 Word 開著）：\n{e}")


def _story_elements(doc):
    """本文 + 所有頁首／頁尾的 XML 根元素（去重）。"""
    yield doc.element.body
    seen = set()
    for s in doc.sections:
        for attr in ("header", "footer", "first_page_header", "first_page_footer",
                     "even_page_header", "even_page_footer"):
            try:
                hf = getattr(s, attr)
            except Exception:
                continue
            el = getattr(hf, "_element", None)
            if el is None or id(el) in seen:
                continue
            seen.add(id(el))
            yield el


def _paragraph_elements(doc):
    """所有段落元素（自動涵蓋表格、巢狀表格、文字方塊、頁首頁尾）。"""
    p_tag = _qn("w:p")
    for root in _story_elements(doc):
        for p in root.iter(p_tag):
            yield p


def _mask_paragraph(p, rules, counts):
    """遮罩一個段落。先逐 run（w:t）取代；整段仍命中就把整段文字塞回第一個 run。"""
    t_tag = _qn("w:t")
    nodes = list(p.iter(t_tag))
    if not nodes:
        return False
    orig = "".join(n.text or "" for n in nodes)
    if not orig.strip():
        return False
    want, _ = mask_text(orig, rules, {})           # 整段的正確答案（先不計數）
    per = [mask_text(n.text or "", rules, {})[0] for n in nodes]
    if "".join(per) == want:
        changed = False
        for n, v in zip(nodes, per):
            if (n.text or "") != v:
                n.text = v
                changed = True
        if changed:
            mask_text(orig, rules, counts)          # 計數以整段為準
        return changed
    # 姓名被拆在多個 run（Word 常見）→ 整段遮罩後放進第一個 run，其餘清空
    if want == orig:
        return False
    nodes[0].text = want
    nodes[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    for n in nodes[1:]:
        n.text = ""
    mask_text(orig, rules, counts)
    return True


def _docx_all_text(doc):
    t_tag = _qn("w:t")
    parts = []
    for root in _story_elements(doc):
        for n in root.iter(t_tag):
            parts.append(n.text or "")
        parts.append("\n")
    return "".join(parts)


# --------------------------------------------------------------------------
# .docx 目標檔
# --------------------------------------------------------------------------
def mask_docx(src, rules, rep):
    doc = _open_docx(src)
    counts = {}
    for p in _paragraph_elements(doc):
        rep.paragraphs += 1
        if _mask_paragraph(p, rules, counts):
            rep.changed += 1
    _apply_counts(rep, counts)

    dst = next_output_path(src, ".docx")
    try:
        doc.save(dst)
    except Exception as e:
        raise MaskError(f"無法儲存新檔（請確認資料夾可寫入、檔案沒有被 Word 開著）：\n{e}")
    rep.dst = dst

    check = _open_docx(dst)
    rep.residue_names, rep.residue_sids = scan_residue(_docx_all_text(check), rules)
    return rep


# --------------------------------------------------------------------------
# .pdf 目標檔
# --------------------------------------------------------------------------
def pdf_text_lines(src):
    try:
        from pypdf import PdfReader
    except ImportError:
        raise MaskError("要處理 PDF 需要 pypdf 套件（請執行：pip install pypdf）。")
    try:
        reader = PdfReader(src)
    except Exception as e:
        raise MaskError(f"這個 PDF 打不開（可能不是有效的 PDF 或已加密）：\n{e}")
    lines = []
    for page in reader.pages:
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        for ln in txt.split("\n"):
            lines.append(ln.rstrip())
    if not any(s.strip() for s in lines):
        raise MaskError(PDF_NO_TEXT_MSG)
    return lines


def _roster_table_rows(lines, codebook):
    """這份 PDF 本身就是名單 → 回傳遮罩後的表格列（學生編號｜遮罩學號｜遮罩姓名｜班級）。"""
    if codebook is None:
        return []
    try:
        recs = R.parse_pdf_records(lines)
    except Exception:
        return []
    rows = []
    for r in recs:
        sid = normalize_id(r.sid)
        code = codebook.code_for(sid, r.name) or "（不在名單）"
        nm, _ = mask_name(r.name or "")
        rows.append((code, "O" * len(sid) if sid else "", nm, r.klass or ""))
    return rows


def mask_pdf(src, rules, rep, codebook=None):
    docx = _docx()
    lines = pdf_text_lines(src)
    rep.lines = len([x for x in lines if x.strip()])

    doc = docx.Document()
    doc.add_paragraph(PDF_NOTE_LINE)
    doc.add_paragraph("原始檔名：" + os.path.basename(src))
    doc.add_paragraph("")
    counts = {}
    for ln in lines:
        if not ln.strip():
            continue
        new, _ = mask_text(ln, rules, counts)
        doc.add_paragraph(new)
        rep.paragraphs += 1
        if new != ln:
            rep.changed += 1
    _apply_counts(rep, counts)

    rows = _roster_table_rows(lines, codebook)
    if rows:
        doc.add_paragraph("")
        doc.add_paragraph("附表：這份 PDF 解析出來的名單（已遮罩）")
        table = doc.add_table(rows=1, cols=len(TABLE_HEADERS))
        try:
            table.style = "Table Grid"
        except Exception:
            pass
        for i, h in enumerate(TABLE_HEADERS):
            table.rows[0].cells[i].text = h
        for row in rows:
            cells = table.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
        rep.table_rows = len(rows)

    dst = next_output_path(src, ".docx")
    try:
        doc.save(dst)
    except Exception as e:
        raise MaskError(f"無法儲存新檔（請確認資料夾可寫入）：\n{e}")
    rep.dst = dst

    check = _open_docx(dst)
    rep.residue_names, rep.residue_sids = scan_residue(_docx_all_text(check), rules)
    return rep


def _apply_counts(rep, counts):
    rep.email_masked += counts.get("email", 0)
    rep.id_masked += counts.get("sid", 0)
    rep.digit_masked += counts.get("digits", 0)
    rep.name_masked += counts.get("name", 0)


# --------------------------------------------------------------------------
# 對外主函式
# --------------------------------------------------------------------------
def process_document(src, course_code, semester, mask_names=True, codebook=None):
    """遮罩一份 Word／PDF 目標檔，回傳 DocReport。失敗丟 MaskError。

    `mask_names` 只是與 `core_mask.process_workbook()` 的介面一致；文件檔沒有
    「姓名欄」，文字裡的姓名一律換成學生編號（與 Excel 流程的自由文字規則相同）。
    """
    validate_course(course_code)
    validate_semester(semester)
    src = os.path.abspath(src)
    if not os.path.isfile(src):
        raise MaskError(f"找不到檔案：{src}")
    low = src.lower()
    if low.endswith(".doc"):
        raise MaskError(DOC_OLD_MSG)
    if not low.endswith(TARGET_EXT):
        raise MaskError("這個模組只處理 .docx 與 .pdf（Excel／CSV 請走原本的流程）。")

    rules = build_rules(codebook)
    rep = DocReport(src=src, kind="pdf" if low.endswith(".pdf") else "docx",
                    roster_mode=codebook is not None,
                    codebook_path=getattr(codebook, "path", "") if codebook else "")
    if codebook is None:
        rep.notes.append("沒有載入名單，只能遮罩電子郵件與 ≥9 碼數字串；"
                         "姓名不會被遮罩。建議先載入名單再處理。")
    elif rules.empty:
        rep.notes.append("對照表裡沒有任何姓名與學號，只遮罩電子郵件與長數字串。")

    if rep.kind == "pdf":
        mask_pdf(src, rules, rep, codebook)
    else:
        mask_docx(src, rules, rep)

    if codebook is not None:
        try:
            codebook.save()
        except Exception as e:  # noqa: BLE001
            rep.notes.append(f"對照表回寫失敗（請確認沒有被 Excel 開著）：{e}")
    return rep


# --------------------------------------------------------------------------
def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    files = [a for a in argv if not a.startswith("-")]
    if not files:
        sys.exit("用法：python doc_mask.py 檔案.(docx|pdf) --codebook 對照表.xlsx "
                 "[--course EC --semester 115-1]")
    course, sem, cbp = "EC", "115-1", ""
    for i, a in enumerate(argv):
        if a == "--course" and i + 1 < len(argv):
            course = argv[i + 1]
        if a == "--semester" and i + 1 < len(argv):
            sem = argv[i + 1]
        if a == "--codebook" and i + 1 < len(argv):
            cbp = argv[i + 1]
    files = [f for f in files if f not in (course, sem, cbp)]
    cb = None
    if cbp:
        cb, msgs = R.build_codebook(codebook_path=cbp, semester=sem, course=course)
        for m in msgs:
            print(m)
    for f in files:
        try:
            rep = process_document(f, course, sem, codebook=cb)
        except MaskError as e:
            print(f"失敗：{f}\n{e}")
            continue
        for line in rep.summary_lines():
            print(line)
    return 0


if __name__ == "__main__":
    main()
