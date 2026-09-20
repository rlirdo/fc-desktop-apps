# -*- coding: utf-8 -*-
"""
roster.py — 原始名單 → 固定學生編號 → 對照表（HomeworkWordCloud 2.1）

**為什麼要有這個模組**

2.0 的學生編號是「依每個檔案內首次出現順序」給的，所以同一位同學在
Q01 可能是 `115-1_EC_3`、在 Q02 卻變成 `115-1_EC_7` —— 沒辦法跨題、跨週追蹤，
這是已經確認的缺陷。2.1 改成：**先拿 TA 手上的原始名單產生一份固定對照表**，
之後每一次處理都照這份對照表給編號，整學期不變。

**流程**

    原始名單（Excel／CSV／東華選課名單 PDF）
        │  parse_roster_file()
        ▼
    [{序號, 學號, 姓名, 班級}, …]
        │  assign_codes()      學生編號＝{學期}_{課程縮寫}_{序號}
        ▼
    對照表 xlsx：{學期}_{課程縮寫}_學生名單與學生編號對照表.xlsx
        工作表「學生編號對照」：序號|學生編號|學號|姓名|班級
        工作表「名單外作答者」：學生編號|學號|姓名|首次出現檔案（101 起，持久登記）
        工作表「說明」：規則與警語

**與「姓名遮罩與學生編號」（NameMasker 2.1）完全相容**：欄位名、檔名、工作表名、
101 起的名單外登記規則都一樣，兩支程式可以共用同一份對照表。

⚠ 對照表含真實姓名與學號，是**再識別鑰匙**，只能留在本機，不得上傳或外流。
"""
import os
import re
import csv
import datetime as dt

import openpyxl

# ------------------------------------------------------------------ 常數（兩支 App 必須一致）
CODEBOOK_SHEET = "學生編號對照"
OUTSIDE_SHEET = "名單外作答者"
NOTE_SHEET = "說明"

CODEBOOK_COLS = ["序號", "學生編號", "學號", "姓名", "班級"]
OUTSIDE_COLS = ["學生編號", "學號", "姓名", "首次出現檔案"]

OUTSIDE_START = 101          # 名單外作答者（退選等）的編號從 101 開始

SID_LEN = 9                  # 東華學號長度
CJK = r"㐀-䶿一-鿿"

# 東華教務系統「選課名單」PDF：抽字後每位學生一行，欄位可能黏在一起也可能用空白分開
#   `王小明9900540011 自資系大三`      → 姓名 + 9 碼學號 + 序號 + 班級（黏在一起）
#   `王小明【註*】9900540015 自資系大三`
#   `甲小明 990000001 1 自資系大三`     → 用空白分開
PDF_DIGITS = re.compile(r"\d+")
PDF_SID_ONLY = re.compile(r"^\d{" + str(SID_LEN) + r"}$")
PDF_NOTE = re.compile(r"【[^】]*】")
PDF_NAME_ONLY = re.compile(r"^[" + CJK + r"A-Za-z·．\.\- ]{2,12}$")
PDF_CLASS_ONLY = re.compile(r"^[" + CJK + r"]{2,12}(?:大|碩|博|專)[一二三四五六七八九十]?$")
PDF_CLASS_HEAD = re.compile(r"^[" + CJK + r"A-Za-z]{2,20}")

# 名單表頭可能的欄位名
H_SID = ("學號", "學生證號", "student id", "studentid", "id")
H_NAME = ("姓名", "名字", "學生姓名", "name")
H_SEQ = ("序號", "編號", "座號", "no", "no.")
H_CLASS = ("班級", "系級", "系所", "班別", "class")
H_CODE = ("學生編號",)
H_MEMO = ("備註", "註記")
H_SRC = ("首次出現檔案",)


class RosterError(Exception):
    """給使用者看的白話錯誤（GUI 直接顯示，不印 traceback）。"""


# ------------------------------------------------------------------ 小工具
def norm_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        v = int(v)
    return str(v).replace("　", " ").strip()


def normalize_sid(v):
    """學號正規化：int/float → str、去掉空白與不可見字元。"""
    s = norm_text(v)
    s = re.sub(r"[\s\-​]+", "", s)
    return s


def norm_course_code(s, default="EC"):
    s = norm_text(s).upper()
    return s if re.match(r"^[A-Z]{2}$", s) else default


def norm_semester(s, default="115-1"):
    s = norm_text(s).replace("／", "/").replace("/", "-")
    return s if re.match(r"^\d{3}-[12]$", s) else default


def make_code(semester, course, n):
    return f"{norm_semester(semester)}_{norm_course_code(course)}_{int(n)}"


def codebook_filename(semester, course):
    return f"{norm_semester(semester)}_{norm_course_code(course)}_學生名單與學生編號對照表.xlsx"


def _match_header(cell, names):
    c = norm_text(cell).lower().replace(" ", "")
    if not c:
        return False
    return any(c == n.lower().replace(" ", "") or n.lower().replace(" ", "") in c
               for n in names)


def _find_header(rows):
    """在前 5 列裡找同時有「學號」與「姓名」的表頭列。回傳 (列索引, 欄位對應) 或 None。"""
    for i, r in enumerate(rows[:5]):
        cols = {}
        for j, c in enumerate(r):
            for key, names in (("學號", H_SID), ("姓名", H_NAME), ("序號", H_SEQ),
                               ("班級", H_CLASS), ("學生編號", H_CODE),
                               ("備註", H_MEMO), ("首次出現檔案", H_SRC)):
                if key not in cols and _match_header(c, names):
                    cols[key] = j
                    break
        if "學號" in cols and "姓名" in cols:
            return i, cols
    return None


# ------------------------------------------------------------------ 讀檔：Excel / CSV
def _sheet_rows(path):
    """回傳 [(工作表名, [[儲存格字串, …], …]), …]。"""
    out = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sn in wb.sheetnames:
            ws = wb[sn]
            rows = [[norm_text(c) for c in row]
                    for row in ws.iter_rows(values_only=True)]
            out.append((sn, rows))
    finally:
        wb.close()
    return out


def _csv_rows(path):
    for enc in ("utf-8-sig", "cp950", "utf-8"):
        try:
            with open(path, encoding=enc, newline="") as f:
                return [[norm_text(c) for c in r] for r in csv.reader(f)]
        except UnicodeDecodeError:
            continue
    with open(path, encoding="utf-8", errors="replace", newline="") as f:
        return [[norm_text(c) for c in r] for r in csv.reader(f)]


def _rows_to_students(rows, hdr_i, cols):
    """把表格列轉成 [{序號, 學號, 姓名, 班級, 學生編號}]。"""
    out = []
    for r in rows[hdr_i + 1:]:
        if not any(r):
            continue
        sid = normalize_sid(r[cols["學號"]]) if cols["學號"] < len(r) else ""
        name = norm_text(r[cols["姓名"]]) if cols["姓名"] < len(r) else ""
        if not sid and not name:
            continue
        # 表頭重複出現（有些名冊每頁一個表頭）
        if sid in ("學號",) or name in ("姓名",):
            continue
        rec = {"學號": sid, "姓名": PDF_NOTE.sub("", name).strip(), "序號": "",
               "班級": "", "學生編號": ""}
        for key in ("序號", "班級", "學生編號"):
            j = cols.get(key)
            if j is not None and j < len(r):
                rec[key] = norm_text(r[j])
        out.append(rec)
    return out


# ------------------------------------------------------------------ 讀檔：PDF
def pdf_text_lines(path):
    """抽出 PDF 的文字行（純 Python 的 pypdf，不用 PyMuPDF，避免 exe 暴增）。"""
    try:
        from pypdf import PdfReader
    except Exception as e:                                   # pragma: no cover
        raise RosterError("這台電腦缺少讀 PDF 的元件（pypdf），無法解析 PDF 名單。\n"
                          "請改用 Excel 名單（需有 學號、姓名 欄）。\n"
                          f"（技術訊息：{e}）")
    try:
        reader = PdfReader(path)
        lines = []
        for page in reader.pages:
            for ln in (page.extract_text() or "").splitlines():
                ln = ln.replace("　", " ").rstrip()
                if ln.strip():
                    lines.append(ln)
        return lines
    except RosterError:
        raise
    except Exception as e:
        raise RosterError(f"這份 PDF 打不開或讀不出文字：{os.path.basename(path)}\n"
                          "如果它是掃描影像的 PDF，請改用 Excel 名單"
                          "（需有 學號、姓名 欄）。\n"
                          f"（技術訊息：{e}）")


def parse_pdf_line(ln):
    """單行 → {序號, 學號, 姓名, 班級}；不是學生資料就回 None。

    以「9 碼（含以上）的數字串」當錨點：
      * 長度 > 9 → 前 9 碼是學號，其餘是序號（真實檔案兩欄是黏在一起的）
      * 長度 = 9 → 序號在同一行的另一個 1–3 位數字串
    數字串前面的文字是姓名，後面的文字是班級；`【註*】` 一律忽略。
    表頭與頁尾雜訊行沒有 9 碼數字串，自然被濾掉。
    """
    s = PDF_NOTE.sub(" ", ln).strip()
    runs = [(m.start(), m.end(), m.group(0)) for m in PDF_DIGITS.finditer(s)]
    big = [r for r in runs if len(r[2]) >= SID_LEN]
    if not big:
        return None
    b = big[0]
    sid, seq = b[2][:SID_LEN], b[2][SID_LEN:]
    if not seq:
        others = [r for r in runs if r is not b and 1 <= len(r[2]) <= 3]
        seq = others[0][2] if others else ""

    name = s[:b[0]].strip(" 　．.-、,")
    if not (name and PDF_NAME_ONLY.match(name)):
        # 姓名不在數字前面（欄位順序不同）→ 找行內第一個不像班級的中文詞
        cand = [t.strip() for t in re.split(r"[\d\s]+", s) if t.strip()]
        cand = [t for t in cand if PDF_NAME_ONLY.match(t) and not PDF_CLASS_ONLY.match(t)]
        name = cand[0] if cand else ""
    name = name.replace(" ", "").strip()
    if not name:
        return None

    tail = re.sub(r"^[\s\d　]+", "", s[b[1]:]).strip()
    m = PDF_CLASS_HEAD.match(tail)
    klass = m.group(0).strip() if m else ""
    return {"序號": int(seq) if str(seq).isdigit() else "", "學號": sid,
            "姓名": name, "班級": klass, "學生編號": ""}


def parse_pdf_lines(lines):
    """把「抽字後的行列表」解析成 [{序號, 學號, 姓名, 班級}]。

    支援兩種排法（真實檔案與規格書各見過一種）：
      A. 單行：`姓名`＋（`【註*】`）＋`9 碼學號`＋`序號`＋`班級`（黏著或用空白分開都可以）
      B. 逐行分開：`姓名` / （`【註*】`）/ `9 碼學號` / `序號` / `班級`
    表頭與頁尾雜訊行（含「選課名單」「第N頁」「註*:」…）一律忽略。
    """
    lines = [ln.strip() for ln in lines if ln and ln.strip()]

    # ---- A. 單行
    out = []
    for ln in lines:
        rec = parse_pdf_line(ln)
        if rec:
            out.append(rec)
    if out:
        return out

    # ---- B. 以「9 碼學號行」為錨點，姓名／序號／班級在鄰近行
    idx = [i for i, ln in enumerate(lines) if PDF_SID_ONLY.match(ln)]
    for k, i in enumerate(idx):
        sid = lines[i]
        # 姓名：往前找最近的純姓名行（跳過【註*】）
        name = ""
        for j in range(i - 1, max(-1, i - 4), -1):
            cand = PDF_NOTE.sub("", lines[j]).strip()
            if not cand:
                continue
            if PDF_NAME_ONLY.match(cand) and not PDF_CLASS_ONLY.match(cand):
                name = cand
                break
        # 序號、班級：往後找
        seq, klass = "", ""
        stop = idx[k + 1] if k + 1 < len(idx) else len(lines)
        for j in range(i + 1, min(stop, i + 5)):
            cand = lines[j]
            if not seq and cand.isdigit():
                seq = cand
                continue
            if not klass and PDF_CLASS_ONLY.match(cand):
                klass = cand
        if not name:
            continue
        out.append({"序號": int(seq) if seq.isdigit() else "", "學號": sid,
                    "姓名": name, "班級": klass, "學生編號": ""})
    return out


def parse_pdf_roster(path):
    students = parse_pdf_lines(pdf_text_lines(path))
    if not students:
        raise RosterError(
            f"這份 PDF 讀不出名單（{os.path.basename(path)}），"
            "請改用 Excel 名單（需有 學號、姓名 欄）。")
    seqs = [s["序號"] for s in students if isinstance(s["序號"], int)]
    if seqs and sorted(seqs) != list(range(1, len(students) + 1)):
        raise RosterError(
            f"這份 PDF 讀不出名單（{os.path.basename(path)}）："
            f"序號不連續（讀到 {len(students)} 人，序號 "
            f"{min(seqs)}–{max(seqs)}，可能有學生的那一行沒被讀到）。\n"
            "請改用 Excel 名單（需有 學號、姓名 欄）。")
    return students


# ------------------------------------------------------------------ 解析名單檔
def parse_roster_file(path):
    """讀一份名單／對照表 → (students, kind, outside)。

    kind:
        'codebook'  這份檔案本身就是對照表（有「學生編號」欄）→ 直接採用、不重編
        'table'     Excel／CSV 名單
        'pdf'       東華選課名單 PDF
    """
    path = os.path.abspath(path)
    if not os.path.exists(path):
        raise RosterError(f"找不到名單檔：{path}")
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        return parse_pdf_roster(path), "pdf", []

    if ext in (".csv", ".txt"):
        rows = _csv_rows(path)
        found = _find_header(rows)
        if not found:
            raise RosterError(
                f"這份 CSV 的前 5 列找不到「學號」與「姓名」欄：{os.path.basename(path)}\n"
                "請確認第一列是欄位名稱（至少要有 學號、姓名）。")
        hdr_i, cols = found
        students = _rows_to_students(rows, hdr_i, cols)
        kind = "codebook" if "學生編號" in cols and any(s["學生編號"] for s in students) \
            else "table"
        return students, kind, []

    if ext not in (".xlsx", ".xlsm"):
        raise RosterError(f"看不懂的名單格式：{os.path.basename(path)}\n"
                          "支援 .xlsx／.csv／.pdf。")

    sheets = _sheet_rows(path)
    outside = []
    # 1) 先找對照表格式的工作表（有「學生編號」欄）
    for sn, rows in sheets:
        found = _find_header(rows)
        if found and "學生編號" in found[1]:
            students = _rows_to_students(rows, found[0], found[1])
            students = [s for s in students if s["學生編號"]]
            if students:
                outside = _read_outside(sheets)
                return students, "codebook", outside
    # 2) 一般名單
    for sn, rows in sheets:
        found = _find_header(rows)
        if found:
            students = _rows_to_students(rows, found[0], found[1])
            if students:
                return students, "table", []
    raise RosterError(
        f"這份 Excel 的任一工作表前 5 列都找不到「學號」與「姓名」欄："
        f"{os.path.basename(path)}\n"
        "請確認名單至少有 學號、姓名 兩欄（可以另有 序號、班級）。")


def _read_outside(sheets):
    """從對照表檔讀「名單外作答者」工作表。"""
    for sn, rows in sheets:
        if sn.strip() != OUTSIDE_SHEET:
            continue
        if not rows:
            return []
        hdr = [norm_text(c) for c in rows[0]]
        try:
            i_code = hdr.index("學生編號")
        except ValueError:
            return []
        i_sid = hdr.index("學號") if "學號" in hdr else -1
        i_name = hdr.index("姓名") if "姓名" in hdr else -1
        i_src = hdr.index("首次出現檔案") if "首次出現檔案" in hdr else -1
        out = []
        for r in rows[1:]:
            if not any(r):
                continue
            code = norm_text(r[i_code]) if i_code < len(r) else ""
            if not code:
                continue
            out.append({
                "學生編號": code,
                "學號": normalize_sid(r[i_sid]) if 0 <= i_sid < len(r) else "",
                "姓名": norm_text(r[i_name]) if 0 <= i_name < len(r) else "",
                "首次出現檔案": norm_text(r[i_src]) if 0 <= i_src < len(r) else "",
            })
        return out
    return []


# ------------------------------------------------------------------ 編號
def assign_codes(students, semester, course):
    """依 §1.2 給固定學生編號。

    有「序號」欄且是 1..N 不重複 → 學生編號＝{學期}_{課程縮寫}_{序號}
    否則依**學號字串遞增排序**給 1..N。
    """
    semester = norm_semester(semester)
    course = norm_course_code(course)

    sids = [s["學號"] for s in students]
    dup = {s for s in sids if s and sids.count(s) > 1}
    if dup:
        raise RosterError("名單裡有重複的學號："
                          + "、".join(sorted(dup)[:5])
                          + ("…" if len(dup) > 5 else "")
                          + "\n請先在名單檔裡處理掉重複的列，再跑一次。")

    seqs = []
    for s in students:
        v = s.get("序號")
        try:
            seqs.append(int(str(v).strip()))
        except (TypeError, ValueError):
            seqs = []
            break
    use_seq = bool(seqs) and sorted(seqs) == list(range(1, len(students) + 1))

    if use_seq:
        rule = "依名單的序號"
        for s in students:
            n = int(str(s["序號"]).strip())
            s["序號"] = n
            s["學生編號"] = make_code(semester, course, n)
        students = sorted(students, key=lambda s: s["序號"])
    else:
        rule = "依學號字串遞增排序"
        students = sorted(students, key=lambda s: (s["學號"], s["姓名"]))
        for n, s in enumerate(students, 1):
            s["序號"] = n
            s["學生編號"] = make_code(semester, course, n)
    return students, rule


# ------------------------------------------------------------------ 對照表
class Codebook:
    """一份固定學生編號對照表（含名單外作答者的持久登記）。"""

    def __init__(self, semester, course, students=None, outside=None, path="",
                 rule=""):
        self.semester = norm_semester(semester)
        self.course = norm_course_code(course)
        self.students = list(students or [])
        self.outside = list(outside or [])
        self.path = path
        self.rule = rule
        self.dirty = False
        self._reindex()

    # -------------------------------------------------- 索引
    def _reindex(self):
        self.by_sid = {}
        self.by_name = {}
        for s in self.students:
            if s.get("學號"):
                self.by_sid[s["學號"]] = s["學生編號"]
            if s.get("姓名"):
                self.by_name.setdefault(s["姓名"], s["學生編號"])
        for o in self.outside:
            if o.get("學號"):
                self.by_sid.setdefault(o["學號"], o["學生編號"])
            if o.get("姓名"):
                self.by_name.setdefault(o["姓名"], o["學生編號"])

    @property
    def count(self):
        """全班人數＝「學生編號對照」工作表的人數。"""
        return len(self.students)

    def names(self):
        """名單內**所有**真實姓名（§1.4：都要納入自由文字替換）。"""
        return [s["姓名"] for s in self.students if s.get("姓名")] + \
               [o["姓名"] for o in self.outside if o.get("姓名")]

    def sids(self):
        return [s["學號"] for s in self.students if s.get("學號")] + \
               [o["學號"] for o in self.outside if o.get("學號")]

    def name_to_code(self):
        return {s["姓名"]: s["學生編號"] for s in self.students if s.get("姓名")} | \
               {o["姓名"]: o["學生編號"] for o in self.outside if o.get("姓名")}

    def codes(self):
        return [s["學生編號"] for s in self.students]

    # -------------------------------------------------- 查表／登記
    def _next_outside_n(self):
        used = [OUTSIDE_START - 1]
        for o in self.outside:
            m = re.search(r"_(\d+)$", o.get("學生編號", "") or "")
            if m and int(m.group(1)) >= OUTSIDE_START:
                used.append(int(m.group(1)))
        return max(used) + 1

    def code_for(self, sid, name, source_file=""):
        """學號優先、姓名次之；名單外的作答者持久登記（101 起）。"""
        sid = normalize_sid(sid)
        name = norm_text(name)
        if sid and sid in self.by_sid:
            return self.by_sid[sid]
        if name and name in self.by_name:
            return self.by_name[name]
        if not sid and not name:
            return ""
        code = make_code(self.semester, self.course, self._next_outside_n())
        self.outside.append({"學生編號": code, "學號": sid, "姓名": name,
                             "首次出現檔案": source_file})
        if sid:
            self.by_sid[sid] = code
        if name:
            self.by_name[name] = code
        self.dirty = True
        return code

    # -------------------------------------------------- 存檔
    def default_path(self, folder):
        return os.path.join(folder, codebook_filename(self.semester, self.course))

    def save(self, path=None, source=""):
        path = os.path.abspath(path or self.path)
        folder = os.path.dirname(path)
        if folder:
            os.makedirs(folder, exist_ok=True)
        wb = openpyxl.Workbook()

        ws = wb.active
        ws.title = CODEBOOK_SHEET
        ws.append(CODEBOOK_COLS)
        for s in self.students:
            ws.append([s.get("序號", ""), s.get("學生編號", ""), s.get("學號", ""),
                       s.get("姓名", ""), s.get("班級", "")])

        ws2 = wb.create_sheet(OUTSIDE_SHEET)
        ws2.append(OUTSIDE_COLS)
        for o in self.outside:
            ws2.append([o.get("學生編號", ""), o.get("學號", ""), o.get("姓名", ""),
                        o.get("首次出現檔案", "")])

        ws3 = wb.create_sheet(NOTE_SHEET)
        for line in [
            f"來源：{source or self.path or '（未記錄）'}",
            f"學生編號規則：{self.semester}_{self.course}_{{序號}}"
            f"（{self.rule or '依名單的序號'}），整學期固定，"
            "不因作業出現順序而改變。",
            f"名單外作答者（退選等）：編號自 {OUTSIDE_START} 起遞增，跨檔跨週一致，"
            "每次處理完都會回寫「名單外作答者」工作表。",
            "用途：Zuvio 作業去識別化（姓名遮罩與學生編號、學生作業文字雲兩支程式共用本檔）。",
            f"產生時間：{dt.datetime.now():%Y-%m-%d %H:%M:%S}"
            f"（HomeworkWordCloud 2.1）",
            "⚠ 本檔含真實姓名與學號＝再識別鑰匙，只留本機，"
            "不得上傳網路、不得放進 git、不得分享。",
        ]:
            ws3.append([line])

        for w, widths in ((ws, (8, 18, 14, 14, 16)), (ws2, (18, 14, 14, 34)),
                          (ws3, (110,))):
            for i, wd in enumerate(widths, 1):
                w.column_dimensions[openpyxl.utils.get_column_letter(i)].width = wd

        wb.save(path)
        wb.close()
        self.path = path
        self.dirty = False
        return path

    def save_if_dirty(self, log=print):
        if self.dirty and self.path:
            self.save(self.path)
            log(f"  對照表已回寫（名單外作答者 {len(self.outside)} 人）：{self.path}")
            return True
        return False


# ------------------------------------------------------------------ 進入點
def load_codebook(path):
    """直接讀一份既有的對照表 xlsx。"""
    students, kind, outside = parse_roster_file(path)
    if kind != "codebook":
        raise RosterError(f"這份檔案不是對照表（沒有「學生編號」欄）：{os.path.basename(path)}")
    sem, course = _guess_sem_course(students)
    cb = Codebook(sem, course, students, outside, path=os.path.abspath(path),
                  rule="沿用既有對照表")
    return cb


def _guess_sem_course(students, semester="", course=""):
    """從既有編號反推學期與課程縮寫（對照表可能來自別的課）。"""
    for s in students:
        m = re.match(r"^(\d{3}-[12])_([A-Z]{2})_\d+$", norm_text(s.get("學生編號")))
        if m:
            return m.group(1), m.group(2)
    return norm_semester(semester), norm_course_code(course)


def prepare_codebook(path, semester, course, overwrite=False, out_dir=None,
                     log=print):
    """TA 給一份「原始名單或學生編號對照表」→ 回傳可用的 Codebook。

    * 檔案本身就是對照表 → 直接採用、不重編。
    * 否則解析名單 → 編號 → 存成 `{學期}_{課程縮寫}_學生名單與學生編號對照表.xlsx`
      （預設放在名單同資料夾）。
      已存在且 overwrite=False → **沿用既有對照表**（保住名單外作答者的登記）。
    """
    path = os.path.abspath(path)
    students, kind, outside = parse_roster_file(path)

    if kind == "codebook":
        sem, crs = _guess_sem_course(students, semester, course)
        cb = Codebook(sem, crs, students, outside, path=path, rule="沿用既有對照表")
        log(f"  對照表：{os.path.basename(path)}（{cb.count} 人，"
            f"名單外 {len(cb.outside)} 人）")
        log(f"  學生編號規則：{cb.semester}_{cb.course}_{{序號}}（沿用既有對照表，不重編）")
        return cb

    if not students:
        raise RosterError(f"這份名單讀不到任何學生：{os.path.basename(path)}")

    students, rule = assign_codes(students, semester, course)
    cb = Codebook(semester, course, students, [], rule=rule)
    folder = out_dir or os.path.dirname(path) or os.getcwd()
    target = cb.default_path(folder)

    if os.path.exists(target) and not overwrite:
        try:
            old = load_codebook(target)
        except RosterError:
            old = None
        if old is not None and old.count:
            log(f"  對照表已存在，直接沿用（不重編）：{target}")
            log(f"    名單檔 {os.path.basename(path)} 讀到 {len(students)} 人、"
                f"既有對照表 {old.count} 人、名單外 {len(old.outside)} 人")
            if old.count != len(students):
                log("    [提醒] 兩者人數不同（可能有加退選）。要改用新名單重編，"
                    "請勾選／加上「覆寫對照表」。")
            return old

    cb.save(target, source=path)
    log(f"  名單檔：{os.path.basename(path)}　讀到 {len(students)} 人")
    log(f"  學生編號規則：{cb.semester}_{cb.course}_{{序號}}（{rule}）")
    log(f"  對照表已產生：{target}")
    return cb


def describe(cb):
    """給 GUI 顯示的一行摘要。"""
    if cb is None:
        return "尚未載入名單／對照表"
    return (f"讀到 {cb.count} 人　編號規則：{cb.semester}_{cb.course}_{{序號}}"
            f"（{cb.rule}）　名單外 {len(cb.outside)} 人")
