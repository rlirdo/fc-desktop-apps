# -*- coding: utf-8 -*-
"""
parse_zuvio.py — 把作業匯出檔轉成每題一份「去識別化」CSV。（2.0）

支援三種輸入：
  (1) **NameMasker 2.0 的輸出 xlsx**（建議走這條）
      整張 sheet 的 A 欄被插入「學生編號」，表頭列 A 欄寫著「學生編號」，
      學號欄已經全部是 O，姓名已遮罩，自由文字裡的同學姓名已換成學生編號。
      檔案裡另有一張工作表「學生編號連結姓名」＝再識別鑰匙，
      **本程式一律忽略那張表，絕不讀入其中任何姓名。**
  (2) 原始 Zuvio 老師端「下載數據」匯出 xlsx（向下相容）
      這時由 core/mask.py 自行產生學生編號並完成遮罩。
  (3) 通用 CSV（欄位含「姓名」與「作答」即可）。

另外兩種會混進資料夾、但**不是作業**的檔案，一律安全處理不崩潰：
  * **名冊檔**：第 1 列表頭像 `[None, '學號', '姓名', None]`，沒有「回答」欄
    → log「非作業檔，略過」，不編題號、不產 CSV。
  * **測驗題目型匯出**：中繼列 R1–R12（資料夾名稱／問題題型／是否分組／是否匿名／
    總分／題幹／平均／分數／人數…），表頭在 R13：
    `學號|姓名|作答時間|總分|第1題|第2題…`，答案是「答對／答錯」這類短詞
    → 視為選擇題型（統計答對率／選項分佈），不做文字雲與概念矩陣。

實測補充：原始 Zuvio 匯出的學號儲存格是 **int**（411354043），
NameMasker 2.0 輸出後是字串 `'OOOOOOOOO'`；本模組一律先 `str()` 再處理。
同一張 sheet 可能有多達 5 個表頭區塊（例：列 12／62／80／132／184），
每個區塊都會各自歸到對應的子題。

學生主鍵：**A 欄「學生編號」優先**；沒有才退回 9 碼學號，再退回姓名。
作答人數 = 不重複學生編號數。

Zuvio xlsx 的兩種資料區塊（已於真實匯出檔確認）：
    A. 上方彙總表：學號 | 姓名 | 作答時間 | 回答 | 回答 | …（一欄一子題）
    B. 下方逐子題明細區塊：開頭是「第N題:題型」，表頭 學號 | 姓名 | … | 回答
    -> 只要 B 區塊存在，一律以 B 為準；沒有 B 才退回 A。

輸出 CSV 欄位：
    題號,子題號,子題題目,學生編號,學號,姓名,作答時間,作答內容
"""
import os
import re
import csv
import glob

import openpyxl

from .mask import (StudentCoder, build_roster, build_replacer, mask_in_text,
                   mask_name, mask_id, is_student_id, normalize_course_code,
                   normalize_semester, NOT_A_NAME)

CSV_COLS = ["題號", "子題號", "子題題目", "學生編號", "學號", "姓名", "作答時間", "作答內容"]

META_KEYS = ("資料夾名稱", "問題題型", "是否分組", "是否匿名", "題幹", "問題敘述",
             "正解", "總分")

# NameMasker 2.0 產生的「再識別鑰匙」工作表 —— 一律跳過，絕不讀入
LINK_SHEET_NAMES = {"學生編號連結姓名", "學生編號對照", "編號連結姓名"}

CODE_HEADER = "學生編號"


def norm(v):
    if v is None:
        return ""
    return str(v).replace("　", " ").strip()


TITLE_MAX = 26          # 題目標題上限（字），超過加刪節號


def clean_title(s, qtype="", limit=TITLE_MAX):
    """把題幹整理成可以當標題的短字串。

    真實匯出檔的「題幹」常常黏著整串教材網址，例如
        `循環式思考_AI協作(練習)_(1.1-1.2)_115-1 GC1. 閱讀文章 Ch 1 https://elearn4…`
    一路灌進標題與總覽表。這裡：
      1. 從第一個 http／https 起整段截掉
      2. 去掉開頭重複的題型字樣（「題組問答介紹一下自己吧！」→「介紹一下自己吧！」）
      3. 壓掉多餘空白，最後截到 limit 個字
    """
    s = str(s or "").strip()
    m = re.search(r"https?://", s, re.I)
    if m:
        s = s[:m.start()]
    s = re.sub(r"[\s　]+", " ", s).strip(" 　-_、,，.。:：")
    qtype = str(qtype or "").strip()
    if qtype and s.startswith(qtype) and len(s) > len(qtype):
        s = s[len(qtype):].lstrip(" 　-_、,，")
    if len(s) > limit:
        s = s[:limit].rstrip()
        # 切一半的括號很難看（…GC(檔），把沒收尾的那一組整段砍掉
        for op, cl in (("(", ")"), ("（", "）"), ("[", "]"), ("【", "】"),
                       ("《", "》"), ("〈", "〉")):
            if s.count(op) > s.count(cl):
                s = s[:s.rfind(op)]
        s = s.rstrip(" 　([（【{《〈_—、,，.。:：;；/|-")
        s += "…"
    return s


def safe_name(s, n=12):
    s = re.sub(r"https?://\S+", "", s or "")
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f\s]+', "", s)
    return s[:n] or "untitled"


def pick_sheet(wb):
    """回傳要解析的工作表名稱；一律避開「學生編號連結姓名」。"""
    for nm in wb.sheetnames:
        if nm.strip() not in LINK_SHEET_NAMES:
            return nm
    return wb.sheetnames[0]


def read_grid(path):
    """讀第一張（非連結表）工作表成 list of list of str。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[pick_sheet(wb)]
    grid = [[norm(c) for c in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    return grid


def split_code_column(grid):
    """若 A 欄是 NameMasker 2.0 插入的「學生編號」欄，把它抽出來。

    回傳 (codes, body, has_codes)：
        codes[i] = 第 i 列的學生編號（可能是空字串／「匿名」）
        body     = 去掉 A 欄之後的表格（欄位位置與原始 Zuvio 匯出一致）
    """
    has = any(r and r[0] == CODE_HEADER for r in grid)
    if not has:
        return [""] * len(grid), grid, False
    codes = [(r[0] if r else "") for r in grid]
    body = [(list(r[1:]) if r else []) for r in grid]
    return codes, body, True


# ------------------------------------------------------------------ 檔案類型
def header_rows(grid):
    """回傳所有「作答明細表頭列」的索引：該列同時有『學號』與『姓名』。"""
    out = []
    for i, r in enumerate(grid):
        cells = [c for c in r if c]
        if "學號" in cells and "姓名" in cells:
            out.append(i)
    return out


def looks_like_roster_file(grid):
    """名冊檔：有 學號/姓名 表頭，但整張表沒有『回答』欄、也沒有『第N題』區塊。"""
    if not header_rows(grid):
        return False
    for r in grid:
        for c in r:
            if c in ("回答", "作答內容", "答案"):
                return False
            if c and re.match(r"^第\d+題", c):
                return False
    return True


def looks_like_quiz(grid):
    """測驗題目型匯出：表頭是 學號|姓名|作答時間|總分|第1題|…，而且沒有『回答』欄。"""
    has_ans_col = any(c == "回答" for r in grid for c in r)
    if has_ans_col:
        return False
    for r in grid:
        if len(r) >= 5 and r[0] == "學號" and r[1] == "姓名" and r[2] == "作答時間":
            if any((c or "").startswith("第") for c in r[3:]):
                return True
    return False


# ------------------------------------------------------------------ B 區塊
def parse_details(grid):
    """逐子題明細區塊（第N題:題型）。

    回傳 [{label, text, answers:[(學號, 姓名, 回答, 列索引)]}]。
    """
    marks = [i for i, r in enumerate(grid) if r and re.match(r"^第\d+題", r[0] or "")]
    out = []
    for k, i in enumerate(marks):
        end = marks[k + 1] if k + 1 < len(marks) else len(grid)
        label = grid[i][0]
        text = ""
        hdr_i = None
        for j in range(i + 1, end):
            r = grid[j]
            if not r:
                continue
            if r[0] == "問題敘述" and len(r) > 1 and not text:
                text = r[1]
            if r[0] == "學號" and len(r) > 1 and r[1] == "姓名":
                hdr_i = j
                break
        if hdr_i is None:
            continue
        hdr = grid[hdr_i]
        try:
            acol = hdr.index("回答")
        except ValueError:
            acol = 2
        answers = []
        for j in range(hdr_i + 1, end):
            r = grid[j]
            if not r or not r[0] or r[0] == "學號":
                break
            ans = r[acol] if acol < len(r) else ""
            answers.append((r[0], r[1] if len(r) > 1 else "", ans, j))
        out.append({"label": label, "text": text, "answers": answers})
    return out


def parse_answer_blocks(grid):
    """備援：沒有「第N題」標記時，把每個含「回答」欄的表頭列各當成一個子題區塊。

    真實匯出檔同一張 sheet 可能有 5 個表頭區塊；有「第N題」標記時走 parse_details，
    沒有標記時走這裡，確保每個區塊都能歸到一個子題而不是被吃掉。
    """
    hdrs = [i for i in header_rows(grid)
            if "回答" in grid[i] and not (len(grid[i]) > 2 and grid[i][2] == "作答時間")]
    out = []
    for k, i in enumerate(hdrs):
        end = hdrs[k + 1] if k + 1 < len(hdrs) else len(grid)
        hdr = grid[i]
        scol, ncol, acol = hdr.index("學號"), hdr.index("姓名"), hdr.index("回答")
        text = ""
        for j in range(i - 1, max(-1, i - 12), -1):
            if grid[j] and grid[j][0] == "問題敘述" and len(grid[j]) > 1:
                text = grid[j][1]
                break
        answers = []
        for j in range(i + 1, end):
            r = grid[j]
            if not r or scol >= len(r) or not r[scol] or r[scol] == "學號":
                break
            answers.append((r[scol], r[ncol] if ncol < len(r) else "",
                            r[acol] if acol < len(r) else "", j))
        if answers:
            out.append({"label": f"第{len(out) + 1}題", "text": text, "answers": answers})
    return out


# ------------------------------------------------------------------ A 區塊
def parse_summary(grid):
    """上方彙總表。回傳 dict(meta, subs, rows, rows_i, unanswered)。"""
    meta = {}
    for r in grid[:14]:
        if len(r) >= 2 and r[0] in META_KEYS:
            meta.setdefault(r[0], r[1])

    types_row = next((r for r in grid[:16] if r and r[0] == "問題類型"), None)
    desc_row = next((r for r in grid[:16] if r and r[0] == "問題敘述" and len(r) > 3), None)

    hdr_i = None
    for i, r in enumerate(grid):
        if len(r) >= 3 and r[0] == "學號" and r[1] == "姓名" and r[2] == "作答時間":
            hdr_i = i
            break
    if hdr_i is None:
        return {"meta": meta, "subs": [], "rows": [], "rows_i": [], "unanswered": []}
    hdr = grid[hdr_i]

    subs = []
    for ci in range(3, len(hdr)):
        if hdr[ci] not in ("回答", "總分") and not hdr[ci].startswith("第"):
            continue                      # 跳過 答對/答錯選項數量 等統計欄
        if hdr[ci] == "總分":
            continue
        label = types_row[ci] if (types_row and ci < len(types_row)) else ""
        text = desc_row[ci] if (desc_row and ci < len(desc_row)) else ""
        if not label and hdr[ci].startswith("第"):
            label = hdr[ci]
        if not label and hdr[ci] == "回答":
            label = "第1題"
        subs.append({"col": ci, "label": label or f"第{len(subs) + 1}題", "text": text})

    rows, rows_i = [], []
    i = hdr_i + 1
    while i < len(grid):
        r = grid[i]
        if not any(r) or (r and r[0] in ("未作答學生", "學號")) or (r and r[0].startswith("第")):
            break
        if r and r[0]:
            rows.append(r)
            rows_i.append(i)
        i += 1

    unanswered = []
    for j in range(i, len(grid)):
        if grid[j] and grid[j][0] == "未作答學生":
            k = j + 2                      # 跳過 學號|姓名 表頭
            while k < len(grid) and grid[k] and grid[k][0] and grid[k][0] != "學號":
                unanswered.append({"學號": grid[k][0],
                                   "姓名": grid[k][1] if len(grid[k]) > 1 else "",
                                   "列": k})
                k += 1
            break
    return {"meta": meta, "subs": subs, "rows": rows, "rows_i": rows_i,
            "unanswered": unanswered}


# ------------------------------------------------------------------ 通用 CSV
def parse_generic_csv(path):
    """其他平台匯出的 CSV：只要欄位名含「姓名」與「作答」就能吃。"""
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    if not rows:
        return None
    cols = list(rows[0].keys())

    def pick(*keys):
        for k in keys:
            for c in cols:
                if k in c:
                    return c
        return None

    c_name = pick("姓名", "名字", "name")
    c_ans = pick("作答", "回答", "答案", "answer")
    if not c_name or not c_ans:
        return None
    c_code = pick("學生編號")
    c_sid = pick("學號", "編號", "id")
    c_time = pick("時間", "time")
    c_sub = pick("子題", "題號", "題目")
    out = []
    for r in rows:
        out.append({
            "學生編號": (r.get(c_code) or "").strip() if c_code else "",
            "學號": (r.get(c_sid) or "").strip() if c_sid else "",
            "姓名": (r.get(c_name) or "").strip(),
            "作答時間": (r.get(c_time) or "").strip() if c_time else "",
            "子題號": (r.get(c_sub) or "第1題").strip() if c_sub else "第1題",
            "作答內容": (r.get(c_ans) or "").strip(),
        })
    return out


# ------------------------------------------------------------------ 主流程
def list_inputs(input_dir):
    """回傳要處理的檔案清單（xlsx 優先，其次 csv），依檔名排序。"""
    files = sorted(glob.glob(os.path.join(input_dir, "*.xlsx")))
    files = [f for f in files if not os.path.basename(f).startswith("~$")]
    files += sorted(glob.glob(os.path.join(input_dir, "*.csv")))
    return files


def _clean_code(c):
    """A 欄的值：只接受真正的學生編號字樣，其餘（匿名／空白）當作沒有編號。"""
    c = (c or "").strip()
    if not c or c == CODE_HEADER or c in NOT_A_NAME:
        return ""
    return c


def parse_file(path, no, semester="", course=""):
    """單一檔案 → (rows, rec)；rows 是已去識別化的 CSV 列，rec 是統計資訊。"""
    qno = f"Q{no:02d}"
    base = os.path.splitext(os.path.basename(path))[0]
    semester = normalize_semester(semester)
    course = normalize_course_code(course)

    # -------------------------------------------------- 通用 CSV
    if path.lower().endswith(".csv"):
        generic = parse_generic_csv(path)
        if generic is None:
            return [], {"題號": qno, "來源檔": os.path.basename(path),
                        "狀態": "失敗", "原因": "CSV 缺少「姓名」或「作答」欄"}
        coder = StudentCoder(semester, course)
        for g in generic:
            if not _clean_code(g["學生編號"]):
                g["學生編號"] = coder.code_for(g["學號"], g["姓名"])
        roster = build_roster([g["姓名"] for g in generic])
        rep = build_replacer(roster, coder.name_to_code())
        out_rows, n_masked = [], 0
        for g in generic:
            if not g["作答內容"]:
                continue
            m = mask_name(g["姓名"])
            n_masked += (m != g["姓名"])
            out_rows.append([qno, g["子題號"], "", g["學生編號"], mask_id(g["學號"]), m,
                             g["作答時間"], mask_in_text(g["作答內容"], rep)])
        keys = {(r[3] or r[4] or r[5]) for r in out_rows}
        rec = {"題號": qno, "來源檔": os.path.basename(path), "狀態": "成功",
               "題目": base, "題型": "通用CSV", "作答人數": len(keys),
               "未作答人數": 0, "全班人數": len(keys),
               "子題數": len({r[1] for r in out_rows}), "有效作答列數": len(out_rows),
               "取用區塊": "通用CSV", "學生編號來源": "程式產生",
               "姓名遮罩": "已套用(第2字改O)", "遮罩列數": n_masked,
               "子題": [{"label": s, "text": "", "作答數": sum(1 for r in out_rows if r[1] == s)}
                        for s in sorted({r[1] for r in out_rows})]}
        return out_rows, rec

    # -------------------------------------------------- xlsx
    raw = read_grid(path)
    codes, grid, has_codes = split_code_column(raw)

    if looks_like_roster_file(grid):
        # 名冊檔不是作業，但它的學生編號是「全班人數」最可靠的來源
        codes = sorted({_clean_code(c) for c in codes if _clean_code(c)},
                       key=lambda c: (len(c), c))
        return [], {"題號": "", "來源檔": os.path.basename(path), "狀態": "略過",
                    "原因": "非作業檔（看起來是名冊：只有學號／姓名，沒有回答欄）",
                    "名冊人數": len(codes), "學生編號清單": codes}

    p = parse_summary(grid)
    details = parse_details(grid)
    if not details and not p["subs"]:
        details = parse_answer_blocks(grid)
    is_quiz = looks_like_quiz(grid)
    qtype_meta = p["meta"].get("問題題型", "")
    title = clean_title(p["meta"].get("題幹") or p["meta"].get("資料夾名稱") or base,
                        qtype_meta) or clean_title(base)

    # 作答時間只在彙總表有 -> 用學號回查
    tmap = {r[0]: (r[2] if len(r) > 2 else "") for r in p["rows"] if r and r[0]}

    # ---- 學生編號：A 欄優先；沒有就自己編（向下相容原始 Zuvio 匯出）
    coder = StudentCoder(semester, course)
    if not has_codes:
        for r in p["rows"]:                       # 依首次出現順序
            coder.code_for(r[0], r[1] if len(r) > 1 else "")
        for u in p["unanswered"]:
            coder.code_for(u["學號"], u["姓名"])
        for d in details:
            for sid, name, _ans, _i in d["answers"]:
                coder.code_for(sid, name)

    def code_at(i, sid, name):
        c = _clean_code(codes[i]) if (has_codes and i < len(codes)) else ""
        return c or coder.code_for(sid, name)

    # ---- 名冊（只在記憶體，用來把自由文字裡的同學姓名換成學生編號）
    roster = build_roster(
        [r[1] for r in p["rows"] if len(r) > 1],
        [u["姓名"] for u in p["unanswered"]],
        [x[1] for d in details for x in d["answers"]],
    )
    rep = build_replacer(roster, coder.name_to_code())

    out_rows, n_masked = [], 0
    seen_codes = set()
    if details:
        source = "逐子題明細區塊(第N題)"
        for d in details:
            for sid, name, ans, ri in d["answers"]:
                code = code_at(ri, sid, name)
                seen_codes.add(code or sid or name)
                if not ans:
                    continue
                m = mask_name(name)
                n_masked += (m != name)
                out_rows.append([qno, d["label"], d["text"], code, mask_id(sid), m,
                                 tmap.get(sid, ""), mask_in_text(ans, rep)])
        subs_meta = [{"label": d["label"], "text": d["text"],
                      "作答數": sum(1 for x in d["answers"] if x[2])} for d in details]
    else:
        source = "彙總表"
        for r, ri in zip(p["rows"], p["rows_i"]):
            sid = r[0]
            name = r[1] if len(r) > 1 else ""
            ts = r[2] if len(r) > 2 else ""
            code = code_at(ri, sid, name)
            seen_codes.add(code or sid or name)
            m = mask_name(name)
            n_masked += (m != name)
            for s in p["subs"]:
                ans = r[s["col"]] if s["col"] < len(r) else ""
                if not ans:
                    continue
                out_rows.append([qno, s["label"], s["text"], code, mask_id(sid), m,
                                 ts, mask_in_text(ans, rep)])
        subs_meta = [{"label": s["label"], "text": s["text"],
                      "作答數": sum(1 for r in p["rows"]
                                 if s["col"] < len(r) and r[s["col"]])}
                     for s in p["subs"]]

    # 這個檔出現過的所有學生編號（作答＋未作答）—— 用來算「全班人數」的共用分母
    file_codes = {c for c in seen_codes if c and c not in NOT_A_NAME}
    for u in p["unanswered"]:
        c = code_at(u.get("列", -1), u["學號"], u["姓名"])
        if c:
            file_codes.add(c)

    anon = sum(1 for r in p["rows"] if r and r[0] in ("匿名作答者", "匿名"))
    answered = len({(r[3] or r[4] or r[5]) for r in out_rows if (r[3] or r[4] or r[5])})
    named = len(p["rows"]) - anon
    rec = {
        "題號": qno,
        "來源檔": os.path.basename(path),
        "狀態": "成功" if out_rows else "無有效作答",
        "題目": title,
        "題型": p["meta"].get("問題題型", "") or ("測驗題目" if is_quiz else ""),
        "測驗型": bool(is_quiz),
        "資料夾": p["meta"].get("資料夾名稱", ""),
        "作答人數": answered or len(p["rows"]),
        "實名作答": named,
        "匿名作答": anon,
        "未作答人數": len(p["unanswered"]),
        "全班人數": named + len(p["unanswered"]),
        "學生編號清單": sorted(file_codes),
        "子題數": len(subs_meta),
        "子題": subs_meta,
        "有效作答列數": len(out_rows),
        "取用區塊": source,
        "學生編號來源": "A 欄（姓名遮罩與學生編號程式）" if has_codes else "程式產生（原始匯出檔）",
        "學生編號數": len({r[3] for r in out_rows if r[3]}),
        "姓名遮罩": "已套用(第2字改O)",
        "遮罩列數": n_masked,
    }
    return out_rows, rec


def parse_dir(input_dir, out_dir, semester="", course="", log=print):
    """把 input_dir 內所有檔案轉成 out_dir 內的 Q0N_*.csv，回傳 index 清單。"""
    os.makedirs(out_dir, exist_ok=True)
    files = list_inputs(input_dir)
    if not files:
        raise SystemExit(f"[錯誤] {input_dir} 裡找不到任何 .xlsx 或 .csv。\n"
                         f"       請先跑「姓名遮罩與學生編號」，把輸出的 xlsx 放進這個資料夾。")
    index = []
    skipped = 0
    no = 0
    union_codes = set()          # 本週所有輸入檔出現過的不重複學生編號
    roster_n = 0                 # 名冊檔的人數（若有名冊，以它為準）
    for path in files:
        no += 1
        rows, rec = parse_file(path, no, semester=semester, course=course)
        union_codes |= set(rec.get("學生編號清單") or [])
        if rec.get("狀態") == "略過":
            no -= 1
            skipped += 1
            roster_n = max(roster_n, int(rec.get("名冊人數") or 0))
            log(f"  [略過] {os.path.basename(path)}　{rec.get('原因', '非作業檔')}"
                + (f"（名冊 {rec['名冊人數']} 人，拿來當全班人數）" if rec.get("名冊人數") else ""))
            continue
        fn = f"{rec['題號']}_{safe_name(rec.get('題目', ''))}.csv"
        rec["檔名"] = fn
        with open(os.path.join(out_dir, fn), "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(CSV_COLS)
            w.writerows(rows)
        index.append(rec)
        log(f"  {rec['題號']}  {rec.get('題目', '')[:24]} -> {fn}"
            f"（作答 {rec.get('作答人數', 0)} 人／子題 {rec.get('子題數', 0)}／"
            f"有效 {rec.get('有效作答列數', 0)} 列／學生編號 {rec.get('學生編號數', 0)} 個"
            f"／{rec.get('學生編號來源', '')}）")
    if skipped:
        log(f"  共略過 {skipped} 個非作業檔（名冊等），實際分析 {len(index)} 題。")
    if not index:
        raise SystemExit(f"[錯誤] {input_dir} 裡的檔案都不是作業匯出檔（全部被略過）。\n"
                         f"       請確認放進來的是 Zuvio「下載數據」或姓名遮罩程式的輸出。")

    # ---- 全班人數：整份輸入資料夾共用同一個分母，不要每題各算各的
    klass = roster_n or len(union_codes)
    src = "名冊檔" if roster_n else "本週所有輸入檔的不重複學生編號"
    for rec in index:
        rec["全班人數"] = klass or rec.get("全班人數", 0)
        rec["全班人數來源"] = src
        rec.pop("學生編號清單", None)      # 只是中間結果，不寫進 index json
    log(f"  全班人數統一為 {klass} 人（來源：{src}）；所有題頁共用這個分母。")
    return index


def roster_from_dir(input_dir):
    """從輸入檔重建名冊（僅供 --check-privacy 在本機比對，不寫出任何檔案）。

    注意：NameMasker 2.0 的輸出檔裡姓名已經是「王O明」，不是純中文，
    build_roster 會自動排除，所以名冊會是空的 —— 這是正確的，代表沒有真名可外洩。
    """
    names = []
    for path in list_inputs(input_dir):
        if path.lower().endswith(".csv"):
            g = parse_generic_csv(path)
            if g:
                names += [x["姓名"] for x in g]
            continue
        try:
            raw = read_grid(path)
        except Exception:
            continue
        _codes, grid, _has = split_code_column(raw)
        p = parse_summary(grid)
        names += [r[1] for r in p["rows"] if len(r) > 1]
        names += [u["姓名"] for u in p["unanswered"]]
        for d in parse_details(grid):
            names += [x[1] for x in d["answers"]]
    return build_roster(names)
