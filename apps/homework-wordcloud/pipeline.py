# -*- coding: utf-8 -*-
"""
pipeline.py — 「學生作業文字雲」流程核心（可被 GUI 或 CLI 呼叫）。

原始版本是命令列工具 ta_wordcloud.py；這裡把流程改寫成純函式：
  * 所有訊息都走 log callback（GUI 可以即時顯示在文字框），不直接 print。
  * 出錯丟 PipelineError，不呼叫 sys.exit（GUI 不能被關掉）。
  * 所有資料檔（jieba 字典、sample_input、logo、config.json）一律用 resource_path()
    定位，PyInstaller onefile 打包後會落在 sys._MEIPASS。

兩步流程（2.1）：
  ① 先用「姓名遮罩與學生編號」（NameMasker 2.1）處理 Zuvio 匯出的 xlsx
  ② 把它的輸出放進 input 資料夾再跑本程式
  （沒有先跑第 ① 步也可以，本程式會用 core/mask.py 自己補上同一套規則；
    但這時**一定要先給「原始名單或學生編號對照表」**，學生編號才會整學期固定。）

2.1 的關鍵改變：
  * 輸入若已是 NameMasker 輸出（A 欄有「學生編號」）→ 照舊，不需要名單。
  * 輸入若是**原始 Zuvio 匯出檔** → 必須提供原始名單（Excel／CSV／Word／東華選課名單 PDF）
    或既有的學生編號對照表；程式會產生／沿用
    `{學期}_{課程縮寫}_學生名單與學生編號對照表.xlsx`，
    名單外的作答者自 101 起持久登記並回寫。
  * 沒有名單而硬要跑 → 要明確選「沒有名單，改依檔案內出現順序編號」（不建議）。

隱私鐵律（不可關閉）：
  * 學生一律用學生編號表示（例 115-1_EC_3）。
  * 姓名遮罩（第 2 字改 O），學號與電子郵件每個字元都改成 O，長度不變。
  * 自由文字裡的名冊姓名換成該生的學生編號。
  * 「學生編號連結姓名」工作表（再識別鑰匙）解析時一律忽略，絕不讀入。
  * 原始 xlsx 只留在使用者自己的輸入資料夾，工具不會上傳任何資料。
"""
import os
import sys
import json
import shutil
import tempfile
import traceback
import datetime as dt

APP_NAME = "HomeworkWordCloud"
APP_TITLE = "學生作業文字雲"
VERSION = "2.3.0"

# --check-privacy 掃描時要看的純文字副檔名
TEXT_EXT = {".csv", ".json", ".md", ".txt"}


class PipelineError(Exception):
    """使用者看得懂的錯誤（GUI 會直接顯示訊息，不印 traceback）。"""


# ------------------------------------------------------------------ 路徑
def is_frozen():
    return bool(getattr(sys, "frozen", False))


def resource_path(*parts):
    """打包後的唯讀資料檔位置（onefile 解壓在 sys._MEIPASS）。"""
    base = getattr(sys, "_MEIPASS", None) or os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, *parts)


def app_home():
    """使用者設定檔存放處（不放在 exe 旁邊，避免 Program Files 沒有寫入權限）。"""
    p = os.path.join(os.path.expanduser("~"), ".homework_wordcloud")
    os.makedirs(p, exist_ok=True)
    return p


def default_output_root():
    desktop = os.path.join(os.path.expanduser("~"), "Desktop")
    if not os.path.isdir(desktop):
        desktop = os.path.expanduser("~")
    return os.path.join(desktop, f"{APP_NAME}_輸出")


def open_folder(path, log=print):
    """用作業系統的檔案總管打開資料夾。"""
    path = os.path.abspath(path)
    if not os.path.isdir(path):
        raise PipelineError(f"資料夾還不存在：{path}")
    try:
        if sys.platform.startswith("win"):
            os.startfile(path)                                 # noqa: S606
        elif sys.platform == "darwin":
            import subprocess
            subprocess.Popen(["open", path])
        else:
            import subprocess
            subprocess.Popen(["xdg-open", path])
    except Exception as e:                                     # pragma: no cover
        raise PipelineError(f"打不開資料夾：{e}")
    log(f"已開啟資料夾：{path}")


# ------------------------------------------------------------------ 設定
DEFAULTS = {
    "course_name": "我的課程", "teacher": "", "ta_name": "",
    "course_code": "EC",              # 課程縮寫（2 個英文大寫字母），學生編號用
    "semester": "115-1",              # 學期（例 115-1），學生編號用
    "semester_start": "", "logo_path": "", "stopwords_extra": [],
    "keep_short": [], "user_words": [], "deck_series": "",
    "synonyms": {},                   # {正式概念: [同義詞, ...]}
    "question_types": {},             # {類別名: [關鍵詞, ...]}，覆寫預設四類（順序即優先序）
    "appendix_matrix": True,          # 要不要做附錄「概念矩陣」頁
    "top_concepts": 6,                # 每題抓幾個重點概念（2.2 起預設 6）
    "min_words_for_cloud": 5, "input_dir": "", "output_dir": "",
    "codebook_path": "",              # 2.1：上次用的「原始名單或學生編號對照表」
    "config_version": 0.0,            # 2.2：設定檔格式版本，用來做一次性遷移
    # ---- 2.3 斷詞品質
    "min_word_len": 2,                # 關鍵字最少幾個中文字
    "max_word_len": 6,                # 關鍵字最多幾個中文字（白名單不受限）
    "keep_pos": [],                   # 保留的詞性；[] = 用內建 KEEP_POS
    "stopwords_remove": [],           # 要「拿掉」的內建停用詞（例如想分析「同學」）
}

CONFIG_VERSION = 2.3                  # 2–6 字關鍵字＋詞性過濾＋停用詞擴充
_V23_KEYS = ("min_word_len", "max_word_len", "keep_pos", "stopwords_remove")


def _migrate_user_cfg(cfg, user_raw, log=None):
    """一次性遷移：舊的使用者設定檔會把新版預設蓋回舊值。

    舊版的 settings.json 存的是當時的預設值，載入順序又排在打包的 config.json
    後面，會把新預設整個蓋掉。因此凡是「看得出來是舊版留下來的預設值」就換成
    新預設；使用者自己改過的值（例如刻意設 top_concepts=4）保留不動。

    2.2：3 個重點 → 6 個重點、兩類提問 → 四類提問。
    2.3：補上 min_word_len／max_word_len／keep_pos／stopwords_remove 四個新鍵。
    """
    ver = float(user_raw.get("config_version") or 0)
    if ver >= CONFIG_VERSION:
        return cfg
    if ver < 2.2:
        if int(cfg.get("top_concepts") or 0) == 3:      # 2.1 的預設值
            cfg["top_concepts"] = DEFAULTS["top_concepts"]
            if log:
                log("  [設定升級] 每題重點數 3 → 6（2.2 預設）")
        qt = cfg.get("question_types") or {}
        if qt and len(qt) < 4:                          # 2.1 只有兩類
            cfg["question_types"] = {}                  # 清空 → 用 2.2 的四類預設
            if log:
                log("  [設定升級] 提問分類 兩類 → 四類（2.2 預設）")
    if ver < 2.3:
        missing = [k for k in _V23_KEYS if k not in user_raw]
        if missing:
            for k in missing:
                cfg[k] = DEFAULTS[k]
            if log:
                log("  [設定升級] 關鍵字長度 2–6 字＋詞性過濾＋停用詞擴充（2.3 預設）")
    cfg["config_version"] = CONFIG_VERSION
    return cfg


def load_config(path=None, log=None):
    """讀設定：先讀打包進來的 config.json 範本，再讓使用者設定檔覆寫。"""
    cfg = dict(DEFAULTS)
    user_p = path or user_config_path()
    for p in [resource_path("config.json"), user_p]:
        if p and os.path.exists(p):
            try:
                with open(p, encoding="utf-8") as f:
                    cfg.update({k: v for k, v in json.load(f).items() if k in DEFAULTS})
            except Exception:
                pass
    user_raw = {}
    if user_p and os.path.exists(user_p):
        try:
            with open(user_p, encoding="utf-8") as f:
                user_raw = json.load(f)
        except Exception:
            user_raw = {}
    cfg = _migrate_user_cfg(cfg, user_raw, log=log)
    # 範本裡的相對路徑（input/output/對照表）在桌面版沒有意義，交給 GUI 自己決定
    for k in ("input_dir", "output_dir", "codebook_path"):
        if cfg.get(k) and not os.path.isabs(cfg[k]):
            cfg[k] = ""
    return cfg


def user_config_path():
    return os.path.join(app_home(), "settings.json")


def save_config(cfg):
    data = {k: cfg.get(k, v) for k, v in DEFAULTS.items()}
    data["config_version"] = CONFIG_VERSION          # 存檔就代表已是最新格式
    with open(user_config_path(), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    return user_config_path()


def norm_course_code(s):
    from core.mask import normalize_course_code
    return normalize_course_code(s)


def norm_semester(s):
    from core.mask import normalize_semester
    return normalize_semester(s)


def logo_file(cfg):
    p = (cfg.get("logo_path") or "").strip()
    if p and os.path.isabs(p) and os.path.exists(p):
        return p
    bundled = resource_path("assets", "logo_placeholder.png")
    return bundled if os.path.exists(bundled) else ""


def sample_input_dir():
    return resource_path("sample_input")


# ------------------------------------------------------------------ 名單／對照表（2.1）
NO_ROSTER_HINT = (
    "這批輸入檔是**原始 Zuvio 匯出檔**（A 欄沒有「學生編號」），"
    "必須先有名單才能處理：\n"
    "  {files}\n\n"
    "請在「原始名單或學生編號對照表」欄選一份檔案：\n"
    "  ・原始名單：Excel／CSV／Word（要有「學號」「姓名」欄），或東華教務系統的選課名單 PDF\n"
    "  ・或既有的「115-1_EC_學生名單與學生編號對照表.xlsx」\n\n"
    "為什麼一定要？沒有名單就只能依「檔案內出現順序」編號，"
    "同一位同學在不同題目會拿到不同編號，沒辦法跨題跨週追蹤。\n"
    "真的沒有名單時，請勾選「沒有名單，改依檔案內出現順序編號」"
    "（CLI：加上 --no-roster）。")


def load_codebook_for(cfg, src_path="", overwrite=False, log=print):
    """讀 TA 給的「原始名單或學生編號對照表」→ Codebook（丟 PipelineError）。"""
    from core import roster as R
    src = (src_path or cfg.get("codebook_path") or "").strip()
    if not src:
        return None
    if not os.path.exists(src):
        raise PipelineError(f"找不到名單／對照表：{src}\n請重新選一次檔案。")
    try:
        return R.prepare_codebook(src, cfg.get("semester"), cfg.get("course_code"),
                                  overwrite=overwrite, log=log)
    except R.RosterError as e:
        raise PipelineError(str(e))


def describe_codebook(cb):
    from core import roster as R
    return R.describe(cb)


def make_codebook_only(cfg, src_path, overwrite=False, log=print):
    """只產生／更新對照表，不處理作業檔（GUI 的「產生／更新對照表」按鈕）。"""
    cb = load_codebook_for(cfg, src_path, overwrite=overwrite, log=log)
    if cb is None:
        raise PipelineError("請先選一份「原始名單或學生編號對照表」。")
    if not cb.path:
        raise PipelineError("對照表沒有存檔路徑，請改選有寫入權限的資料夾。")
    log("")
    log("=" * 60)
    log(f"對照表：{cb.path}")
    log(f"  {describe_codebook(cb)}")
    log("  ⚠ 這份檔案含真實姓名與學號（再識別鑰匙），只留本機，不要上傳或分享。")
    log("=" * 60)
    return cb


# ------------------------------------------------------------------ jieba
_JIEBA_READY = False


def init_jieba(log=print):
    """打包後 jieba 找不到自己的 dict.txt，必須明確指定路徑。"""
    global _JIEBA_READY
    if _JIEBA_READY:
        return
    import logging
    import jieba
    jieba.setLogLevel(logging.WARN)
    d = resource_path("jieba", "dict.txt")
    if os.path.exists(d):
        try:
            jieba.set_dictionary(d)
        except Exception as e:
            log(f"  [提醒] 指定 jieba 字典失敗，改用預設字典：{e}")
    jieba.dt.tmp_dir = tempfile.gettempdir()
    _JIEBA_READY = True


# ------------------------------------------------------------------ 週次
def week_start(today=None):
    """回傳 today 所在那一週的星期日（週期＝週日～週六）。"""
    today = today or dt.date.today()
    return today - dt.timedelta(days=(today.weekday() + 1) % 7)


def parse_semester_start(s):
    s = (s or "").strip().replace("/", "-")
    if not s:
        raise PipelineError("「學期起日」是空的，無法自動判週。\n"
                            "請填第 1 週的星期日（例如 2026-09-06），或改用「指定週次」。")
    try:
        d = dt.date.fromisoformat(s)
    except ValueError:
        raise PipelineError(f"學期起日格式看不懂：{s}\n請用 YYYY-MM-DD，例如 2026-09-06。")
    return d - dt.timedelta(days=(d.weekday() + 1) % 7)      # 校正成星期日


def week_number(semester_start, d=None):
    s = parse_semester_start(semester_start)
    return (week_start(d) - s).days // 7 + 1


def week_folder(n, semester_start, d=None):
    """Wn-MMDD-MMDD（週日～週六）。"""
    if (semester_start or "").strip():
        sun = parse_semester_start(semester_start) + dt.timedelta(weeks=(n - 1))
    else:
        sun = week_start(d)
    sat = sun + dt.timedelta(days=6)
    return f"W{n}-{sun:%m%d}-{sat:%m%d}", sun, sat


# ------------------------------------------------------------------ 主流程
def _core():
    from core import parse_zuvio as PZ
    from core import analyze as AN
    from core import wordcloud_gen as WG
    from core import deck as DK
    return PZ, AN, WG, DK


def can_make_thumbs():
    """縮圖靠 PowerPoint COM：只有 Windows 而且裝了 PowerPoint 才做得到。"""
    if not sys.platform.startswith("win"):
        return False
    try:
        import win32com.client                                # noqa: F401
    except Exception:
        return False
    return True


def run_pipeline(cfg, input_dir, out_dir, week_label, deck_name,
                 thumbs=True, log=print, codebook=None, no_roster=False):
    """完整跑一次：解析 → 統計 → 文字雲 → 簡報 → QA。回傳結果摘要 dict。

    codebook   固定學生編號對照表（core.roster.Codebook）；原始 Zuvio 檔必須有它。
    no_roster  使用者明確選「沒有名單，改依檔案內出現順序編號」（不建議）。
    """
    PZ, AN, WG, DK = _core()
    input_dir = os.path.abspath(input_dir)
    out_dir = os.path.abspath(out_dir)
    if not os.path.isdir(input_dir):
        raise PipelineError(f"找不到輸入資料夾：{input_dir}")
    if not PZ.list_inputs(input_dir):
        raise PipelineError(
            f"輸入資料夾裡沒有任何 .xlsx 或 .csv：\n  {input_dir}\n\n"
            "請先到 Zuvio 老師端「下載數據」，把 xlsx 放進這個資料夾再試一次。")

    # ---- 2.1 §1.5：原始 Zuvio 檔一定要先有名單／對照表
    _files, raw_files = PZ.scan_inputs(input_dir)
    if raw_files and codebook is None and not no_roster:
        names = "\n  ".join(os.path.basename(p) for p in raw_files[:8])
        if len(raw_files) > 8:
            names += f"\n  …等共 {len(raw_files)} 個檔"
        raise PipelineError(NO_ROSTER_HINT.format(files=names))
    if raw_files and codebook is None and no_roster:
        log("  [注意] 沒有名單，改依檔案內出現順序編號："
            "同一位同學在不同題目會拿到不同編號，只能看單題，不能跨題跨週比對。")

    os.makedirs(out_dir, exist_ok=True)
    init_jieba(log)
    AN.load_dict(cfg.get("user_words"), cfg.get("stopwords_extra"), cfg.get("keep_short"),
                 synonyms=cfg.get("synonyms"), question_types=cfg.get("question_types"),
                 min_word_len=cfg.get("min_word_len"), max_word_len=cfg.get("max_word_len"),
                 keep_pos=cfg.get("keep_pos"), stopwords_remove=cfg.get("stopwords_remove"))
    log(f"  關鍵字規則：中文 {AN.MIN_LEN}–{AN.MAX_LEN} 字、英文 ≤{AN.MAX_EN_WORD_LEN} 字母，"
        f"詞性保留 {len(AN.POS_KEEP)} 類，停用詞 {len(AN.STOPWORDS)} 個，"
        f"白名單 {len(AN.ALWAYS_KEEP)} 個（不受長度與詞性限制）")

    from core.mask import normalize_course_code, normalize_semester
    course_code = normalize_course_code(cfg.get("course_code"))
    semester = normalize_semester(cfg.get("semester"))

    log("")
    log("[1/5] 解析匯出檔（學生編號優先）並完成去識別化")
    log(f"  學生編號格式：{semester}_{course_code}_n"
        f"（輸入檔若已有 A 欄「學生編號」就直接沿用）")
    index = PZ.parse_dir(input_dir, out_dir, semester=semester, course=course_code,
                         log=log, codebook=codebook)
    if codebook is not None:
        codebook.save_if_dirty(log=log)
    with open(os.path.join(out_dir, "questions_index.json"), "w", encoding="utf-8") as f:
        json.dump(index, f, ensure_ascii=False, indent=2)

    log("")
    log("[2/5] 斷詞、六個重點／概念矩陣、四類提問分類")
    result = AN.analyse_dir(out_dir, index=index,
                            top_concepts=int(cfg.get("top_concepts", 6) or 6), log=log)
    AN.write_outputs(result, out_dir, cfg.get("course_name", ""), week_label)

    log("")
    log("[3/5] 產生文字雲")
    wc_map = WG.generate_all(out_dir, min_words=int(cfg.get("min_words_for_cloud", 5) or 5),
                             log=log)

    log("")
    log("[4/5] 組裝簡報")
    meta = {
        "course_name": cfg.get("course_name", ""),
        "teacher": cfg.get("teacher", ""),
        "ta_name": cfg.get("ta_name", ""),
        "week_label": week_label,
        "date": dt.date.today().strftime("%Y/%m/%d"),
        "series": cfg.get("deck_series") or f"{cfg.get('course_name', '')}　同學作答分析",
        "logo": logo_file(cfg),
        "appendix_matrix": bool(cfg.get("appendix_matrix", True)),
    }
    pptx_path = os.path.join(out_dir, deck_name)
    prs, total = DK.build_deck(result, pptx_path, meta, wc_map, log=log)

    log("")
    log("[5/5] 版面 QA")
    qa_dir = os.path.join(out_dir, "qa")
    rep = DK.layout_scan(prs, os.path.join(qa_dir, "layout_report.json"))
    log(f"  幾何掃描：出界 {len(rep['out_of_bounds'])} 處、"
        f"重疊>8% {len(rep['overlaps'])} 處、刻意出血裝飾 {len(rep['刻意出血裝飾'])} 個")
    for o in rep["out_of_bounds"][:6]:
        log(f"    出界：{o}")
    for o in rep["overlaps"][:6]:
        log(f"    重疊：{o}")

    if thumbs:
        if can_make_thumbs():
            DK.export_thumbs(pptx_path, qa_dir, log=log)
        else:
            log("  [略過] 這台電腦沒有 PowerPoint（或非 Windows），不產生縮圖，不影響簡報。")

    with open(os.path.join(out_dir, "run_meta.json"), "w", encoding="utf-8") as f:
        json.dump({
            "產生時間": dt.datetime.now().isoformat(timespec="seconds"),
            "工具": f"{APP_TITLE} {VERSION}",
            "課程": cfg.get("course_name", ""), "週次標籤": week_label,
            # 只記資料夾名稱，不記完整本機路徑（輸出可能會分享出去，
            # 而且路徑裡的數字會被隱私掃描誤判）
            "輸入資料夾": os.path.basename(os.path.normpath(input_dir)),
            "輸出資料夾": os.path.basename(os.path.normpath(out_dir)),
            "簡報": deck_name, "頁數": total, "題數": len(result["題目"]),
            "學生編號格式": f"{semester}_{course_code}_n",
            "學生編號來源": ("固定對照表（名單 "
                       f"{codebook.count} 人、名單外 {len(codebook.outside)} 人）"
                       if codebook is not None else
                       ("依檔案內出現順序（未用名單）" if raw_files else
                        "輸入檔 A 欄（姓名遮罩與學生編號程式）")),
            "去識別化": "姓名第 2 字改 O；學號與 email 全部改 O；主鍵用學生編號",
            "幾何QA": {"出界": len(rep["out_of_bounds"]), "重疊": len(rep["overlaps"])},
        }, f, ensure_ascii=False, indent=2)

    log("")
    log("=" * 60)
    log(f"完成！輸出資料夾：{out_dir}")
    log(f"  簡報：{deck_name}（{total} 頁，每頁備忘稿都有逐字稿）")
    log("  摘要：summary.md　資料：analysis.json、Q*.csv")
    log("  每題兩份分析表：Q0N_concept_matrix.csv（概念矩陣）、Q0N_questions.csv（提問分類）")
    log("=" * 60)
    return {"out_dir": out_dir, "pptx": pptx_path, "pages": total,
            "questions": len(result["題目"]), "wordclouds": len(wc_map),
            "analysis": result}


def run_demo(cfg, out_root, thumbs=True, log=print):
    """用內建合成資料（虛構姓名）跑一次示範。"""
    out_dir = os.path.join(os.path.abspath(out_root), "示範")
    if os.path.isdir(out_dir):
        for fn in os.listdir(out_dir):
            p = os.path.join(out_dir, fn)
            try:
                shutil.rmtree(p) if os.path.isdir(p) else os.remove(p)
            except OSError:
                pass
    log("=" * 60)
    log("示範模式：使用內建合成資料（虛構姓名，不是真實學生）")
    log("=" * 60)
    demo_cfg = dict(cfg)
    demo_cfg.setdefault("course_name", "我的課程")
    # 內建示範檔已經是「姓名遮罩與學生編號」的輸出（A 欄有學生編號），不需要名單。
    return run_pipeline(demo_cfg, sample_input_dir(), out_dir,
                        "示範週", "同學作答分析-示範.pptx", thumbs=thumbs, log=log)


def pick_input_dir(input_dir, wfolder):
    """輸入資料夾底下若有 Wn-MMDD-MMDD 子資料夾就優先用它。"""
    base = os.path.abspath(input_dir)
    sub = os.path.join(base, wfolder)
    try:
        from core import parse_zuvio as PZ
        if os.path.isdir(sub) and PZ.list_inputs(sub):
            return sub
    except Exception:
        pass
    return base


def run_week(cfg, input_dir, out_root, week_no=None, thumbs=True, log=print,
             codebook=None, no_roster=False):
    """跑某一週（week_no=None 代表依學期起日自動判斷本週）。"""
    from core import deck as DK
    n = int(week_no) if week_no else week_number(cfg.get("semester_start"))
    if n < 1:
        raise PipelineError(f"算出來的週次是第 {n} 週（今天早於學期起日）。\n"
                            "請確認「學期起日」，或改用「指定週次」。")
    wfolder, sun, sat = week_folder(n, cfg.get("semester_start"))
    week_label = f"第{DK.cn_week(n)}週（{sun:%Y/%m/%d}–{sat:%m/%d}）"
    deck_name = f"同學作答分析-第{DK.cn_week(n)}周.pptx"
    in_dir = pick_input_dir(input_dir, wfolder)
    out_dir = os.path.join(os.path.abspath(out_root), wfolder)

    log("=" * 60)
    log(f"課程：{cfg.get('course_name', '')}　週次：{week_label}")
    log(f"輸入：{in_dir}")
    log(f"輸出：{out_dir}")
    log("=" * 60)
    r = run_pipeline(cfg, in_dir, out_dir, week_label, deck_name, thumbs=thumbs,
                     log=log, codebook=codebook, no_roster=no_roster)
    r["week"] = n
    return r


# ------------------------------------------------------------------ 隱私掃描
def pptx_texts(path):
    from pptx import Presentation
    out = []
    prs = Presentation(path)
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append((f"slide{i}", sh.text_frame.text))
        if s.has_notes_slide:
            t = s.notes_slide.notes_text_frame.text
            if t.strip():
                out.append((f"slide{i}/notes", t))
    return out


def check_privacy(target, input_dir=None, log=print, require_roster=False,
                  codebook=None):
    """掃描輸出資料夾，看有沒有漏掉的個資。回傳命中筆數（0 才算通過）。

    比對四件事：
      1. 名冊上的真實姓名（從原始輸入檔重建，只留在記憶體）
      2. **對照表裡的所有真實姓名與學號**（2.1 新增；有給對照表才比）
      3. **9 碼學號**（`(?<!\\d)\\d{9}(?!\\d)`）
      4. **電子郵件樣式**
    學生編號（例 115-1_EC_3）是刻意要出現的，不算命中。
    """
    from core import parse_zuvio as PZ
    from core.mask import SID9_IN_TEXT, EMAIL, STUDENT_CODE_IN_TEXT
    target = os.path.abspath(target)
    if not os.path.isdir(target):
        raise PipelineError(f"找不到要掃描的資料夾：{target}")

    if not input_dir:
        meta_p = os.path.join(target, "run_meta.json")
        if os.path.exists(meta_p):
            try:
                with open(meta_p, encoding="utf-8") as f:
                    cand = json.load(f).get("input_dir") or ""
                if cand and os.path.isdir(cand):
                    input_dir = cand
            except Exception:
                pass
    reals = []
    if input_dir and os.path.isdir(input_dir):
        reals = [r for r, _ in PZ.roster_from_dir(input_dir)]
        log(f"名冊比對來源：{input_dir}（{len(reals)} 位真實姓名）")
        if not reals:
            log("  （輸入檔的姓名已經是遮罩版，名冊自然是空的 —— 這是正確的。）")
    elif require_roster:
        raise PipelineError(
            "找不到原始輸入資料夾，無法比對名冊。\n"
            "請先在上方把「輸入資料夾」指到當初那批 xlsx，再按一次「隱私檢查」。\n"
            "（提示：也可以直接選當週的輸出資料夾，例如 W3-0920-0926）")
    else:
        log("名冊比對來源：無（只做 9 碼學號與電子郵件樣式掃描）")

    cb_names, cb_sids = [], []
    if codebook is not None:
        cb_names = sorted({n for n in codebook.names() if n and len(n) >= 2})
        cb_sids = sorted({s for s in codebook.sids() if s})
        log(f"對照表比對來源：{os.path.basename(codebook.path or '（記憶體）')}"
            f"（{len(cb_names)} 個真實姓名、{len(cb_sids)} 個學號）")

    log(f"掃描目標：{target}")
    log("允許出現：學生編號（例 115-1_EC_3）")

    hits = []

    def scan(where, txt):
        for name in reals:
            if name and name in txt:
                hits.append({"檔案": where, "類型": "名冊姓名", "命中": name})
        for name in cb_names:
            if name in txt:
                hits.append({"檔案": where, "類型": "對照表姓名", "命中": name})
        for sid in cb_sids:
            if sid in txt:
                hits.append({"檔案": where, "類型": "對照表學號", "命中": sid})
        # 學生編號本身不含 9 碼連續數字，先挖掉以策安全
        clean_txt = STUDENT_CODE_IN_TEXT.sub(" ", txt)
        for m in SID9_IN_TEXT.finditer(clean_txt):
            hits.append({"檔案": where, "類型": "9 碼學號", "命中": m.group(0)})
        for m in EMAIL.finditer(clean_txt):
            hits.append({"檔案": where, "類型": "電子郵件", "命中": m.group(0)})

    for root, _dirs, files in os.walk(target):
        for fn in files:
            p = os.path.join(root, fn)
            ext = os.path.splitext(fn)[1].lower()
            rel = os.path.relpath(p, target)
            chunks = []
            if ext in TEXT_EXT:
                try:
                    with open(p, encoding="utf-8-sig", errors="replace") as f:
                        chunks = [(rel, f.read())]
                except Exception as e:
                    log(f"  [略過] 讀不到 {rel}：{e}")
            elif ext == ".pptx":
                try:
                    chunks = [(f"{rel}#{w}", t) for w, t in pptx_texts(p)]
                except Exception as e:
                    log(f"  [略過] 讀不到 {rel}：{e}")
            for where, txt in chunks:
                scan(where, txt)

    log("")
    log("=" * 60)
    if hits:
        kinds = {}
        for h in hits:
            kinds[h["類型"]] = kinds.get(h["類型"], 0) + 1
        log(f"[不通過] 發現 {len(hits)} 處可能的個資（"
            + "、".join(f"{k} {v} 處" for k, v in kinds.items()) + "）：")
        for h in hits[:40]:
            log(f"  - {h['檔案']}　{h['類型']}：「{h['命中']}」")
        log("請先不要把輸出交出去，並回報這個問題。")
    else:
        log(f"[通過] 0 命中。比對 {len(reals)} 個名冊姓名、"
            f"{len(cb_names)} 個對照表姓名、{len(cb_sids)} 個對照表學號、"
            "9 碼學號與電子郵件樣式，輸出資料夾內都沒有出現。")
    log("=" * 60)
    return len(hits)


# ------------------------------------------------------------------ 自我測試
def _link_sheet_leak(input_dir, out_dir, log=print):
    """驗證「學生編號連結姓名」工作表的姓名／學號／email 完全沒有進到輸出。"""
    import openpyxl
    from core.parse_zuvio import LINK_SHEET_NAMES
    secrets = set()
    for fn in os.listdir(input_dir):
        if not fn.lower().endswith(".xlsx"):
            continue
        try:
            wb = openpyxl.load_workbook(os.path.join(input_dir, fn), read_only=True,
                                        data_only=True)
        except Exception:
            continue
        for nm in wb.sheetnames:
            if nm.strip() not in LINK_SHEET_NAMES:
                continue
            for row in wb[nm].iter_rows(min_row=2, values_only=True):
                for c in (row[1:4] if row else []):
                    c = str(c or "").strip()
                    if len(c) >= 2 and c not in ("姓名", "學號（原）", "電子郵件（原）"):
                        secrets.add(c)
        wb.close()
    if not secrets:
        log("  連結表檢查：輸入檔沒有「學生編號連結姓名」工作表（略過）")
        return 0
    n = 0
    for root, _d, files in os.walk(out_dir):
        for fn in files:
            p = os.path.join(root, fn)
            ext = os.path.splitext(fn)[1].lower()
            texts = []
            if ext in TEXT_EXT:
                try:
                    with open(p, encoding="utf-8-sig", errors="replace") as f:
                        texts = [f.read()]
                except Exception:
                    pass
            elif ext == ".pptx":
                try:
                    texts = [t for _w, t in pptx_texts(p)]
                except Exception:
                    pass
            for t in texts:
                for sec in secrets:
                    if sec in t:
                        n += 1
                        log(f"    [外洩] {os.path.relpath(p, out_dir)} 出現「{sec}」")
    log(f"  連結表檢查：{len(secrets)} 個機敏字串，輸出命中 {n} 處")
    return n


def _selftest_questions(log):
    """提問偵測與**四類**分類的規則驗證（2.2）。

    規格：每一類至少 3 個問句，外加陳述句干擾；「其他」要 ≤ 20%。
    """
    from core import analyze as AN
    cases = [
        # ---- 陳述句干擾（含疑問詞但不是提問）
        ("了解物質是由原子組成，以及原子如何組成不同物質", False, None),
        ("理解物質的微觀結構如何決定性質", False, None),
        ("我知道要怎麼做實驗了", False, None),
        ("這次實驗的步驟我都記下來了。", False, None),
        ("原子核帶正電，電子帶負電。", False, None),
        ("原子是什麼", False, None),                  # 沒問號、疑問詞不在句首
        ("好想知道更多", False, None),
        # ---- 計算數據類（≥3）
        ("這個公式要怎麼算？", True, "計算數據類"),
        ("濃度的單位是什麼？", True, "計算數據類"),
        ("莫耳數要怎麼換算？", True, "計算數據類"),
        ("有效數字要取到第幾位？", True, "計算數據類"),
        # ---- 操作應用類（≥3）
        ("這個實驗的步驟要怎麼做？", True, "操作應用類"),
        ("那要用什麼儀器？", True, "操作應用類"),
        ("要怎麼觀察這個現象？", True, "操作應用類"),
        ("How do we measure it?", True, "操作應用類"),
        # ---- 延伸探究類（≥3）
        ("如果換成替代溶劑會不會比較好？", True, "延伸探究類"),
        ("這個對環境的影響是什麼？", True, "延伸探究類"),
        ("未來產業會怎麼因應這個政策？", True, "延伸探究類"),
        ("在生活中有沒有實際的案例？", True, "延伸探究類"),
        # ---- 概念理解類（≥3）
        ("為什麼原子會結合成分子？", True, "概念理解類"),
        ("什麼是原子", True, "概念理解類"),           # 句首疑問詞
        ("下列何者正確？", True, "概念理解類"),
        ("有什麼不同？", True, "概念理解類"),
        ("Why?", True, "概念理解類"),
        ("Why important?", True, "概念理解類"),
        ("這樣對嗎", True, "概念理解類"),             # 句尾「嗎」→ 走後備規則
        ("我想知道為什麼會這樣", True, "概念理解類"),
    ]
    problems = []
    n_q = n_other = 0
    per_cat = {}
    for s, want_q, want_c in cases:
        got_q = AN.is_question_sentence(s)
        if got_q != want_q:
            problems.append(f"提問偵測錯誤（應為 {want_q}）：{s}")
            continue
        if not got_q:
            continue
        n_q += 1
        got_c = AN.classify_question(s)[0]
        per_cat[got_c] = per_cat.get(got_c, 0) + 1
        if got_c == AN.OTHER_TYPE:
            n_other += 1
        if got_c != want_c:
            problems.append(f"分類錯誤（應為 {want_c}，得到 {got_c}）：{s}")
    want_cats = ["計算數據類", "操作應用類", "延伸探究類", "概念理解類"]
    if list(AN.QUESTION_TYPES.keys()) != want_cats:
        problems.append(f"提問類別應該是四類 {want_cats}，實際 "
                        f"{list(AN.QUESTION_TYPES.keys())}")
    for c in want_cats:
        if per_cat.get(c, 0) < 3:
            problems.append(f"{c} 的正確分類只有 {per_cat.get(c, 0)} 個，規格要求 ≥3")
    if n_q and n_other * 100.0 / n_q > 20:
        problems.append(f"「其他」佔 {n_other * 100.0 / n_q:.0f}%，超過兩成")
    log(f"  提問規則（四類）：{len(cases)} 個案例，問句 {n_q} 個，"
        + "、".join(f"{c} {per_cat.get(c, 0)} 個" for c in want_cats)
        + f"，其他 {n_other} 個，錯誤 {len(problems)} 個")
    return problems


def _selftest_tokens(log):
    """2.3 斷詞規則單元測試（合成句子，無任何真實個資）。

    1. 題幹／客套用語（下列／何者／關於／請問／同學）不可以出現在關鍵字裡
    2. 所有關鍵字的中文字數都在 min_word_len–max_word_len 之間（白名單除外）
    3. user_words 的詞就算超過 6 字也要保留
    4. 代詞／助詞（我們／的／了）不可以出現
    5. load_dict() 重新呼叫後斷詞快取要清掉（改了停用詞，結果要跟著變）
    """
    from core import analyze as AN
    problems = []

    AN.load_dict()                     # 回到內建預設（不吃使用者 config）
    texts = [
        "下列何者正確？關於綠色化學十二原則，請問同學有什麼想法？",
        "我們覺得生成式AI的提示詞很重要，pH 值也要會看。",
        "實驗時要先稀釋再中和，接著觀察氧化還原反應的顏色變化。",
        "這一題的答案我寫在心得裡，謝謝老師的說明。",
    ]
    toks = []
    for t in texts:
        toks += AN.tokens(t)

    banned = ["下列", "何者", "關於", "請問", "同學"]
    for b in banned:
        if b in toks:
            problems.append(f"停用詞「{b}」仍出現在關鍵字裡")
    for p in ["我們", "的", "了", "這樣", "可以"]:
        if p in toks:
            problems.append(f"代詞／助詞「{p}」仍出現在關鍵字裡")

    bad_len = [w for w in toks
               if w not in AN.ALWAYS_KEEP
               and AN.zh_len(w) and not (AN.MIN_LEN <= AN.zh_len(w) <= AN.MAX_LEN)]
    if bad_len:
        problems.append(f"這些關鍵字的中文字數不在 {AN.MIN_LEN}–{AN.MAX_LEN} 之間："
                        f"{sorted(set(bad_len))}")

    # 白名單超長詞：內建的「綠色化學十二原則」（8 字）要留著
    if "綠色化學十二原則" not in toks:
        problems.append("白名單的長詞「綠色化學十二原則」被長度規則砍掉了")
    # 化學動詞（v）要留下來
    for v in ("稀釋", "中和", "氧化還原"):
        if v not in toks:
            problems.append(f"化學操作詞「{v}」沒有被保留（詞性過濾把動詞砍掉了？）")
    # nr（人名）不得列入保留詞性——學生姓名漏遮罩時的隱私防線
    if "nr" in AN.KEEP_POS or "nr" in AN.POS_KEEP:
        problems.append("KEEP_POS 不可以包含 nr（人名），會把漏遮罩的學生姓名當成概念")

    # user_words 的超長詞（> 6 字）要保留
    long_word = "一個超過六個字的專有名詞"
    AN.load_dict(user_words=[long_word])
    if long_word not in AN.tokens(f"我們在報告裡討論了{long_word}的應用。"):
        problems.append(f"user_words 的長詞「{long_word}」沒有被保留")

    # 快取有沒有跟著 load_dict 清掉：改停用詞，同一段文字結果要不一樣
    probe = "原子經濟性和廢棄物減量都很重要"
    AN.load_dict()
    before = AN.tokens(probe)
    AN.load_dict(stopwords_extra=["原子經濟性"])
    after = AN.tokens(probe)
    if "原子經濟性" not in before:
        problems.append("快取測試的基準有問題：預設沒有切出「原子經濟性」")
    if "原子經濟性" in after:
        problems.append("load_dict() 之後斷詞快取沒有清掉（改了停用詞結果卻沒變）")
    # stopwords_remove：把剛加的停用詞再拿掉，應該又回來
    AN.load_dict(stopwords_extra=["原子經濟性"], stopwords_remove=["原子經濟性"])
    if "原子經濟性" not in AN.tokens(probe):
        problems.append("stopwords_remove 沒有把停用詞拿掉")

    AN.load_dict()                     # 收尾：恢復內建預設
    log(f"  斷詞規則（2.3）：{len(texts)} 句合成文字 → {len(toks)} 個關鍵字，"
        f"長度 {AN.MIN_LEN}–{AN.MAX_LEN} 字、詞性 {len(AN.POS_KEEP)} 類、"
        f"停用詞 {len(AN.STOPWORDS)} 個，錯誤 {len(problems)} 個")
    log(f"       範例關鍵字：{'、'.join(dict.fromkeys(toks))}")
    return problems


def _selftest_compat(tmp, cfg, log):
    """向下相容檢查：直接餵一份「原始 Zuvio 匯出」格式（學號是 int、姓名沒遮罩），
    確認程式會自己補學生編號、把學號改成 O、把自由文字裡的同學姓名換成編號。"""
    from core import parse_zuvio as PZ
    from core.mask import SID9_IN_TEXT
    import openpyxl

    names = ["甲小明", "乙小華", "丙大文"]
    sids = [990000041, 990000042, 990000043]        # 刻意用 int，複刻真實匯出（虛構學號）
    ans = ["原子經濟性讓我很有感，反應物要盡量都變成產物。為什麼一定要這樣算？",
           "我和甲小明討論過，實驗室要先算好用量再倒試劑。",
           "廢棄物減量最重要，源頭減量比事後回收有效。這個公式怎麼算？"]

    path = os.path.join(tmp, "compat_raw.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["資料夾名稱", "相容測試"])
    ws.append(["問題題型", "問答"])
    ws.append(["題幹", "向下相容測試題"])
    ws.append([])
    ws.append(["問題類型", "", "", "第1題:問答題"])
    ws.append(["問題敘述", "", "", "向下相容測試題"])
    ws.append([])
    ws.append(["學號", "姓名", "作答時間", "回答"])
    for i, n in enumerate(names):
        ws.append([sids[i], n, "2026-09-08 19:00:00", ans[i]])
    wb.save(path)

    rows, rec = PZ.parse_file(path, 1, semester="115-1", course="GC")
    problems = []
    if rec.get("學生編號來源", "").startswith("A 欄"):
        problems.append("相容模式誤判成已有學生編號欄")
    codes = {r[3] for r in rows}
    if len(codes) != 3 or not all(c.startswith("115-1_GC_") for c in codes):
        problems.append(f"相容模式的學生編號不對：{sorted(codes)}")
    for r in rows:
        if r[4] != "OOOOOOOOO":
            problems.append(f"相容模式沒有把學號改成全 O：{r[4]}")
            break
    blob = "\n".join(r[7] for r in rows) + "\n" + "\n".join(r[5] for r in rows)
    for n in names:
        if n in blob:
            problems.append(f"相容模式沒有換掉自由文字／姓名欄的「{n}」")
    if SID9_IN_TEXT.search(blob):
        problems.append("相容模式的輸出仍有 9 碼學號")
    if "115-1_GC_1" not in blob:
        problems.append("相容模式沒有把自由文字裡的同學姓名換成學生編號")
    log(f"  向下相容（原始 Zuvio 匯出、學號 int）：學生編號 {sorted(codes)}、"
        f"學號欄 {rows[0][4]}、姓名欄 {rows[0][5]}")
    return problems


# ------------------------------------------------------------------ 2.1 名單自我測試
# 以下全部是**合成資料**（虛構姓名與學號），不是真實學生。
_R_SEM = "115-2"
_R_CRS = "ZZ"
_R_STUDENTS = [
    (1, "990000001", "甲小明", "自資系大三"),
    (2, "990000002", "乙小華", "自資系大三"),
    (3, "990000003", "丙大文", "自資系大二"),
    (4, "990000004", "丁美玲", "自資系大二"),
    (5, "990000005", "戊建國", "資工系大二"),
    (6, "990000006", "己淑芬", "自資系大二"),
]
_R_OUTSIDER = ("990000099", "庚志強")        # 名單外作答者（退選），應該拿 101


def _write_roster_xlsx(path, with_seq=True, shuffled=False):
    """合成 Excel 名單：with_seq=False 時沒有序號欄（要改用學號排序規則）。"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "名單"
    ws.append(["國立東華大學合成名單（虛構資料）"])
    ws.append([])
    ws.append((["序號"] if with_seq else [None]) + ["學號", "姓名", "班級"])
    rows = list(_R_STUDENTS)
    if shuffled:
        rows = rows[::-1]
    for seq, sid, name, klass in rows:
        ws.append(([seq] if with_seq else [None]) + [int(sid), name, klass])
    wb.save(path)
    wb.close()
    return path


def _write_raw_zuvio(path, title, people, answers):
    """合成**原始 Zuvio 匯出檔**（沒有 A 欄學生編號、學號是 int、姓名沒遮罩）。"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["資料夾名稱", title])
    ws.append(["問題題型", "問答"])
    ws.append(["題幹", title])
    ws.append([])
    ws.append(["問題類型", "", "", "第1題:問答題"])
    ws.append(["問題敘述", "", "", title])
    ws.append([])
    ws.append(["學號", "姓名", "作答時間", "回答"])
    for (sid, name), ans in zip(people, answers):
        ws.append([int(sid), name, "2026-09-14 20:00:00", ans])
    ws.append([])
    ws.append(["未作答學生"])
    ws.append(["學號", "姓名"])
    answered = {s for s, _n in people}
    for _seq, sid, name, _k in _R_STUDENTS:
        if sid not in answered:
            ws.append([int(sid), name])
    wb.save(path)
    wb.close()
    return path


def _selftest_roster(tmp, cfg, log):
    """SPEC v2.1 §T2 的名單驗收（全部用合成資料）。"""
    from core import roster as R
    problems = []
    rcfg = dict(cfg)
    rcfg.update({"semester": _R_SEM, "course_code": _R_CRS,
                 "course_name": "合成測試課程", "appendix_matrix": True})

    # ---- 1. PDF 行解析：三種排法都要讀得出來
    line_cases = {
        "黏在一起（真實東華選課名單）": ["甲小明990000001 1 自資系大三",
                             "乙小華 【註*】99000000210 自資系大三"],
        "空白分隔": ["甲小明 990000001 1 自資系大三",
                 "乙小華 【註*】 990000002 2 自資系大三"],
        "逐行分開": ["甲小明", "990000001", "1", "自資系大三",
                 "乙小華", "【註*】", "990000002", "2", "自資系大三"],
    }
    for label, lines in line_cases.items():
        got = R.parse_pdf_lines(lines)
        if len(got) != 2 or got[0]["學號"] != "990000001" or got[0]["姓名"] != "甲小明":
            problems.append(f"PDF 行解析（{label}）失敗：{got}")

    # 2.2：完全黏連（姓名／學號／序號中間不留空白）＋ 含英文字母的學號
    #      ＋ 跨行被切斷的【註*】＋ 讀不到姓名的那一列
    glued_lines = [
        "班    級 姓      名 序號 學    號 1 2 3 4 5 6 7 8 9 10 出席 平時",
        "甲小明9900000011 自資系大三",            # 姓名＋9 碼學號＋1 碼序號 黏在一起
        "乙小華 【註*】9900000022 自資系大三",     # 註記夾在中間
        "丙大文99000A0033 資工系大二",            # 學號第 6 碼是英文字母
        "丁美玲A990000044 自資系大二",            # 學號＝1 字母＋8 數字
        "9900000055 自資系大三",                  # 這一列讀不到姓名
        "戊小強 Ming",                            # 姓名折行（中文名＋英文名前半）
        "Wu9900000066 環科系碩一",                # …英文名後半＋學號＋序號
        "註*:前一學期GPA平均未達2.0者。 共1頁,第1頁2026/9/15 列印 12:59:02",
    ]
    glued_want = [
        (1, "990000001", "甲小明"), (2, "990000002", "乙小華"),
        (3, "99000A003", "丙大文"), (4, "A99000004", "丁美玲"),
        (5, "990000005", ""), (6, "990000006", "戊小強 Ming Wu"),
    ]
    got_g = R.parse_pdf_lines(glued_lines)
    if len(got_g) != len(glued_want):
        problems.append(f"PDF 行解析（黏連格式）應讀到 {len(glued_want)} 人，"
                        f"實際 {len(got_g)}")
    else:
        for g, (seq, sid, name) in zip(got_g, glued_want):
            if (g["序號"], g["學號"], g["姓名"]) != (seq, sid, name):
                problems.append("PDF 行解析（黏連格式）不符："
                                f"{(g['序號'], g['學號'], g['姓名'])} != "
                                f"{(seq, sid, name)}")
    log(f"  PDF 行解析（黏連格式）：{len(got_g)} 人，含字母學號、缺姓名列、姓名折行")
    noise = R.parse_pdf_lines([
        "班    級 姓      名序號 學    號 1 2 3 4 5 6 7 8 9 10 平時 期中 期未 成績總計 備註",
        "國立東華大學115學年度第1學期選課名單",
        "NRES20050環境化學科目名稱: 科目代號:",
        "註*:前一學期GPA平均未達2.0之學生列為『學習追蹤對象』。 第1頁,共1頁2026/9/15 下午 12:59:02",
    ])
    if noise:
        problems.append(f"PDF 表頭／頁尾雜訊沒有被濾掉：{noise}")
    log(f"  PDF 行解析：{len(line_cases)} 種排法都讀得出來，雜訊行 0 誤判")

    # ---- 2. 端對端：合成 PDF 名單檔
    pdf = resource_path("selftest", "樣本_選課名單_合成.pdf")
    if os.path.exists(pdf):
        try:
            cb_pdf = R.prepare_codebook(pdf, _R_SEM, _R_CRS,
                                        out_dir=os.path.join(tmp, "cb_pdf"),
                                        log=lambda *a: None)
            if cb_pdf.count != len(_R_STUDENTS):
                problems.append(f"合成 PDF 名單讀到 {cb_pdf.count} 人，應該是 "
                                f"{len(_R_STUDENTS)} 人")
            if cb_pdf.by_sid.get("990000003") != f"{_R_SEM}_{_R_CRS}_3":
                problems.append(f"合成 PDF 名單的編號不對：{cb_pdf.by_sid.get('990000003')}")
            log(f"  合成 PDF 名單（端對端）：{cb_pdf.count} 人，"
                f"990000003 → {cb_pdf.by_sid.get('990000003')}")
        except R.RosterError as e:
            problems.append(f"合成 PDF 名單解析失敗：{e}")
    else:
        log("  [略過] 找不到 selftest\\樣本_選課名單_合成.pdf，只測行解析")

    # ---- 3. Excel 名單：有序號 → 依序號；無序號 → 依學號排序
    p1 = _write_roster_xlsx(os.path.join(tmp, "名單_有序號.xlsx"), with_seq=True)
    cb1 = R.prepare_codebook(p1, _R_SEM, _R_CRS, out_dir=os.path.join(tmp, "cb1"),
                             log=lambda *a: None)
    for seq, sid, _n, _k in _R_STUDENTS:
        if cb1.by_sid.get(sid) != f"{_R_SEM}_{_R_CRS}_{seq}":
            problems.append(f"有序號名單編號錯：{sid} → {cb1.by_sid.get(sid)}")
    p2 = _write_roster_xlsx(os.path.join(tmp, "名單_無序號.xlsx"), with_seq=False,
                            shuffled=True)
    cb2 = R.prepare_codebook(p2, _R_SEM, _R_CRS, out_dir=os.path.join(tmp, "cb2"),
                             log=lambda *a: None)
    for seq, sid, _n, _k in _R_STUDENTS:          # 學號本來就是遞增，答案一樣
        if cb2.by_sid.get(sid) != f"{_R_SEM}_{_R_CRS}_{seq}":
            problems.append(f"無序號名單（學號排序）編號錯：{sid} → {cb2.by_sid.get(sid)}")
    if "學號" not in cb2.rule:
        problems.append(f"無序號名單應該走「依學號字串遞增排序」，實際是：{cb2.rule}")
    log(f"  Excel 名單：有序號 → {cb1.rule}；無序號 → {cb2.rule}（各 {cb1.count} 人）")

    # ---- 4. 兩個原始 Zuvio 合成檔（出現順序不同）＋名單外作答者
    in_dir = os.path.join(tmp, "raw_in")
    os.makedirs(in_dir, exist_ok=True)
    o_sid, o_name = _R_OUTSIDER
    q1_people = [("990000003", "丙大文"), ("990000001", "甲小明"),
                 (o_sid, o_name), ("990000002", "乙小華")]
    q1_ans = [
        "原子經濟性很重要，反應物要盡量都變成產物，不要變成廢棄物。我和丁美玲討論過這題。",
        "綠色化學十二原則裡我最有感的是源頭減量，先算好用量再倒試劑。為什麼一定要這樣算？",
        "我覺得綠色溶劑可以取代有機溶劑，實驗室會比較安全。",
        "廢棄物減量最重要，源頭減量比事後回收有效，這個公式怎麼算？",
    ]
    q2_people = [("990000006", "己淑芬"), (o_sid, o_name),
                 ("990000001", "甲小明"), ("990000003", "丙大文")]
    q2_ans = [
        "水質檢測要先看pH值，再看溶氧量，這樣才知道水污染的程度。",
        "生成式AI可以幫忙整理文獻，但是資料還是要自己查證。",
        "原子經濟性的計算讓我更了解反應效率，綠色化學真的要從設計開始。",
        "我想知道為什麼廢棄物減量要放在十二原則的第一條？",
    ]
    _write_raw_zuvio(os.path.join(in_dir, "Q1_合成問答一.xlsx"), "合成問答一",
                     q1_people, q1_ans)
    _write_raw_zuvio(os.path.join(in_dir, "Q2_合成問答二.xlsx"), "合成問答二",
                     q2_people, q2_ans)

    # 4a. 未給名單 → 必須被擋下
    out_blocked = os.path.join(tmp, "out_blocked")
    try:
        run_pipeline(rcfg, in_dir, out_blocked, "合成週", "擋下測試.pptx",
                     thumbs=False, log=lambda *a: None)
        problems.append("原始 Zuvio 檔沒有名單竟然跑完了（§1.5 的擋下沒有生效）")
    except PipelineError as e:
        if "沒有名單" not in str(e) and "名單" not in str(e):
            problems.append(f"擋下的訊息看不出要先給名單：{e}")
        log("  未給名單：已正確擋下並提示要先載入名單／對照表")

    # 4b. --no-roster 放行（2.0 行為）
    one_dir = os.path.join(tmp, "raw_one")
    os.makedirs(one_dir, exist_ok=True)
    shutil.copy(os.path.join(in_dir, "Q1_合成問答一.xlsx"), one_dir)
    try:
        run_pipeline(rcfg, one_dir, os.path.join(tmp, "out_noroster"), "合成週",
                     "無名單測試.pptx", thumbs=False, log=lambda *a: None,
                     no_roster=True)
        log("  --no-roster：放行（依檔案內出現順序編號，2.0 行為）")
    except PipelineError as e:
        problems.append(f"--no-roster 應該放行卻被擋：{e}")

    # 4c. 正式跑一次（用合成名單產生的對照表）
    cb = R.prepare_codebook(p1, _R_SEM, _R_CRS, out_dir=os.path.join(tmp, "cb_run"),
                            log=lambda *a: None)
    out_dir = os.path.join(tmp, "out_roster")
    r = run_pipeline(rcfg, in_dir, out_dir, "合成週", "名單測試.pptx",
                     thumbs=False, log=lambda *a: None, codebook=cb)

    want = {sid: f"{_R_SEM}_{_R_CRS}_{seq}" for seq, sid, _n, _k in _R_STUDENTS}
    out_code = f"{_R_SEM}_{_R_CRS}_{R.OUTSIDE_START}"
    per_file = {}
    for fn in sorted(os.listdir(out_dir)):
        if not (fn.startswith("Q0") and fn.endswith(".csv")) or "_concept" in fn \
                or "_questions" in fn:
            continue
        import csv as _csv
        with open(os.path.join(out_dir, fn), encoding="utf-8-sig", newline="") as f:
            rows = list(_csv.DictReader(f))
        per_file[fn] = {r_["學生編號"] for r_ in rows if r_.get("學生編號")}
    if len(per_file) != 2:
        problems.append(f"應該產生 2 題 CSV，實際 {len(per_file)}：{list(per_file)}")

    # 編號與名單一致：Q1 出現的 4 個編號必須正好對上
    q1_expect = {want[s] if s in want else out_code for s, _n in q1_people}
    q2_expect = {want[s] if s in want else out_code for s, _n in q2_people}
    got = list(per_file.values())
    if got and got[0] != q1_expect:
        problems.append(f"第一題的學生編號與名單不符：{sorted(got[0])} != {sorted(q1_expect)}")
    if len(got) > 1 and got[1] != q2_expect:
        problems.append(f"第二題的學生編號與名單不符：{sorted(got[1])} != {sorted(q2_expect)}")
    # 兩題同人同號（甲小明、丙大文、名單外的庚志強都出現在兩題）
    if len(got) > 1:
        both = got[0] & got[1]
        for sid in ("990000001", "990000003"):
            if want[sid] not in both:
                problems.append(f"同一位同學在兩題的編號不一致：{want[sid]} 不在交集")
        if out_code not in both:
            problems.append(f"名單外作答者在兩題的編號不一致（應該都是 {out_code}）")

    # 全班人數 = 名單人數
    klass = {q["全班人數"] for q in r["analysis"]["題目"]}
    if klass != {len(_R_STUDENTS)}:
        problems.append(f"全班人數應該等於名單人數 {len(_R_STUDENTS)}，實際 {klass}")

    # 對照表已回寫（重新讀檔）
    cb_again = R.load_codebook(cb.path)
    if len(cb_again.outside) != 1 or cb_again.outside[0].code != out_code:
        problems.append(f"名單外作答者沒有正確回寫對照表：{len(cb_again.outside)} 筆")
    elif not cb_again.outside[0].src:
        problems.append("名單外作答者沒有記下「首次出現檔案」")

    # 附錄概念矩陣：學生依編號數字排序
    for q in r["analysis"]["題目"]:
        codes = list(q.get("概念矩陣", {}).keys())
        nums = [int(c.rsplit("_", 1)[1]) for c in codes if c.rsplit("_", 1)[-1].isdigit()]
        if nums != sorted(nums):
            problems.append(f"{q['題號']} 概念矩陣沒有依編號數字排序：{codes}")

    # 隱私：比對對照表內所有真名與學號
    hits = check_privacy(out_dir, codebook=cb_again, log=lambda *a: None)
    if hits:
        problems.append(f"名單模式的隱私檢查命中 {hits} 處")
        check_privacy(out_dir, codebook=cb_again, log=log)

    log(f"  名單模式：{len(per_file)} 題、全班人數 {sorted(klass)}、"
        f"名單外作答者 {out_code}、對照表回寫 {len(cb_again.outside)} 人、隱私命中 {hits} 處")
    return problems


def selftest(log=print, keep=False):
    """用內建合成資料（NameMasker 2.0 格式）跑完整流程並驗收。

    驗收項目：
      1. 簡報頁數 = 3（封面／總覽／結尾）＋ 題數（＋附錄頁，每頁 2 題、最多 3 頁）
      2. 每題都有 Q0N_concept_matrix.csv 與 Q0N_questions.csv
      3. 每個覆蓋率都介於 0–1，且「≥1 個 ≥ ≥3 個 ≥ 全部都提到」
      3b.**2.2**：每題開放題重點數 = min(6, 候選數)；概念矩陣欄數 = 1 + 重點數 + 1；
          候選不足 6 個時有幾個列幾個（不補空字串、不報錯）
      4. 幾何 QA：出界 0、重疊 0
      5. 隱私檢查 0 命中（名冊姓名＋9 碼學號＋email 樣式）
      6. 文字雲 ≥ 3 張
      7. 每一頁都有備忘稿逐字稿，且 ≤ 350 字
      7b.**2.2**：四類提問的分類單元測試（每類 ≥3 個問句＋陳述句干擾，其他 ≤20%）
      7c.**2.3**：斷詞規則單元測試（2–6 字、詞性過濾、停用詞、白名單、快取清除）
      8. 向下相容（原始 Zuvio 匯出格式）
      9. **2.1 名單流程**：PDF／Excel 名單解析、固定編號、兩題同人同號、
         名單外作答者 101 起且回寫對照表、全班人數＝名單人數、
         未給名單被擋、--no-roster 放行、隱私（含對照表真名學號）0 命中
    """
    from pptx import Presentation
    tmp = tempfile.mkdtemp(prefix="hwwc_selftest_[課程] ")   # 刻意含中括號與空白：回歸測試 glob 跳脫
    try:
        log(f"暫存資料夾：{tmp}")
        cfg = load_config()
        cfg["course_name"] = cfg.get("course_name") or "我的課程"
        r = run_demo(cfg, tmp, thumbs=False, log=log)

        out_dir = r["out_dir"]
        pptx = r["pptx"]
        result = r.get("analysis") or {}
        qs = result.get("題目", [])
        n_q = len(qs)
        problems = []

        # ---- 1. 頁數（2.2：附錄每頁 2 題、最多 3 頁）
        from core import deck as _DK
        n_apx_q = sum(1 for q in qs if q.get("重點概念"))
        n_apx = 0
        if bool(cfg.get("appendix_matrix", True)) and n_apx_q:
            n_apx = min(-(-n_apx_q // _DK.APX_PER_PAGE), _DK.APX_MAX_PAGES)
        expect = 3 + n_q + n_apx
        if not os.path.exists(pptx):
            problems.append("沒有產生 pptx")
            pages = 0
            prs = None
        else:
            prs = Presentation(pptx)
            pages = len(prs.slides)
        if pages != expect:
            problems.append(f"簡報 {pages} 頁，應該是 {expect} 頁"
                            f"（3 + 題數 {n_q}{f' + 附錄 {n_apx}' if n_apx else ''}）")

        # ---- 2. 每題兩個 CSV
        for q in qs:
            for fn in (f"{q['題號']}_concept_matrix.csv", f"{q['題號']}_questions.csv"):
                if not os.path.exists(os.path.join(out_dir, fn)):
                    problems.append(f"少了 {fn}")

        # ---- 3. 覆蓋率介於 0–1
        for q in qs:
            vals = [q.get("整體覆蓋率", 0), q.get("提到三個以上比例", 0),
                    q.get("全部提到比例", 0)] + \
                   [f["覆蓋率"] for f in q.get("重點概念", [])]
            for v in vals:
                if not (0.0 <= float(v) <= 1.0):
                    problems.append(f"{q['題號']} 覆蓋率 {v} 不在 0–1 之間")
            # 覆蓋率的單調關係：整體（≥1 個）≥ ≥3 個 ≥ 全部都提到
            if not (q.get("整體覆蓋率", 0) >= q.get("提到三個以上比例", 0)
                    >= q.get("全部提到比例", 0)):
                problems.append(f"{q['題號']} 三種覆蓋率不符合 ≥1 ≥ ≥3 ≥ 全部 的關係")

        # ---- 3b. 2.2：每題重點數 = min(6, 候選數)；概念矩陣欄數要對得上
        import csv as _csv2
        n_focus = 0
        for q in qs:
            focus = q.get("重點概念", [])
            cand = int(q.get("候選概念數", 0))
            if not focus:
                continue                       # 選擇題／測驗題不做重點概念
            want_n = min(6, cand)
            if len(focus) != want_n:
                problems.append(f"{q['題號']} 重點數 {len(focus)}，"
                                f"應該是 min(6, 候選 {cand}) = {want_n}")
            if len(focus) == 6:
                n_focus += 1
            cm = os.path.join(out_dir, f"{q['題號']}_concept_matrix.csv")
            if os.path.exists(cm):
                with open(cm, encoding="utf-8-sig", newline="") as f:
                    head = next(_csv2.reader(f), [])
                if len(head) != len(focus) + 2:
                    problems.append(f"{q['題號']}_concept_matrix.csv 有 {len(head)} 欄，"
                                    f"應該是學生編號 + {len(focus)} 概念 + 提到重點數")
                if head[:1] != ["學生編號"] or head[-1:] != ["提到重點數"]:
                    problems.append(f"{q['題號']}_concept_matrix.csv 欄位標題不對：{head}")
        if n_focus < 1:
            problems.append("沒有任何一題抓到 6 個重點概念")

        # ---- 3c. 2.2：候選不足 6 個時要「有幾個列幾個」，不補空字串、不報錯
        from core import analyze as _AN
        few = [{"學生編號": "115-1_ZZ_1", "作答內容": "原子經濟 廢棄物減量"},
               {"學生編號": "115-1_ZZ_2", "作答內容": "原子經濟 綠色溶劑"}]
        few_c, few_ps, _cf, few_spk = _AN.pick_concepts(few, top_n=6)
        if len(few_c) != min(6, len(few_spk)) or any(not c for c in few_c):
            problems.append(f"候選不足 6 個時重點數不對：{few_c}（候選 {len(few_spk)}）")
        _keys, _mat = _AN.concept_matrix(few_c, few_ps)
        if any(len(v) != len(few_c) for v in _mat.values()):
            problems.append("候選不足 6 個時概念矩陣欄數與重點數對不上")

        # ---- 4. 幾何 QA
        rep_p = os.path.join(out_dir, "qa", "layout_report.json")
        oob = ovl = -1
        if os.path.exists(rep_p):
            with open(rep_p, encoding="utf-8") as f:
                rep = json.load(f)
            oob, ovl = len(rep["out_of_bounds"]), len(rep["overlaps"])
            if oob:
                problems.append(f"版面出界 {oob} 處")
                for o in rep["out_of_bounds"][:5]:
                    log(f"    出界：{o}")
            if ovl:
                problems.append(f"版面重疊 {ovl} 處")
                for o in rep["overlaps"][:5]:
                    log(f"    重疊：{o}")
        else:
            problems.append("沒有產生 layout_report.json")

        # ---- 5. 文字雲
        wc_dir = os.path.join(out_dir, "wc")
        pngs = [f for f in os.listdir(wc_dir)] if os.path.isdir(wc_dir) else []
        pngs = [f for f in pngs if f.lower().endswith(".png")]
        if len(pngs) < 3:
            problems.append(f"文字雲只有 {len(pngs)} 張，少於 3 張")

        # ---- 6. 每頁備忘稿（一定要有，而且不能超過 350 字）
        if prs is not None:
            empty, toolong = [], []
            for i, s in enumerate(prs.slides, 1):
                t = (s.notes_slide.notes_text_frame.text
                     if s.has_notes_slide else "")
                if not t.strip():
                    empty.append(i)
                elif len(t) > 350:
                    toolong.append((i, len(t)))
            if empty:
                problems.append(f"第 {empty} 頁沒有備忘稿逐字稿")
            if toolong:
                problems.append(f"逐字稿超過 350 字：{toolong}")

        # ---- 6b. 提問偵測與分類規則
        log("")
        log("提問規則檢查")
        problems += _selftest_questions(log)

        # ---- 6c. 2.3：斷詞規則（長度／詞性／停用詞／白名單／快取）
        log("")
        log("斷詞規則檢查（2.3）")
        try:
            problems += _selftest_tokens(log)
        except Exception as e:
            problems.append(f"斷詞規則檢查失敗：{e}")
            log(traceback.format_exc())

        # ---- 7. 隱私
        hits = check_privacy(out_dir, log=log)
        if hits:
            problems.append(f"隱私檢查命中 {hits} 處")

        # ---- 7b.「學生編號連結姓名」工作表絕不可被讀進來
        leaked = _link_sheet_leak(sample_input_dir(), out_dir, log=log)
        if leaked:
            problems.append(f"「學生編號連結姓名」的內容外洩 {leaked} 處")

        # ---- 8. 向下相容（原始 Zuvio 匯出格式）
        log("")
        log("向下相容檢查")
        try:
            problems += _selftest_compat(tmp, cfg, log)
        except Exception as e:
            problems.append(f"向下相容檢查失敗：{e}")

        # ---- 9. 2.1 名單 → 固定學生編號 → 遮罩
        log("")
        log("名單流程檢查（2.1）")
        try:
            problems += _selftest_roster(tmp, cfg, log)
        except Exception as e:
            problems.append(f"名單流程檢查失敗：{e}")
            log(traceback.format_exc())

        log("")
        log(f"檢查結果：簡報 {pages}/{expect} 頁、題數 {n_q}、每題重點 6 個的有 {n_focus} 題、"
            f"文字雲 {len(pngs)} 張、出界 {oob} 處、重疊 {ovl} 處、隱私命中 {hits} 處")
        if problems:
            for p in problems:
                log(f"  [失敗] {p}")
            log("SELFTEST FAILED")
            return 1
        log("SELFTEST OK")
        return 0
    finally:
        if not keep:
            shutil.rmtree(tmp, ignore_errors=True)
